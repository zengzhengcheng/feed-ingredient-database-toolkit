from __future__ import annotations

import argparse
import csv
import json
import math
import re
import statistics
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SEVERITY_ORDER = {"BLOCKER": 0, "ERROR": 1, "WARNING": 2, "INFO": 3}
BLOCKING_SEVERITIES = {"BLOCKER", "ERROR"}
OWNERS = {"SOURCE_EXCEL", "IMPORTER_RULE", "DATABASE_QC", "MANUAL_REVIEW"}


@dataclass
class Issue:
    severity: str
    owner: str
    code: str
    collection: str
    workbook: str
    sheet: str = ""
    cell: str = ""
    message: str = ""
    evidence: str = ""
    suggested_action: str = ""


@dataclass
class WorkbookRecord:
    collection: str
    relative_path: str
    file_role: str
    workbook_status: str
    sheets: str
    parameter_sheet: str
    data_sheet: str
    title: str
    doi_raw: str
    doi_normalized: str
    ingredient_columns: int
    observation_cells: int
    issue_count: int = 0


@dataclass
class Observation:
    workbook: str
    collection: str
    sheet: str
    row: int
    column: int
    metric: str
    section: str
    subsection: str
    context: str
    unit: str
    basis: str
    value: Any


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def normalize_label(value: object) -> str:
    text = unicodedata.normalize("NFKC", clean(value)).casefold()
    for character in ("-", "–", "—", "_", "‘", "’"):
        text = text.replace(character, " ")
    return " ".join(text.split())


def normalize_title(value: object) -> str:
    text = unicodedata.normalize("NFKC", clean(value)).casefold()
    return "".join(character for character in text if character.isalnum())


def normalize_doi(value: object) -> str:
    text = unicodedata.normalize("NFKC", clean(value)).casefold().strip()
    text = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", text)
    text = re.sub(r"^doi\s*:\s*", "", text)
    text = re.sub(r"\s+", "", text)
    return text.rstrip(".,;。；")


def numeric(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = clean(value).replace(",", "")
    try:
        number = float(text)
        return number if math.isfinite(number) else None
    except ValueError:
        return None


def read_rules(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def find_sheet(workbook, accepted: Iterable[str]):
    accepted_names = {normalize_label(name) for name in accepted}
    return next((sheet for sheet in workbook.worksheets if normalize_label(sheet.title) in accepted_names), None)


def classify_file(path: Path, rules: dict[str, Any]) -> str:
    suffix = path.suffix.casefold()
    if path.name.startswith("~$") and suffix in {".xlsx", ".xlsm"}:
        return "excel_temporary_lock"
    if suffix in {".xlsx", ".xlsm"}:
        if any(marker in path.stem for marker in rules["raw_workbook_name_markers"]):
            return "raw_historical_workbook"
        return "curated_historical_workbook"
    if suffix == ".pdf":
        return "article_pdf_copy" if "副本" in path.stem else "article_pdf"
    if suffix in {".txt", ".md", ".csv"}:
        return "supporting_text"
    return "other_file"


def collection_name(root: Path, path: Path) -> str:
    relative = path.relative_to(root)
    return relative.parent.as_posix()


def metadata_from_sheet(sheet, rules: dict[str, Any]) -> tuple[str, str]:
    headers = [normalize_label(cell.value) for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]] if sheet.max_row >= 2 else []
    data = dict(zip(headers, values))
    title = next((clean(data.get(normalize_label(field))) for field in rules["title_fields"] if clean(data.get(normalize_label(field)))), "")
    doi = next((clean(data.get(normalize_label(field))) for field in rules["doi_fields"] if clean(data.get(normalize_label(field)))), "")
    return title, doi


def detect_unit(text: str, unit_tokens: list[str]) -> str:
    normalized = unicodedata.normalize("NFKC", text)
    for token in sorted(unit_tokens, key=len, reverse=True):
        if re.search(rf"(?<![A-Za-z]){re.escape(token)}(?![A-Za-z])", normalized, flags=re.I):
            return token
    return ""


def detect_basis(text: str, basis_tokens: list[str]) -> str:
    normalized = unicodedata.normalize("NFKC", text)
    for token in sorted(basis_tokens, key=len, reverse=True):
        if re.search(rf"(?<![A-Za-z]){re.escape(token)}(?![A-Za-z])", normalized, flags=re.I):
            return token
    return ""


def row_is_subsection(sheet, row: int, width: int) -> bool:
    if width >= 2 and normalize_label(sheet.cell(row, 2).value) == "data":
        return True
    values = [sheet.cell(row, column).value for column in range(2, width + 1)]
    return bool(values) and all(value is None or clean(value) == "" for value in values)


def measurement_context(section: str, subsection: str, rules: dict[str, Any]) -> str:
    normalized = normalize_label(subsection)
    if section == "能量":
        return "energy"
    if section == "营养物质消化率":
        return rules["digestibility_contexts"].get(normalized, normalized.upper()) if normalized else "ATTD"
    if section == "非表观营养价值":
        return rules["digestibility_contexts"].get(normalized, normalized.upper()) if normalized else ""
    return "composition"


def issue(severity: str, owner: str, code: str, collection: str, workbook: str, **kwargs: str) -> Issue:
    if severity not in SEVERITY_ORDER:
        raise ValueError(severity)
    if owner not in OWNERS:
        raise ValueError(owner)
    return Issue(severity, owner, code, collection, workbook, **kwargs)


def audit_workbook(root: Path, path: Path, rules: dict[str, Any]) -> tuple[WorkbookRecord, list[Issue], list[Observation]]:
    relative = path.relative_to(root).as_posix()
    collection = collection_name(root, path)
    issues: list[Issue] = []
    observations: list[Observation] = []
    try:
        formulas = load_workbook(path, read_only=False, data_only=False)
        values = load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:
        record = WorkbookRecord(collection, relative, "curated_historical_workbook", "unreadable", "", "", "", "", "", "", 0, 0)
        issues.append(issue("BLOCKER", "SOURCE_EXCEL", "WORKBOOK_UNREADABLE", collection, relative, message=f"工作簿无法读取：{type(exc).__name__}: {exc}", suggested_action="修复或重新导出该工作簿。"))
        return record, issues, observations

    try:
        sheet_names = [sheet.title for sheet in values.worksheets]
        parameter_sheet = find_sheet(values, rules["accepted_parameter_sheets"])
        data_sheet = find_sheet(values, rules["accepted_ingredient_sheets"])
        title = doi_raw = ""
        if parameter_sheet is None:
            issues.append(issue("BLOCKER", "SOURCE_EXCEL", "MISSING_METADATA_SHEET", collection, relative, message="缺少可识别的论文元数据 sheet。", suggested_action="增加“论文参数”sheet，并填写至少 title。"))
        else:
            title, doi_raw = metadata_from_sheet(parameter_sheet, rules)
            if not title:
                issues.append(issue("BLOCKER", "SOURCE_EXCEL", "MISSING_TITLE", collection, relative, sheet=parameter_sheet.title, cell="row 2", message="论文题目为空；题目是必填身份字段。", suggested_action="补录论文题目；DOI 可以为空。"))
        doi_normalized = normalize_doi(doi_raw)

        for formula_sheet in formulas.worksheets:
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "e":
                        issues.append(issue("ERROR", "SOURCE_EXCEL", "EXCEL_ERROR_CELL", collection, relative, sheet=formula_sheet.title, cell=cell.coordinate, message="单元格保存为 Excel 错误值。", evidence=clean(cell.value), suggested_action="回到源表修复公式或错误值。"))
                    elif cell.data_type == "f":
                        cached = values[formula_sheet.title][cell.coordinate].value
                        if cached is None:
                            issues.append(issue("WARNING", "SOURCE_EXCEL", "FORMULA_WITHOUT_CACHED_VALUE", collection, relative, sheet=formula_sheet.title, cell=cell.coordinate, message="公式没有可读取的缓存结果。", evidence=clean(cell.value), suggested_action="用 Excel/LibreOffice 重算并保存，或改为经核对的固定值。"))

        ingredient_count = 0
        if data_sheet is None:
            issues.append(issue("BLOCKER", "SOURCE_EXCEL", "MISSING_DATA_SHEET", collection, relative, message="缺少可识别的成分及营养价值数据 sheet。", suggested_action="增加或重命名为“成分及营养价值数据”。"))
        else:
            width = data_sheet.max_column
            last_nonblank_header = max((column for column in range(2, width + 1) if clean(data_sheet.cell(1, column).value)), default=1)
            ingredient_columns = [column for column in range(2, last_nonblank_header + 1) if clean(data_sheet.cell(1, column).value)]
            ingredient_count = len(ingredient_columns)
            if not ingredient_columns:
                issues.append(issue("BLOCKER", "SOURCE_EXCEL", "NO_INGREDIENT_COLUMNS", collection, relative, sheet=data_sheet.title, cell="row 1", message="没有识别到原料列。", suggested_action="第1行第2列起填写原料名称。"))
            for column in range(2, last_nonblank_header + 1):
                if not clean(data_sheet.cell(1, column).value):
                    has_values = any(clean(data_sheet.cell(row, column).value) for row in range(2, data_sheet.max_row + 1))
                    if has_values:
                        issues.append(issue("ERROR", "SOURCE_EXCEL", "BLANK_INGREDIENT_HEADER", collection, relative, sheet=data_sheet.title, cell=f"{get_column_letter(column)}1", message="列标题为空但列内有数据。", suggested_action="补录原料名称或清除误填数据。"))
            names = defaultdict(list)
            for column in ingredient_columns:
                names[normalize_label(data_sheet.cell(1, column).value)].append(column)
            for normalized_name, columns in names.items():
                if normalized_name and len(columns) > 1:
                    issues.append(issue("WARNING", "MANUAL_REVIEW", "DUPLICATE_INGREDIENT_NAME", collection, relative, sheet=data_sheet.title, cell=" | ".join(f"{get_column_letter(column)}1" for column in columns), message="同一工作簿含重复原料名称；可能是不同样品，也可能是重复列。", evidence=clean(data_sheet.cell(1, columns[0]).value), suggested_action="依据原料编号和论文表格人工确认。"))

            section = "基础数据"
            subsection = ""
            reported_basis = {column: "" for column in ingredient_columns}
            source_number_by_column: dict[int, str] = {}
            trial_number_by_column: dict[int, str] = {}
            metric_rows: dict[tuple[int, str, str], list[Observation]] = defaultdict(list)
            for row in range(2, data_sheet.max_row + 1):
                raw_label = clean(data_sheet.cell(row, 1).value)
                if not raw_label:
                    continue
                normalized = normalize_label(raw_label)
                if normalized in {normalize_label(key) for key in rules["section_labels"]}:
                    canonical_sections = {normalize_label(key): value for key, value in rules["section_labels"].items()}
                    section = canonical_sections[normalized]
                    subsection = ""
                    continue
                if row_is_subsection(data_sheet, row, max(last_nonblank_header, 2)):
                    subsection = raw_label
                    continue
                unit = detect_unit(" | ".join(filter(None, [raw_label, subsection, section])), rules["unit_tokens"])
                label_basis = detect_basis(" | ".join(filter(None, [raw_label, subsection, section])), rules["basis_tokens"])
                context = measurement_context(section, subsection, rules)
                for column in ingredient_columns:
                    value = data_sheet.cell(row, column).value
                    if value is None or clean(value) == "" or normalize_label(value) == "data":
                        continue
                    cell = f"{get_column_letter(column)}{row}"
                    if section == "基础数据":
                        if normalized in {"dry base", "dry matter basis", "basis", "基础"}:
                            reported_basis[column] = detect_basis(clean(value), rules["basis_tokens"]) or clean(value)
                        if normalized in {"原料编号", "ingredient no", "ingredient number"}:
                            source_number_by_column[column] = clean(value)
                        elif normalized in {"试验编号", "trial no", "trial number"}:
                            trial_number_by_column[column] = clean(value)
                        continue
                    basis = label_basis or reported_basis[column]
                    obs = Observation(relative, collection, data_sheet.title, row, column, raw_label, section, subsection, context, unit, basis, value)
                    observations.append(obs)
                    metric_rows[(row, section, subsection)].append(obs)
                    if isinstance(value, float) and not math.isfinite(value):
                        issues.append(issue("ERROR", "SOURCE_EXCEL", "NONFINITE_NUMBER", collection, relative, sheet=data_sheet.title, cell=cell, message="数值不是有限数。", evidence=repr(value), suggested_action="核对并改为原文中的有效值或空值。"))

            reused_numbers: dict[str, list[int]] = defaultdict(list)
            for column, source_no in source_number_by_column.items():
                reused_numbers[normalize_label(source_no)].append(column)
            reused_numbers = {key: columns for key, columns in reused_numbers.items() if key and len(columns) > 1}
            if reused_numbers:
                evidence = "; ".join(
                    f"{key}:" + ",".join(
                        f"{get_column_letter(column)}(trial={trial_number_by_column.get(column, '') or 'blank'})"
                        for column in columns
                    )
                    for key, columns in sorted(reused_numbers.items())
                )
                issues.append(issue("INFO", "IMPORTER_RULE", "REUSED_INGREDIENT_NUMBER", collection, relative, sheet=data_sheet.title, message="原料编号在不同列中重复；该编号不是跨试验或跨处理的全局唯一键。", evidence=evidence, suggested_action="入库时用来源工作簿、源列、试验编号和原料编号组成可追溯身份，不得仅按原料编号合并。"))

            explicit_metric_units: set[tuple[int, str]] = set()
            missing_unit_metrics: set[tuple[int, str, str]] = set()
            missing_basis_metrics: set[tuple[int, str, str]] = set()
            for obs in observations:
                if obs.unit:
                    explicit_metric_units.add((obs.row, obs.metric))
                elif obs.context in {"energy", "AID", "SID", "ATTD", "STTD"}:
                    missing_unit_metrics.add((obs.row, obs.metric, obs.context))
                if not obs.basis and obs.context in {"composition", "energy"}:
                    missing_basis_metrics.add((obs.row, obs.metric, obs.context))
            for row, metric, context in sorted(missing_unit_metrics):
                issues.append(issue("WARNING", "IMPORTER_RULE", "UNIT_NOT_EXPLICIT", collection, relative, sheet=data_sheet.title, cell=f"A{row}", message="指标及相邻结构标签未明确写出单位。", evidence=f"context={context}; metric={metric}", suggested_action="优先查历史指标/单位字典；仍无法确定时回查原文并保持缺失，禁止猜测。"))
            for row, metric, context in sorted(missing_basis_metrics):
                issues.append(issue("WARNING", "SOURCE_EXCEL", "BASIS_NOT_EXPLICIT", collection, relative, sheet=data_sheet.title, cell=f"A{row}", message="成分或能量值的计量基础未明确。", evidence=f"context={context}; metric={metric}", suggested_action="依据原文补录 dry matter/as-fed；原文未说明时保持缺失。"))
            for row, metric in sorted(explicit_metric_units):
                issues.append(issue("INFO", "IMPORTER_RULE", "UNIT_EMBEDDED_IN_LABEL", collection, relative, sheet=data_sheet.title, cell=f"A{row}", message="单位可从指标或结构标签中提取，导入器必须保留并规范化。", evidence=metric, suggested_action="导入时拆分 raw_metric 与 raw_unit，但保留原始标签和定位。"))

            for (_, _, _), row_observations in metric_rows.items():
                numbers = [(obs, numeric(obs.value)) for obs in row_observations]
                numbers = [(obs, number) for obs, number in numbers if number is not None]
                finite_values = [abs(number) for _, number in numbers if number != 0]
                if len(finite_values) >= 3:
                    median = statistics.median(finite_values)
                    if median > 0:
                        for obs, number in numbers:
                            if abs(number) >= median * 100 and abs(number) - median >= 100:
                                severity = "BLOCKER" if obs.context == "energy" and median < 100 and abs(number) > 1000 else "WARNING"
                                owner = "SOURCE_EXCEL" if severity == "BLOCKER" else "MANUAL_REVIEW"
                                code = "ENERGY_SCALE_OUTLIER" if severity == "BLOCKER" else "ROW_SCALE_OUTLIER"
                                issues.append(issue(severity, owner, code, collection, relative, sheet=obs.sheet, cell=f"{get_column_letter(obs.column)}{obs.row}", message="同一指标行出现至少100倍量级离群值，疑似小数点、单位、处理效应或列位差异。", evidence=f"metric={obs.metric}; value={number:g}; row_median={median:g}", suggested_action="对照原始论文表格核对；审查器只定位，不自动更正。"))
                for obs, number in numbers:
                    metric_normalized = normalize_label(obs.metric)
                    unit_normalized = normalize_label(obs.unit)
                    cell = f"{get_column_letter(obs.column)}{obs.row}"
                    is_energy = obs.context == "energy" or any(token == metric_normalized or token in metric_normalized.split() for token in rules["energy_metric_tokens"])
                    if is_energy and unit_normalized == normalize_label("MJ/kg") and abs(number) > 100:
                        issues.append(issue("BLOCKER", "SOURCE_EXCEL", "ENERGY_MJ_IMPLAUSIBLE", collection, relative, sheet=obs.sheet, cell=cell, message="以 MJ/kg 标示的能量值超过100，疑似小数点或单位错误。", evidence=f"metric={obs.metric}; value={number:g}; unit={obs.unit}", suggested_action="对照原文核对，禁止程序自动缩放。"))
                    digestibility_metric = any(token in metric_normalized for token in ("digestibility", "coefficient", "消化率"))
                    if obs.context in {"AID", "SID", "ATTD", "STTD"} and abs(number) > 100:
                        if unit_normalized == normalize_label("%") or digestibility_metric:
                            issues.append(issue("ERROR", "SOURCE_EXCEL", "DIGESTIBILITY_PERCENT_OVER_100", collection, relative, sheet=obs.sheet, cell=cell, message="明确作为消化率/百分比的数据绝对值超过100。", evidence=f"context={obs.context}; metric={obs.metric}; value={number:g}; unit={obs.unit}", suggested_action="核对单位、百分号、小数位及列对应关系。"))
                        else:
                            issues.append(issue("WARNING", "MANUAL_REVIEW", "DIGESTIBILITY_CONTEXT_VALUE_OVER_100", collection, relative, sheet=obs.sheet, cell=cell, message="AID/SID/ATTD/STTD 区域数值超过100，但单位未证明它是百分比；也可能是可消化含量。", evidence=f"context={obs.context}; metric={obs.metric}; value={number:g}", suggested_action="先确认指标含义和单位，再决定是否修正；不得仅因超过100自动判错。"))
                    if obs.context == "composition" and unit_normalized == normalize_label("%") and (number < 0 or number > 100):
                        issues.append(issue("ERROR", "SOURCE_EXCEL", "PERCENT_OUT_OF_RANGE", collection, relative, sheet=obs.sheet, cell=cell, message="百分比成分值超出0–100。", evidence=f"metric={obs.metric}; value={number:g}", suggested_action="核对原值、小数位和单位。"))

        record = WorkbookRecord(collection, relative, "curated_historical_workbook", "recognized" if data_sheet and parameter_sheet else "unrecognized", " | ".join(sheet_names), parameter_sheet.title if parameter_sheet else "", data_sheet.title if data_sheet else "", title, doi_raw, doi_normalized, ingredient_count, len(observations))
        return record, issues, observations
    finally:
        formulas.close()
        values.close()


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    names = fieldnames or (list(rows[0]) if rows else [])
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=names)
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def markdown_report(summary: dict[str, Any], issues: list[Issue]) -> str:
    counts = summary["issue_counts_by_severity"]
    owner_counts = summary["issue_counts_by_owner"]
    lines = [
        "# 原始工作簿独立审查报告",
        "",
        f"- 审查结论：**{'允许进入建库流程' if summary['can_build_database'] else '阻断建库'}**",
        f"- 原始目录：`{summary['source_root']}`",
        f"- 审查工作簿：{summary['curated_workbooks']} 个",
        f"- 唯一规范化题目：{summary['unique_normalized_titles']} 个",
        f"- 唯一规范化 DOI：{summary['unique_normalized_dois']} 个（DOI 可空）",
        f"- 观察值单元格：{summary['observation_cells']} 个",
        f"- 问题：BLOCKER {counts.get('BLOCKER', 0)}；ERROR {counts.get('ERROR', 0)}；WARNING {counts.get('WARNING', 0)}；INFO {counts.get('INFO', 0)}",
        "",
        "## 问题归属",
        "",
    ]
    for owner in ("SOURCE_EXCEL", "IMPORTER_RULE", "DATABASE_QC", "MANUAL_REVIEW"):
        lines.append(f"- {owner}: {owner_counts.get(owner, 0)}")
    lines.extend(["", "## 阻断项与错误", ""])
    blocking = [item for item in issues if item.severity in BLOCKING_SEVERITIES]
    if not blocking:
        lines.append("无。")
    else:
        for item in blocking[:200]:
            location = "!".join(part for part in (item.workbook, item.sheet, item.cell) if part)
            lines.append(f"- [{item.severity}] `{item.code}` — `{location}` — {item.message} {item.evidence}".rstrip())
        if len(blocking) > 200:
            lines.append(f"- 其余 {len(blocking) - 200} 项见 `all_issues.csv`。")
    lines.extend([
        "",
        "## 门禁解释",
        "",
        "只有在不存在 BLOCKER/ERROR，且不存在未解决的题目或 DOI 重复时，`release_gate.json` 中的 `can_build_database` 才为 true。WARNING 不自动阻断，但需要按台账决定是否修正。审查器不修改 Excel、不建立数据库、不解析或发布论文 PDF，也不计算文件哈希。",
        "",
    ])
    return "\n".join(lines)


def audit_tree(source_root: Path, output: Path, rules: dict[str, Any]) -> dict[str, Any]:
    root = source_root.resolve()
    if not root.is_dir():
        raise NotADirectoryError(root)
    output = output.resolve()
    if output == root or root in output.parents:
        raise ValueError("输出目录不能位于源数据目录内部，以免下一次审查把报告误当成源文件。")

    files = sorted(path for path in root.rglob("*") if path.is_file())
    inventory: list[dict[str, Any]] = []
    by_collection: dict[str, list[tuple[Path, str]]] = defaultdict(list)
    for path in files:
        role = classify_file(path, rules)
        collection = collection_name(root, path)
        by_collection[collection].append((path, role))
        inventory.append({"relative_path": path.relative_to(root).as_posix(), "collection": collection, "extension": path.suffix.casefold(), "file_role": role, "size_bytes": path.stat().st_size})

    issues: list[Issue] = []
    for collection, items in sorted(by_collection.items()):
        curated = [path for path, role in items if role == "curated_historical_workbook"]
        pdfs = [path for path, role in items if role in {"article_pdf", "article_pdf_copy"}]
        locks = [path for path, role in items if role == "excel_temporary_lock"]
        if not curated:
            issues.append(issue("WARNING", "MANUAL_REVIEW", "COLLECTION_WITHOUT_CURATED_WORKBOOK", collection, "", message="目录中没有待入库的整理后工作簿。", suggested_action="确认该目录是否只是辅助材料。"))
        if len(curated) > 1:
            issues.append(issue("BLOCKER", "SOURCE_EXCEL", "MULTIPLE_CURATED_WORKBOOKS", collection, " | ".join(path.name for path in curated), message="同一来源目录含多个整理后工作簿，存在重复入库风险。", suggested_action="明确保留版本并隔离重复文件。"))
        if not pdfs:
            issues.append(issue("WARNING", "MANUAL_REVIEW", "SOURCE_PDF_NOT_PRESENT", collection, "", message="本地来源目录没有论文 PDF；不影响公开包，但会影响人工溯源核对。", suggested_action="如有访问权，仅在内部核对区保留；公开包只放引文/DOI。"))
        if len(pdfs) > 1:
            issues.append(issue("WARNING", "MANUAL_REVIEW", "MULTIPLE_SOURCE_PDFS", collection, " | ".join(path.name for path in pdfs), message="同一来源目录含多个 PDF 文件。", suggested_action="确认是否为副本；公开包不得包含论文 PDF。"))
        for path in locks:
            issues.append(issue("WARNING", "SOURCE_EXCEL", "EXCEL_TEMPORARY_LOCK", collection, path.relative_to(root).as_posix(), message="发现 Excel 临时锁文件。", suggested_action="关闭 Excel 后重新审查；锁文件不得入库或发布。"))

    workbook_records: list[WorkbookRecord] = []
    all_observations: list[Observation] = []
    curated_paths = [path for path in files if classify_file(path, rules) == "curated_historical_workbook"]
    for path in curated_paths:
        record, workbook_issues, observations = audit_workbook(root, path, rules)
        workbook_records.append(record)
        issues.extend(workbook_issues)
        all_observations.extend(observations)

    title_groups: dict[str, list[WorkbookRecord]] = defaultdict(list)
    doi_groups: dict[str, list[WorkbookRecord]] = defaultdict(list)
    for record in workbook_records:
        if normalize_title(record.title):
            title_groups[normalize_title(record.title)].append(record)
        if record.doi_normalized:
            doi_groups[record.doi_normalized].append(record)
    for group_type, groups in (("TITLE", title_groups), ("DOI", doi_groups)):
        for key, records in groups.items():
            if len(records) <= 1:
                continue
            files_text = " | ".join(record.relative_path for record in records)
            issues.append(issue("BLOCKER", "SOURCE_EXCEL", f"DUPLICATE_{group_type}", "multiple", files_text, message=f"多个整理后工作簿具有相同的规范化{group_type}。", evidence=key, suggested_action="逐篇核对后只保留正确版本，或在去重台账中明确排除。"))

    by_workbook = Counter(item.workbook for item in issues)
    for record in workbook_records:
        record.issue_count = by_workbook[record.relative_path]
    issues.sort(key=lambda item: (SEVERITY_ORDER[item.severity], item.owner, item.workbook, item.sheet, item.cell, item.code))
    severity_counts = Counter(item.severity for item in issues)
    owner_counts = Counter(item.owner for item in issues)
    blocking = [item for item in issues if item.severity in BLOCKING_SEVERITIES]
    duplicate_blockers = [item for item in blocking if item.code in {"DUPLICATE_TITLE", "DUPLICATE_DOI"}]
    role_counts = Counter(row["file_role"] for row in inventory)
    summary = {
        "auditor_version": "0.1.0",
        "generated_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source_root": root.as_posix(),
        "output_root": output.as_posix(),
        "read_only_source": True,
        "hashes_calculated": False,
        "article_pdfs_parsed": False,
        "total_files": len(inventory),
        "file_roles": dict(sorted(role_counts.items())),
        "collections": len(by_collection),
        "curated_workbooks": len(workbook_records),
        "recognized_workbooks": sum(record.workbook_status == "recognized" for record in workbook_records),
        "unique_normalized_titles": len(title_groups),
        "unique_normalized_dois": len(doi_groups),
        "observation_cells": len(all_observations),
        "issue_counts_by_severity": dict(sorted(severity_counts.items(), key=lambda item: SEVERITY_ORDER[item[0]])),
        "issue_counts_by_owner": dict(sorted(owner_counts.items())),
        "blocking_issue_count": len(blocking),
        "unresolved_duplicate_identity_count": len(duplicate_blockers),
        "can_build_database": not blocking and not duplicate_blockers,
    }

    output.mkdir(parents=True, exist_ok=True)
    metric_groups: dict[tuple[str, str, str, str], list[Observation]] = defaultdict(list)
    for observation in all_observations:
        metric_groups[(observation.context, observation.metric, observation.unit, observation.basis)].append(observation)
    metric_profiles: list[dict[str, Any]] = []
    for (context, metric, unit, basis), group in sorted(metric_groups.items(), key=lambda item: tuple(normalize_label(value) for value in item[0])):
        values = [number for number in (numeric(observation.value) for observation in group) if number is not None]
        metric_profiles.append({
            "context": context,
            "raw_metric": metric,
            "explicit_unit": unit,
            "explicit_basis": basis,
            "observation_count": len(group),
            "numeric_count": len(values),
            "numeric_min": min(values) if values else "",
            "numeric_max": max(values) if values else "",
            "source_workbook_count": len({observation.workbook for observation in group}),
            "example_workbook": group[0].workbook,
            "example_cell": f"{group[0].sheet}!{get_column_letter(group[0].column)}{group[0].row}",
            "unit_resolution_needed": "yes" if not unit and context in {"energy", "AID", "SID", "ATTD", "STTD"} else "no",
        })

    write_csv(output / "file_inventory.csv", inventory)
    write_csv(output / "workbook_inventory.csv", [asdict(record) for record in workbook_records], list(WorkbookRecord.__dataclass_fields__))
    write_csv(output / "metric_profile.csv", metric_profiles)
    issue_rows = [asdict(item) for item in issues]
    write_csv(output / "all_issues.csv", issue_rows, list(Issue.__dataclass_fields__))
    for owner, filename in (("SOURCE_EXCEL", "source_excel_issues.csv"), ("IMPORTER_RULE", "importer_rule_issues.csv"), ("DATABASE_QC", "database_qc_issues.csv"), ("MANUAL_REVIEW", "manual_review_issues.csv")):
        write_csv(output / filename, [asdict(item) for item in issues if item.owner == owner], list(Issue.__dataclass_fields__))
    write_json(output / "summary.json", summary)
    gate = {
        "can_build_database": summary["can_build_database"],
        "blocking_issue_count": summary["blocking_issue_count"],
        "unresolved_duplicate_identity_count": summary["unresolved_duplicate_identity_count"],
        "rule": "Build is allowed only when no BLOCKER/ERROR issue and no unresolved duplicate title/DOI remains.",
        "source_was_modified": False,
    }
    write_json(output / "release_gate.json", gate)
    (output / "AUDIT_REPORT.md").write_text(markdown_report(summary, issues), encoding="utf-8")
    return summary


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="独立、只读的历史饲料原料工作簿审查器")
    parser.add_argument("source", type=Path, help="原始数据根目录")
    parser.add_argument("output", type=Path, help="审查报告输出目录（不得位于源目录内部）")
    parser.add_argument("--rules", type=Path, default=Path(__file__).with_name("default_rules.json"), help="JSON 审查规则")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        summary = audit_tree(args.source, args.output, read_rules(args.rules))
    except Exception as exc:
        print(f"审查器运行失败：{type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    print(f"workbooks={summary['curated_workbooks']}")
    print(f"observations={summary['observation_cells']}")
    print(f"blocking_issues={summary['blocking_issue_count']}")
    print(f"can_build_database={str(summary['can_build_database']).lower()}")
    print(f"report={(args.output.resolve() / 'AUDIT_REPORT.md')}")
    return 0 if summary["can_build_database"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
