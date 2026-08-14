from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import auditor


class AuditorTests(unittest.TestCase):
    def make_workbook(self, root: Path, name: str, *, title: str, doi: str, bad_energy: bool = False) -> Path:
        collection = root / "原料" / "录入人" / name
        collection.mkdir(parents=True)
        path = collection / f"{name}.xlsx"
        book = Workbook()
        meta = book.active
        meta.title = "论文参数"
        meta.append(["title", "doi"])
        meta.append([title, doi])
        sheet = book.create_sheet("成分及营养价值数据")
        sheet.append(["", "样品1", "样品2", "样品3"])
        sheet.append(["基础数据", None, None, None])
        sheet.append(["dry base", "DM", "DM", "DM"])
        sheet.append(["能量", None, None, None])
        sheet.append(["AME(MJ/kg)", 13.8, 14.1, 14019 if bad_energy else 14.0])
        book.save(path)
        return path

    def test_duplicate_and_energy_are_blocking_and_source_is_unchanged(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            source = base / "source"
            output = base / "output"
            first = self.make_workbook(source, "a", title="Same title", doi="https://doi.org/10.1/X", bad_energy=True)
            self.make_workbook(source, "b", title="Same title", doi="doi:10.1/x")
            before = first.stat().st_mtime_ns
            rules = auditor.read_rules(Path(auditor.__file__).with_name("default_rules.json"))
            summary = auditor.audit_tree(source, output, rules)
            self.assertFalse(summary["can_build_database"])
            self.assertEqual(before, first.stat().st_mtime_ns)
            issue_codes = {row["code"] for row in self.read_csv(output / "all_issues.csv")}
            self.assertIn("DUPLICATE_TITLE", issue_codes)
            self.assertIn("DUPLICATE_DOI", issue_codes)
            self.assertIn("ENERGY_SCALE_OUTLIER", issue_codes)
            self.assertIn("ENERGY_MJ_IMPLAUSIBLE", issue_codes)
            gate = json.loads((output / "release_gate.json").read_text(encoding="utf-8"))
            self.assertFalse(gate["source_was_modified"])

    def test_title_required_doi_optional(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            source = base / "source"
            output = base / "output"
            self.make_workbook(source, "no_doi", title="A valid title", doi="")
            rules = auditor.read_rules(Path(auditor.__file__).with_name("default_rules.json"))
            summary = auditor.audit_tree(source, output, rules)
            codes = {row["code"] for row in self.read_csv(output / "all_issues.csv")}
            self.assertNotIn("MISSING_DOI", codes)
            self.assertEqual(summary["unique_normalized_dois"], 0)

    def test_reused_ingredient_number_and_ambiguous_sid_value_do_not_block(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            source = base / "source"
            output = base / "output"
            path = self.make_workbook(source, "reused", title="Valid study", doi="")
            book = load_workbook(path)
            sheet = book["成分及营养价值数据"]
            sheet.insert_rows(4, amount=2)
            sheet["A4"] = "原料编号"
            sheet["B4"] = 1
            sheet["C4"] = 1
            sheet["D4"] = 2
            sheet["A5"] = "试验编号"
            sheet["B5"] = 1
            sheet["C5"] = 2
            sheet["D5"] = 1
            sheet["A8"] = "非表观营养价值"
            sheet["A9"] = "SID"
            sheet["B9"] = "data"
            sheet["A10"] = "Leucine"
            sheet["B10"] = 235.2
            sheet["C10"] = 242.4
            sheet["D10"] = 240.1
            book.save(path)
            book.close()
            rules = auditor.read_rules(Path(auditor.__file__).with_name("default_rules.json"))
            summary = auditor.audit_tree(source, output, rules)
            rows = self.read_csv(output / "all_issues.csv")
            codes = {row["code"] for row in rows}
            self.assertIn("REUSED_INGREDIENT_NUMBER", codes)
            self.assertIn("DIGESTIBILITY_CONTEXT_VALUE_OVER_100", codes)
            self.assertNotIn("DUPLICATE_INGREDIENT_NUMBER", codes)
            self.assertEqual(summary["blocking_issue_count"], 0)

    def test_output_cannot_be_inside_source(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source"
            self.make_workbook(source, "a", title="Valid", doi="")
            rules = auditor.read_rules(Path(auditor.__file__).with_name("default_rules.json"))
            with self.assertRaises(ValueError):
                auditor.audit_tree(source, source / "audit_output", rules)

    @staticmethod
    def read_csv(path: Path) -> list[dict[str, str]]:
        import csv
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))


if __name__ == "__main__":
    unittest.main()
