"""Finalize independent ScienceDB-data and GitHub-code upload candidates."""

from __future__ import annotations

import argparse
import csv
import json
import shutil
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

VERSION_DATE = "2026-08-14"
PUBLIC_SHEETS = ["Article Metadata", "Ingredient Data"]
TEMPLATE_SHEETS = ["Instructions_说明", *PUBLIC_SHEETS]
GITHUB_DIRECTORIES = (
    "config", "docs", "examples", "feeddb", "integrations", "source_auditor",
    "templates", "tests", "tools",
)
GITHUB_FILES = (
    ".gitignore", "CODE_LICENSE_REVIEW_REQUIRED.txt", "pyproject.toml", "README.md", "run_gui.bat",
)
IGNORED_DIRECTORY_NAMES = {"__pycache__", ".pytest_cache", "release_working", "tmp", ".git"}
IGNORED_SUFFIXES = {".pyc", ".pyo"}


def _write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _copy_tree(source: Path, destination: Path) -> None:
    def ignore(directory: str, names: list[str]) -> set[str]:
        ignored = set()
        for name in names:
            path = Path(directory) / name
            if name in IGNORED_DIRECTORY_NAMES or path.suffix.casefold() in IGNORED_SUFFIXES:
                ignored.add(name)
        return ignored

    shutil.copytree(source, destination, ignore=ignore)


def _manifest(root: Path, path: Path) -> None:
    rows = []
    for item in sorted(candidate for candidate in root.rglob("*") if candidate.is_file()):
        if item.resolve() == path.resolve():
            continue
        relative = item.relative_to(root).as_posix()
        top = relative.split("/", 1)[0]
        purpose = {
            "data": "editable data products and exchange tables",
            "database": "read-only SQLite data product",
            "dictionaries": "name, metric, unit, and basis mappings",
            "templates": "bilingual contribution formats and examples",
            "documentation": "data understanding, citation, and reuse documentation",
            "validation": "quality-control and release-gate evidence",
            "feeddb": "application and library source code",
            "tests": "automated verification",
            "examples": "read-only queries and reproducible reuse examples",
        }.get(top, "package metadata or software support file")
        rows.append({
            "relative_path": relative,
            "size_bytes": item.stat().st_size,
            "purpose": purpose,
            "version_date": VERSION_DATE,
        })
    _write_csv(path, rows, ["relative_path", "size_bytes", "purpose", "version_date"])


def _database_gate(database: Path) -> dict[str, object]:
    with sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True) as connection:
        connection.execute("PRAGMA query_only=ON")
        actual = {
            "studies": connection.execute("SELECT count(*) FROM studies").fetchone()[0],
            "ingredient_samples": connection.execute("SELECT count(*) FROM ingredient_samples").fetchone()[0],
            "logical_observations": connection.execute("SELECT count(*) FROM observations").fetchone()[0],
            "provenance_records": connection.execute("SELECT count(*) FROM source_provenance").fetchone()[0],
            "held": connection.execute("SELECT count(*) FROM observations WHERE validation_status='held'").fetchone()[0],
            "rejected": connection.execute("SELECT count(*) FROM observations WHERE validation_status='rejected'").fetchone()[0],
            "unapproved_metric": connection.execute(
                "SELECT count(*) FROM observations WHERE metric_mapping_status<>'approved'"
            ).fetchone()[0],
            "unapproved_unit": connection.execute(
                "SELECT count(*) FROM observations WHERE unit_mapping_status<>'approved'"
            ).fetchone()[0],
            "unapproved_basis": connection.execute(
                "SELECT count(*) FROM observations WHERE basis_mapping_status<>'approved'"
            ).fetchone()[0],
            "missing_canonical_unit": connection.execute(
                "SELECT count(*) FROM observations WHERE trim(coalesce(canonical_unit,''))=''"
            ).fetchone()[0],
            "missing_titles": connection.execute(
                "SELECT count(*) FROM studies WHERE trim(coalesce(title,''))=''"
            ).fetchone()[0],
            "blank_doi": connection.execute(
                "SELECT count(*) FROM studies WHERE trim(coalesce(doi,''))=''"
            ).fetchone()[0],
            "invalid_public_source_sheets": connection.execute(
                "SELECT count(*) FROM source_provenance WHERE trim(coalesce(source_sheet,''))<>'' AND source_sheet NOT IN (?,?)",
                PUBLIC_SHEETS,
            ).fetchone()[0],
        }
        actual["sqlite_integrity"] = connection.execute("PRAGMA integrity_check").fetchone()[0]
        actual["foreign_key_errors"] = len(connection.execute("PRAGMA foreign_key_check").fetchall())
    expected = {
        "studies": 110,
        "ingredient_samples": 690,
        "logical_observations": 20247,
        "provenance_records": 27793,
        "held": 0,
        "rejected": 0,
        "unapproved_metric": 0,
        "unapproved_unit": 0,
        "unapproved_basis": 0,
        "missing_canonical_unit": 0,
        "missing_titles": 0,
        "invalid_public_source_sheets": 0,
        "sqlite_integrity": "ok",
        "foreign_key_errors": 0,
    }
    comparisons = {key: actual[key] == value for key, value in expected.items()}
    return {"expected": expected, "actual": actual, "comparisons": comparisons, "passed": all(comparisons.values())}


def _workbook_gate(science: Path) -> dict[str, object]:
    editable = science / "data" / "editable_workbooks"
    files = sorted(editable.glob("study_*.xlsx"))
    failures = []
    for path in files:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if workbook.sheetnames != PUBLIC_SHEETS:
            failures.append({"file": path.name, "sheets": workbook.sheetnames})
        workbook.close()
    template = science / "templates" / "feed_ingredient_contribution_template.xlsx"
    workbook = load_workbook(template, read_only=True, data_only=True)
    template_sheets = workbook.sheetnames
    workbook.close()
    return {
        "workbook_count": len(files),
        "expected_workbook_count": 110,
        "public_workbook_sheet_failures": failures,
        "template_sheets": template_sheets,
        "passed": len(files) == 110 and not failures and template_sheets == TEMPLATE_SHEETS,
    }


def _science_boundary_gate(science: Path) -> dict[str, object]:
    forbidden_suffixes = {
        ".py", ".pyc", ".pyo", ".ipynb", ".joblib", ".pkl", ".pdf",
        ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".sql",
    }
    forbidden = []
    for path in science.rglob("*"):
        if path.is_dir() and path.name == "__pycache__":
            forbidden.append(path.relative_to(science).as_posix())
        elif path.is_file() and (
            path.suffix.casefold() in forbidden_suffixes
            or path.name.casefold().startswith("test_")
            or "ocr" in path.name.casefold()
        ):
            forbidden.append(path.relative_to(science).as_posix())
    database_files = [
        path.relative_to(science).as_posix()
        for path in (science / "database").rglob("*")
        if path.is_file() and path.suffix.casefold() != ".sqlite"
    ]
    return {
        "forbidden_files": sorted(forbidden),
        "non_sqlite_files_in_database_directory": sorted(database_files),
        "passed": not forbidden and not database_files,
    }


def finalize_sciencedb(science: Path) -> dict[str, object]:
    index = science / "data" / "study_source_index.csv"
    with index.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    references = [
        {
            "public_workbook_id": row["public_workbook_id"],
            "doi": row["doi"],
            "title": row["title"],
            "journal": row["journal"],
            "publication_year": row["publication_year"],
            "country": row["country"],
            "source_reference": row["source_reference"],
        }
        for row in rows
    ]
    _write_csv(
        science / "documentation" / "source_literature_references.csv",
        references,
        ["public_workbook_id", "doi", "title", "journal", "publication_year", "country", "source_reference"],
    )
    (science / "documentation" / "repository_metadata.md").write_text(
        """# ScienceDB candidate metadata

- Version date: 2026-08-14
- Related article or project: Not applicable
- Data license: Pending author decision
- Repository record URL: Pending author decision
- Associated article setting: Do not set an associated article
""",
        encoding="utf-8",
    )
    database_gate = _database_gate(science / "database" / "feed_ingredient_release.sqlite")
    workbook_gate = _workbook_gate(science)
    boundary_gate = _science_boundary_gate(science)
    report = {
        "version_date": VERSION_DATE,
        "database_gate": database_gate,
        "workbook_gate": workbook_gate,
        "boundary_gate": boundary_gate,
        "passed": database_gate["passed"] and workbook_gate["passed"] and boundary_gate["passed"],
    }
    _write_json(science / "validation" / "upload_candidate_validation.json", report)
    actual = database_gate["actual"]
    (science / "validation" / "FINAL_RELEASE_REPORT.md").write_text(
        f"""# ScienceDB upload-candidate validation report

- Version date: {VERSION_DATE}
- Public workbooks: {workbook_gate['workbook_count']}; exact sheets: Article Metadata, Ingredient Data
- Studies: {actual['studies']}
- Ingredient samples: {actual['ingredient_samples']}
- Logical observations: {actual['logical_observations']}
- Provenance records: {actual['provenance_records']}
- Held/rejected: {actual['held']}/{actual['rejected']}
- Unapproved metric/unit/basis: {actual['unapproved_metric']}/{actual['unapproved_unit']}/{actual['unapproved_basis']}
- Missing canonical unit: {actual['missing_canonical_unit']}
- SQLite integrity/foreign-key errors: {actual['sqlite_integrity']}/{actual['foreign_key_errors']}
- ScienceDB boundary scan: {'PASS' if boundary_gate['passed'] else 'FAIL'}
- Overall data and boundary gates: {'PASS' if report['passed'] else 'FAIL'}

Licensing and repository record URL remain author decisions. Related article or project is Not applicable.
""",
        encoding="utf-8",
    )
    _manifest(science, science / "documentation" / "release_contents.csv")
    return report


def validate_github(github: Path) -> dict[str, object]:
    forbidden = []
    for path in github.rglob("*"):
        if path.is_dir() and path.name in {"__pycache__", ".pytest_cache"}:
            forbidden.append(path.relative_to(github).as_posix())
        elif path.is_file() and (
            path.suffix.casefold() in {".sqlite", ".db", ".pdf", ".pyc", ".pyo"}
            or (path.suffix.casefold() == ".xlsx" and path.name != "feed_ingredient_contribution_template.xlsx")
            or (path.name.startswith("study_") and path.suffix.casefold() == ".xlsx")
        ):
            forbidden.append(path.relative_to(github).as_posix())
    report = {
        "version_date": VERSION_DATE,
        "forbidden_formal_data_or_source_files": sorted(forbidden),
        "query_rule_files": sorted(
            path.relative_to(github).as_posix() for path in github.rglob("*.sql")
        ),
        "passed": not forbidden,
    }
    _write_json(github / "RELEASE_VALIDATION.json", report)
    _manifest(github, github / "FILE_MANIFEST.csv")
    return report


def assemble_github(source: Path, github: Path) -> dict[str, object]:
    if github.exists():
        raise FileExistsError(f"GitHub output directory already exists: {github}")
    github.mkdir(parents=True)
    for name in GITHUB_DIRECTORIES:
        path = source / name
        if path.is_dir():
            _copy_tree(path, github / name)
    for name in GITHUB_FILES:
        path = source / name
        if path.is_file():
            shutil.copy2(path, github / name)
    (github / "REPOSITORY_METADATA.md").write_text(
        """# GitHub repository metadata

- Version date: 2026-08-14
- Repository: zengzhengcheng/feed-ingredient-database-toolkit
- Repository URL: https://github.com/zengzhengcheng/feed-ingredient-database-toolkit
- Visibility: private
- Default branch: main
- Related article or project: Not applicable
- Associated article setting: Do not set an associated article
- Code license: Pending author decision; see CODE_LICENSE_REVIEW_REQUIRED.txt
""",
        encoding="utf-8",
    )
    (github / "RELEASE_BOUNDARY.md").write_text(
        """# Release boundary

This repository contains software, tests, editable read-only SQL query rules, blank contribution templates, seed dictionaries, and reproducible reuse examples. The formal 110 public workbooks, exchange tables, data dictionaries, and release SQLite database are data products distributed separately through ScienceDB (`ScienceDB URL: pending author release`).

Excluded here: the 110 formal workbooks, the formal release SQLite database, ScienceDB exchange tables, source article PDFs/full text, user working databases, diet/formulation data, and the complete ScienceDB data package.

Related article or project: Not applicable.
""",
        encoding="utf-8",
    )
    return validate_github(github)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("science_candidate", type=Path)
    parser.add_argument("github_source", type=Path)
    parser.add_argument("github_candidate", type=Path)
    args = parser.parse_args()
    science = args.science_candidate.resolve()
    github_source = args.github_source.resolve()
    github = args.github_candidate.resolve()
    science_report = finalize_sciencedb(science)
    github_report = assemble_github(github_source, github)
    print(json.dumps({"science": science_report, "github": github_report}, ensure_ascii=False, indent=2))
    return int(not science_report["passed"] or not github_report["passed"])


if __name__ == "__main__":
    raise SystemExit(main())
