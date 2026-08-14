from __future__ import annotations

import json
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from feeddb.bilingual_template import (
    ARTICLE_METADATA_FIELDS,
    ARTICLE_METADATA_SHEET,
    INGREDIENT_DATA_FIELDS,
    INGREDIENT_DATA_SHEET,
    INSTRUCTIONS_SHEET,
)
from feeddb.contribution_file import load_contribution_file
from feeddb.excel_io import import_workbook


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "feed_ingredient_contribution_template.xlsx"


def _filled_template(path: Path) -> None:
    workbook = load_workbook(TEMPLATE)
    metadata = workbook[ARTICLE_METADATA_SHEET]
    metadata.append(["S001", "", "Bilingual template import", "Example Journal", 2025, "Pig", "China"])
    data = workbook[INGREDIENT_DATA_SHEET]
    positions = {name: index for index, name in enumerate(INGREDIENT_DATA_FIELDS)}

    def row(record_type: str, **values):
        result = [""] * len(INGREDIENT_DATA_FIELDS)
        result[positions["record_type"]] = record_type
        result[positions["study_key"]] = "S001"
        for key, value in values.items():
            result[positions[key]] = value
        data.append(result)

    row("trial", trial_key="T001", source_trial_no="1", treatment="Control")
    row(
        "ingredient", ingredient_key="I001", source_ingredient_no="1", raw_name="Corn",
        reported_basis="dry matter", dry_matter_value=88.5,
    )
    row(
        "animal_group", trial_key="T001", animal_group_key="A001", species="Pig",
        animal_type="growing pig", body_weight_kg=25,
    )
    row(
        "composition", ingredient_key="I001", measurement_context="composition",
        raw_metric="CP", value=9.5, unit="%", basis="dry matter",
    )
    row(
        "nutrition", ingredient_key="I001", animal_group_key="A001",
        measurement_context="energy", raw_metric="ME", value=14.2,
        unit="MJ/kg", basis="dry matter",
    )
    workbook.save(path)
    workbook.close()


def test_blank_template_has_stable_bilingual_three_sheet_contract():
    workbook = load_workbook(TEMPLATE, read_only=True, data_only=True)
    assert workbook.sheetnames == [INSTRUCTIONS_SHEET, ARTICLE_METADATA_SHEET, INGREDIENT_DATA_SHEET]
    assert tuple(cell.value for cell in workbook[ARTICLE_METADATA_SHEET][1]) == ARTICLE_METADATA_FIELDS
    assert tuple(cell.value for cell in workbook[INGREDIENT_DATA_SHEET][1]) == INGREDIENT_DATA_FIELDS
    workbook.close()


def test_bilingual_excel_and_json_have_equivalent_batch_semantics(tmp_path: Path):
    excel_path = tmp_path / "contribution.xlsx"
    _filled_template(excel_path)
    excel = load_contribution_file(excel_path)
    assert excel.source_kind == "bilingual_two_sheet_excel"
    assert excel.batch.title == "Bilingual template import"
    assert excel.batch.doi == ""
    assert [item.raw_metric for item in excel.batch.composition_metrics] == ["CP"]
    assert [item.raw_metric for item in excel.batch.nutrition_metrics] == ["ME"]
    assert excel.batch.ingredients[0].values == ["9.5"]
    assert excel.batch.nutrition_rows[0].values == ["14.2"]

    payload = {
        "format_version": "1.0",
        "metadata": {
            "title": "Bilingual template import", "doi": None, "journal": "Example Journal",
            "publication_year": 2025, "species": "Pig", "country": "China",
        },
        "composition_metrics": [{"metric": "CP", "unit": "%", "basis": "dry matter"}],
        "ingredients": [{
            "ingredient_key": "I001", "raw_name": "Corn", "source_ingredient_no": "1",
            "reported_basis": "dry matter", "dry_matter_value": 88.5, "values": [9.5],
        }],
        "animals": [{
            "animal_group_key": "A001", "trial_no": "1", "species": "Pig",
            "animal_type": "growing pig", "body_weight_kg": 25,
        }],
        "nutrition_metrics": [{"metric": "ME", "context": "energy", "unit": "MJ/kg", "basis": "dry matter"}],
        "nutrition_rows": [{"ingredient_key": "I001", "animal_group_key": "A001", "values": [14.2]}],
    }
    json_path = tmp_path / "contribution.json"
    json_path.write_text(json.dumps(payload), encoding="utf-8")
    json_batch = load_contribution_file(json_path).batch
    assert json_batch.title == excel.batch.title
    assert json_batch.doi == excel.batch.doi
    assert json_batch.composition_metrics == excel.batch.composition_metrics
    assert json_batch.nutrition_metrics == excel.batch.nutrition_metrics
    assert json_batch.ingredients[0].raw_name == excel.batch.ingredients[0].raw_name
    assert json_batch.ingredients[0].values == excel.batch.ingredients[0].values
    assert json_batch.nutrition_rows[0].values == excel.batch.nutrition_rows[0].values


def test_bilingual_excel_cli_import_accepts_blank_doi_and_preserves_values(tmp_path: Path):
    source = tmp_path / "contribution.xlsx"
    database = tmp_path / "working.sqlite"
    _filled_template(source)
    result = import_workbook(source, database)
    assert not [issue for issue in result.issues if issue.level == "error"]
    assert (result.studies, result.ingredients, result.component_observations, result.nutrition_observations) == (1, 1, 1, 1)
    with sqlite3.connect(database) as connection:
        title, doi = connection.execute("SELECT title,doi FROM studies").fetchone()
        component = connection.execute("SELECT value_num,canonical_unit FROM component_observations").fetchone()
        nutrition = connection.execute("SELECT value_num,canonical_unit FROM nutrition_observations").fetchone()
    assert title == "Bilingual template import" and doi is None
    assert component == (9.5, "%")
    assert nutrition == (14.2, "MJ/kg")


def test_title_is_required_but_measurement_value_may_be_blank(tmp_path: Path):
    blank = tmp_path / "blank.xlsx"
    workbook = load_workbook(TEMPLATE)
    workbook[ARTICLE_METADATA_SHEET].append(["S001", "", "", "", "", "", ""])
    workbook.save(blank)
    workbook.close()
    result = import_workbook(blank, tmp_path / "blank.sqlite")
    assert any("title" in issue.message.casefold() or "题" in issue.message for issue in result.issues)
