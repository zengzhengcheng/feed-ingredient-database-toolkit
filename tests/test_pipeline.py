import os
import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from feeddb.database import apply_dictionary_mapping, canonicalize, connect, initialize, lookup_dictionary, seed_dictionaries, upsert_metric_catalog
from feeddb.contribution_api import batch_from_payload, import_payload, read_payload, validate_payload
from feeddb.contribution_file import load_contribution_file
from feeddb.excel_io import discover_legacy_workbooks, import_workbook, validate_standard_workbook
from feeddb.legacy_parser import parse_legacy_ingredient_sheet
from feeddb.release_audit import audit_source_tree
from feeddb.release_stats import collect_release_statistics
from feeddb.release_builder import build_release_package
from feeddb.release_check import check_release_package
from feeddb.contribution_batch import (
    AnimalGridRow, ContributionBatch, IngredientGridRow, MetricSpec, NutritionGridRow,
    save_contribution_batch, validate_contribution_batch,
)
from feeddb.query import search_database
from feeddb.user_query import execute_readonly_query, export_readonly_query, validate_readonly_sql
from feeddb.database_admin import (
    clear_database, delete_database_file, delete_study, preview_delete_study, rebuild_database,
)
from feeddb.term_review import (
    create_standard_from_candidate, list_candidate_terms, resolve_candidate_as_alias,
)


def test_standard_template_import(tmp_path: Path):
    book=Workbook(); ws=book.active; ws.title="Studies"; ws.append(["study_key","doi","title","journal","publication_year","species","country"]); ws.append(["S1","10.1/test","Test","Journal",2024,"pig","China"])
    ws=book.create_sheet("Trials"); ws.append(["trial_key","study_key","source_trial_no"]); ws.append(["T1","S1","1"])
    ws=book.create_sheet("Ingredients"); ws.append(["ingredient_key","study_key","raw_name","feed_type"]); ws.append(["I1","S1","corn","yellow dent"])
    ws=book.create_sheet("AnimalGroups"); ws.append(["animal_group_key","trial_key"])
    ws=book.create_sheet("CompositionValues"); ws.append(["ingredient_key","raw_metric","value","unit","basis","qualifier"]); ws.append(["I1","CP",9.5,"%","DM",""])
    ws=book.create_sheet("NutritionValues"); ws.append(["ingredient_key","animal_group_key","measurement_context","raw_metric","value"])
    path=tmp_path/"input.xlsx"; db=tmp_path/"feed.sqlite"; book.save(path)
    result=import_workbook(path,db)
    assert (result.studies,result.ingredients,result.observations,result.held)==(1,1,1,0)
    with connect(db) as conn:
        row=conn.execute("""SELECT i.canonical_name,o.canonical_metric,o.value_num,
                                   o.raw_unit,o.canonical_unit,o.unit_resolution_method,o.unit_rule_reference
                          FROM ingredient_samples i JOIN component_observations o USING(ingredient_sample_id)""").fetchone()
    assert tuple(row)==("Corn","CP",9.5,"%","%","explicit_field","source unit field")


def test_standard_loader_rejects_duplicate_trial_keys_and_accepts_casefolded_references(tmp_path: Path):
    def workbook_with_trial_keys(keys: list[str], reference: str, filename: str) -> Path:
        book = Workbook(); studies = book.active; studies.title = "Studies"
        studies.append(["study_key", "doi", "title"]); studies.append(["S1", "", "Key checks"])
        trials = book.create_sheet("Trials"); trials.append(["trial_key", "study_key", "source_trial_no"])
        for index, key in enumerate(keys, start=1): trials.append([key, "S1", str(index)])
        ingredients = book.create_sheet("Ingredients"); ingredients.append(["ingredient_key", "study_key", "raw_name"])
        ingredients.append(["I1", "S1", "corn"])
        animals = book.create_sheet("AnimalGroups"); animals.append(["animal_group_key", "trial_key"])
        animals.append(["A1", reference])
        composition = book.create_sheet("CompositionValues"); composition.append(["ingredient_key", "raw_metric", "value"])
        composition.append(["i1", "CP", 9.5])
        nutrition = book.create_sheet("NutritionValues")
        nutrition.append(["ingredient_key", "animal_group_key", "measurement_context", "raw_metric", "value"])
        path = tmp_path / filename; book.save(path); return path

    with pytest.raises(ValueError, match="trial_key 重复"):
        load_contribution_file(workbook_with_trial_keys(["T1", "t1"], "T1", "duplicates.xlsx"))
    loaded = load_contribution_file(workbook_with_trial_keys(["t1"], "T1", "casefold.xlsx"))
    assert loaded.batch.ingredients[0].values == ["9.5"]
    assert loaded.batch.animals[0].trial_no == "1"


def test_standard_preflight_warns_about_merged_input_cells(tmp_path: Path):
    book = Workbook(); studies = book.active; studies.title = "Studies"
    studies.append(["study_key", "doi", "title"]); studies.append(["S1", "", "Merged check"])
    studies.merge_cells("B2:C2")
    for sheet_name, headers in (
        ("Trials", ["trial_key", "study_key", "source_trial_no"]),
        ("Ingredients", ["ingredient_key", "study_key", "raw_name"]),
        ("AnimalGroups", ["animal_group_key", "trial_key"]),
        ("CompositionValues", ["ingredient_key", "raw_metric", "value"]),
        ("NutritionValues", ["ingredient_key", "animal_group_key", "measurement_context", "raw_metric", "value"]),
    ):
        book.create_sheet(sheet_name).append(headers)
    path = tmp_path / "merged.xlsx"; book.save(path)
    issues = validate_standard_workbook(path)
    assert any(issue.level == "warning" and "Merged cells" in issue.message for issue in issues)


def test_dictionary_approval_updates_mapping_without_holding_valid_data(tmp_path: Path):
    book=Workbook(); ws=book.active; ws.title="Studies"; ws.append(["study_key","doi","title"]); ws.append(["S1","10.1/unknown","Test"])
    ws=book.create_sheet("Trials"); ws.append(["trial_key","study_key","source_trial_no"]); ws.append(["T1","S1","1"])
    ws=book.create_sheet("Ingredients"); ws.append(["ingredient_key","study_key","raw_name"]); ws.append(["I1","S1","Novel ingredient"])
    ws=book.create_sheet("AnimalGroups"); ws.append(["animal_group_key","trial_key"])
    ws=book.create_sheet("CompositionValues"); ws.append(["ingredient_key","raw_metric","value","unit","basis"]); ws.append(["I1","Novel metric",1.2,"%","DM"])
    ws=book.create_sheet("NutritionValues"); ws.append(["ingredient_key","animal_group_key","measurement_context","raw_metric","value"])
    path=tmp_path/"candidate.xlsx"; db=tmp_path/"feed.sqlite"; book.save(path)
    result=import_workbook(path,db); assert result.held==0; assert result.unmapped==1
    with connect(db) as conn:
        apply_dictionary_mapping(conn,"ingredient","Novel ingredient","Novel ingredient","approved")
        apply_dictionary_mapping(conn,"metric","Novel metric","Novel metric","approved")
        status=conn.execute("SELECT validation_status,metric_mapping_status FROM component_observations").fetchone()
    assert tuple(status)==("accepted","approved")


def test_source_audit_accounts_for_every_file(tmp_path: Path):
    collection = tmp_path / "玉米" / "Contributor" / "data1"
    collection.mkdir(parents=True)
    book = Workbook(); book.active.title = "成分及营养价值数据"
    book.save(collection / "data1.xlsx")
    book.save(collection / "data1原始.xlsx")
    (collection / "article.pdf").write_bytes(b"%PDF-placeholder")
    (collection / "~$data1.xlsx").write_bytes(b"lock")
    records, collections, summary = audit_source_tree(tmp_path)
    assert len(records) == 4
    assert len(collections) == 1
    assert summary["total_files"] == 4
    assert summary["file_roles"]["curated_historical_workbook"] == 1
    assert summary["file_roles"]["raw_historical_workbook"] == 1
    assert summary["file_roles"]["article_pdf"] == 1
    assert summary["file_roles"]["excel_temporary_lock"] == 1


def _legacy_book(path: Path, doi: str, include_diet: bool = False):
    book = Workbook(); ws = book.active; ws.title = "成分及营养价值数据"
    ws.append(["ingredient", "Corn"]); ws.append(["原料编号", 1]); ws.append(["饲料组成", "data"]); ws.append(["CP (%)", 9.5])
    meta = book.create_sheet("论文参数"); meta.append(["doi", "title"]); meta.append([doi, "Test"])
    if include_diet:
        diet = book.create_sheet("配方数据库"); diet.append(["ingredient", "diet value"]); diet.append(["Corn", 99])
    book.save(path)


def test_duplicate_doi_resolution_is_explicit(tmp_path: Path):
    first = tmp_path / "玉米" / "A" / "data1" / "data1.xlsx"; first.parent.mkdir(parents=True)
    second = tmp_path / "玉米" / "A" / "data2" / "data2.xlsx"; second.parent.mkdir(parents=True)
    _legacy_book(first, "10.1/duplicate"); _legacy_book(second, "10.1/duplicate")
    resolution = tmp_path / "duplicates.csv"
    resolution.write_text(
        "doi,preferred_relative_path,excluded_relative_path,resolution_reason\n"
        "10.1/duplicate,玉米/A/data1/data1.xlsx,玉米/A/data2/data2.xlsx,test choice\n",
        encoding="utf-8",
    )
    assert discover_legacy_workbooks(tmp_path, resolution) == [first]


def test_blank_doi_uses_required_title_as_duplicate_identity(tmp_path: Path):
    first = tmp_path / "玉米" / "A" / "data1" / "data1.xlsx"; first.parent.mkdir(parents=True)
    second = tmp_path / "玉米" / "A" / "data2" / "data2.xlsx"; second.parent.mkdir(parents=True)
    _legacy_book(first, ""); _legacy_book(second, "")
    second_book = load_workbook(second); second_book["论文参数"]["B2"] = "Another title"; second_book.save(second); second_book.close()
    assert discover_legacy_workbooks(tmp_path) == [first, second]


def test_release_builder_creates_sanitized_two_sheet_package(tmp_path: Path):
    source_root = tmp_path / "source"
    source = source_root / "玉米" / "Contributor" / "data1" / "data1.xlsx"
    source.parent.mkdir(parents=True)
    book = Workbook(); data = book.active; data.title = "成分及营养价值数据"
    data.append(["ingredient", "Corn"]); data.append(["原料编号", 1]); data.append(["饲料组成", "data"]); data.append(["CP (%)", 9.5])
    metadata = book.create_sheet("论文参数")
    metadata.append(["doi", "title", "journal", "luru_preson", "cunyi"])
    metadata.append(["10.1/release-builder", "Release builder test", "Journal", "Internal user", "Internal note"])
    diet = book.create_sheet("配方数据库"); diet.append(["ingredient", "value"]); diet.append(["Corn", 99])
    book.save(source); book.close()
    (source.parent / "article.pdf").write_bytes(b"%PDF-placeholder")

    working = tmp_path / "working.sqlite"
    imported = import_workbook(source, working, mode="legacy")
    assert imported.observations == 1
    rights = tmp_path / "rights.csv"
    rights.write_text(
        "public_workbook_id,doi,rights_status,rights_note\n"
        "study_0001,10.1/release-builder,approved,Author confirmed structured data release\n",
        encoding="utf-8",
    )
    output = tmp_path / "release"
    result = build_release_package(
        source_root, working, output, rights_review_file=rights,
        release_version="v0.5-test", data_license="CC BY 4.0",
    )
    assert result.ready_for_repository is True
    clean = load_workbook(output / "data" / "editable_workbooks" / "study_0001.xlsx", read_only=True, data_only=True)
    assert clean.sheetnames == ["Article Metadata", "Ingredient Data"]
    assert [cell.value for cell in clean["Article Metadata"][1]] == ["doi", "title", "journal"]
    clean.close()
    clean_report = (output / "validation" / "clean_workbook_report.csv").read_text(encoding="utf-8-sig")
    assert "removed_internal_metadata_fields" in clean_report
    assert "luru_preson" not in clean_report and "cunyi" not in clean_report
    assert not list(output.rglob("*.pdf"))
    assert (output / "validation" / "duplicate_decisions.csv").exists()
    with connect(output / "database" / "feed_ingredient_release.sqlite") as conn:
        assert conn.execute("SELECT DISTINCT source_workbook FROM studies").fetchone()[0] == "study_0001.xlsx"
        assert conn.execute("SELECT count(*) FROM observations").fetchone()[0] == 1
    readiness = check_release_package(output, write_report=False)
    assert readiness["ready_for_repository"] is True
    assert len(list((output / "data" / "exchange_tables").glob("*.csv"))) >= 14


def test_release_builder_reports_pending_rights_and_license(tmp_path: Path):
    source_root = tmp_path / "source"
    source = source_root / "玉米" / "Contributor" / "data1" / "data1.xlsx"
    source.parent.mkdir(parents=True); _legacy_book(source, "10.1/pending")
    working = tmp_path / "working.sqlite"; import_workbook(source, working, mode="legacy")
    output = tmp_path / "release"
    result = build_release_package(source_root, working, output)
    assert result.ready_for_repository is False
    report = check_release_package(output, write_report=False)
    failed_codes = {item["code"] for item in report["checks"] if item["status"] == "fail"}
    assert {"rights_review", "data_license"} <= failed_codes


def test_diet_sheet_is_never_imported(tmp_path: Path):
    workbook = tmp_path / "legacy.xlsx"; database = tmp_path / "feed.sqlite"
    _legacy_book(workbook, "10.1/no-diet", include_diet=True)
    result = import_workbook(workbook, database, mode="legacy")
    assert result.observations == 1
    with connect(database) as conn:
        row = conn.execute(
            "SELECT raw_metric,canonical_metric,raw_unit,canonical_unit,unit_resolution_method "
            "FROM observations"
        ).fetchone()
    assert tuple(row) == ("CP (%)", "CP", "%", "%", "embedded_in_metric")


def test_legacy_dictionary_is_context_aware(tmp_path: Path):
    database = tmp_path / "feed.sqlite"; initialize(database)
    with connect(database) as conn:
        seed_dictionaries(conn)
        assert canonicalize(conn, "ingredient", "sbm") == ("soybean meal", "approved")
        assert canonicalize(conn, "ingredient", "yellow corn") == ("Corn", "approved")
        assert canonicalize(conn, "metric", "Crude Protein", "composition") == ("CP", "approved")
        assert canonicalize(conn, "metric", "metabolism energy", "energy") == ("ME", "approved")
        assert canonicalize(conn, "metric", "crude protrin", "AID") == ("CP", "approved")
        assert canonicalize(conn, "metric", "Lys", "AID") == ("Lysine", "approved")
        assert canonicalize(conn, "metric", "Lys", "SID") == ("Lysine", "approved")
        assert canonicalize(conn, "metric", "ame", "energy") == ("AME", "approved")


def test_legacy_digestibility_sections_are_not_observations(tmp_path: Path):
    workbook = tmp_path / "legacy_context.xlsx"; database = tmp_path / "feed.sqlite"
    book = Workbook(); ws = book.active; ws.title = "成分及营养价值数据"
    ws.append(["ingredient", "Soybean Meal"]); ws.append(["原料编号", 1]); ws.append(["饲料组成", "data"]); ws.append(["CP (%)", 47.2])
    ws.append(["非表观营养价值", "data"]); ws.append(["AID", "data"]); ws.append(["Lys", 0.89]); ws.append(["SID", "data"]); ws.append(["Lys", 0.94])
    meta = book.create_sheet("论文参数"); meta.append(["doi", "title"]); meta.append(["10.1/context", "Test"])
    book.save(workbook)
    result = import_workbook(workbook, database, mode="legacy")
    assert result.observations == 3
    with connect(database) as conn:
        rows = [tuple(row) for row in conn.execute("SELECT measurement_context,canonical_metric,value_num FROM observations ORDER BY source_row")]
    assert rows == [("composition", "CP", 47.2), ("AID", "Lysine", 0.89), ("SID", "Lysine", 0.94)]


def test_legacy_parser_restores_structure_basis_animals_and_digestibility():
    book = Workbook(); ws = book.active; ws.title = "成分及营养价值数据"
    ws.append(["ingredient", "Soybean Meal"])
    ws.append(["原料编号", 7])
    ws.append(["试验编号", 2])
    ws.append(["breed", "Duroc × Landrace × Yorkshire"])
    ws.append(["weight", 25.0])
    ws.append(["dry base", False])
    ws.append(["dry matter", 88.0])
    ws.append(["饲料组成", "data"])
    ws.append(["CP (%)", 47.2])
    ws.append(["特殊成分", "data"])
    ws.append(["NIR", "data"])
    ws.append(["Starch", 2.1])
    ws.append(["能量", "data"])
    ws.append(["ME (MJ/kg)", 13.4])
    ws.append(["营养物质消化率", "data"])
    ws.append(["CP", 0.82])
    ws.append(["非表观营养价值", "data"])
    ws.append(["AID", "data"])
    ws.append(["Lys", 0.89])
    ws.append(["SID", "data"])
    ws.append(["Lys", 0.94])
    ws.append(["STTD", "data"])
    ws.append(["P", 0.61])

    parsed = parse_legacy_ingredient_sheet(ws)
    assert len(parsed.ingredients) == 1
    ingredient = parsed.ingredients[0]
    assert ingredient.source_ingredient_no == "7"
    assert ingredient.trial_no == "2"
    assert ingredient.reported_basis == "as-fed"
    assert ingredient.ingredient_fields["dry matter"] == 88.0
    assert ingredient.animal_fields["breed"] == "Duroc × Landrace × Yorkshire"
    assert ingredient.animal_fields["weight"] == 25.0

    observations = [
        (row.raw_metric, row.section, row.subsection, row.measurement_context, row.value)
        for row in ingredient.observations
    ]
    assert observations == [
        ("CP (%)", "饲料组成", None, "composition", 47.2),
        ("Starch", "特殊成分", "NIR", "composition", 2.1),
        ("ME (MJ/kg)", "能量", None, "energy", 13.4),
        ("CP", "营养物质消化率", None, "ATTD", 0.82),
        ("Lys", "非表观营养价值", "AID", "AID", 0.89),
        ("Lys", "非表观营养价值", "SID", "SID", 0.94),
        ("P", "非表观营养价值", "STTD", "STTD", 0.61),
    ]
    structural_labels = {row.raw_label for row in parsed.structural_rows}
    assert {"饲料组成", "特殊成分", "NIR", "能量", "营养物质消化率", "非表观营养价值", "AID", "SID", "STTD"} <= structural_labels


def test_release_statistics_enforce_scope(tmp_path: Path):
    workbook = tmp_path / "legacy.xlsx"; database = tmp_path / "feed.sqlite"
    _legacy_book(workbook, "10.1/release", include_diet=True)
    import_workbook(workbook, database, mode="legacy")
    statistics = collect_release_statistics(database)
    assert statistics["quality_checks"]["duplicate_dois"] == 0
    assert statistics["quality_checks"]["diet_sheet_exclusion_passed"] is True
    assert statistics["copyright_boundary"].startswith("No article PDF")
    assert "name_mapping" not in statistics
    terminology = statistics["terminology_review"]
    assert "candidate_terms_requiring_review" in terminology
    assert "candidate_metric_terms_by_measurement_context" not in terminology
    assert "measurement_context_note" not in terminology


def test_batch_grid_allows_blank_doi_blank_cells_and_fuzzy_search(tmp_path: Path):
    database = tmp_path / "batch.sqlite"
    batch = ContributionBatch(
        title="中文期刊无 DOI 的饲料原料研究",
        doi="",
        species="pig",
        composition_metrics=[MetricSpec("CP", "composition", "%", "DM"), MetricSpec("Starch", "composition", "%", "DM")],
        ingredients=[
            IngredientGridRow("I001", "corn", values=["9.5", "63.0"]),
            IngredientGridRow("I002", "wheat", values=["12.1", ""]),
        ],
        animals=[AnimalGridRow("A001", "1", "pig", "Duroc", "male", body_weight_kg="25")],
        nutrition_metrics=[
            MetricSpec("Lysine", "SID", "%", "DM"),
            MetricSpec("Methionine", "SID", "%", "DM"),
        ],
        nutrition_rows=[
            NutritionGridRow("I001", "A001", ["0.94", "0.91"]),
            NutritionGridRow("I002", "A001", ["0.90", ""]),
        ],
    )
    check = validate_contribution_batch(database, batch)
    assert check.errors == []
    result = save_contribution_batch(database, batch, approve_new_terms=True)
    assert (result.ingredient_samples, result.animal_groups, result.component_observations, result.nutrition_observations) == (2, 1, 3, 3)
    with connect(database) as conn:
        study = conn.execute("SELECT title,doi FROM studies").fetchone()
    assert tuple(study) == (batch.title, None)
    columns, rows = search_database(database, "中文 SID Duroc")
    assert "title" in columns
    assert len(rows) == 3


def test_fuzzy_search_expands_dictionary_aliases(tmp_path: Path):
    database = tmp_path / "alias-search.sqlite"
    batch = ContributionBatch(
        title="Corn composition",
        composition_metrics=[MetricSpec("CP", "composition", "%", "dry matter")],
        ingredients=[IngredientGridRow("I1", "corn", values=["9.5"])],
    )
    save_contribution_batch(database, batch)
    with connect(database) as conn:
        apply_dictionary_mapping(conn, "ingredient", "玉米", "Corn", "approved")
    _, rows = search_database(database, "玉米")
    assert len(rows) == 1
    assert rows[0][3] == "Corn"


def test_editable_metric_catalog_has_scientific_groups_and_units(tmp_path: Path):
    database = tmp_path / "catalog.sqlite"; initialize(database)
    with connect(database) as conn:
        seed_dictionaries(conn)
        essential = conn.execute(
            "SELECT count(*) FROM metric_catalog WHERE context='SID' AND group_name='必需氨基酸' AND active=1"
        ).fetchone()[0]
        nonessential = conn.execute(
            "SELECT count(*) FROM metric_catalog WHERE context='SID' AND group_name='非必需氨基酸' AND active=1"
        ).fetchone()[0]
        energy_units = {row[0] for row in conn.execute(
            "SELECT default_unit FROM metric_catalog WHERE context='energy' AND active=1"
        )}
        groups = {row[0] for row in conn.execute(
            "SELECT DISTINCT group_name FROM metric_catalog WHERE context='composition' AND active=1"
        )}
        aflatoxin_unit = conn.execute(
            "SELECT default_unit FROM metric_catalog WHERE canonical_metric='Aflatoxin B1'"
        ).fetchone()[0]
    assert (essential, nonessential) == (10, 8)
    assert energy_units == {"MJ/kg"}
    assert {"常用组分", "微量元素", "维生素", "抗营养因子", "霉菌毒素"} <= groups
    assert aflatoxin_unit == "µg/kg"


def test_ai_json_interface_reuses_batch_validation_and_writer(tmp_path: Path):
    database = tmp_path / "agent.sqlite"
    payload = {
        "format_version": "1.0",
        "metadata": {"title": "AI-assisted contribution", "doi": None, "species": "pig"},
        "composition_metrics": [
            {"metric": "CP", "unit": "%", "basis": "dry matter"},
            {"metric": "Zinc", "unit": "mg/kg", "basis": "dry matter"},
        ],
        "ingredients": [
            {"ingredient_key": "I001", "raw_name": "corn", "values": [9.5, None]},
        ],
        "animals": [
            {"animal_group_key": "A001", "trial_no": "1", "species": "pig"},
        ],
        "nutrition_metrics": [
            {"metric": "ME", "context": "energy", "unit": "kJ/kg", "basis": "dry matter"},
        ],
        "nutrition_rows": [
            {"ingredient_key": "I001", "animal_group_key": "A001", "values": [13500]},
        ],
    }
    _, check = validate_payload(database, payload)
    assert check.errors == []
    saved = import_payload(database, payload)
    assert saved["saved"] is True
    with connect(database) as conn:
        counts = tuple(conn.execute(
            "SELECT (SELECT count(*) FROM component_observations),"
            "       (SELECT count(*) FROM nutrition_observations)"
        ).fetchone())
        sources = {row[0] for row in conn.execute("SELECT DISTINCT source_workbook FROM source_provenance")}
        units = [tuple(row) for row in conn.execute(
            "SELECT raw_unit,canonical_unit,unit_resolution_method FROM observations ORDER BY observation_id"
        )]
    assert counts == (1, 1)
    assert sources == {"ai-json/AI-assisted contribution"}
    assert units == [("%", "%", "explicit_field"), ("kJ/kg", "kJ/kg", "explicit_field")]


def test_json_omitted_ame_unit_uses_historical_rule_without_changing_value(tmp_path: Path):
    database = tmp_path / "agent-unit.sqlite"
    payload = {
        "format_version": "1.0",
        "metadata": {"title": "AME unit rule"},
        "composition_metrics": [],
        "ingredients": [{"ingredient_key": "I1", "raw_name": "corn", "values": []}],
        "animals": [{"animal_group_key": "A1", "trial_no": "1", "species": "chicken"}],
        "nutrition_metrics": [{"metric": "AME", "context": "energy"}],
        "nutrition_rows": [{"ingredient_key": "I1", "animal_group_key": "A1", "values": [13.25]}],
    }
    result = import_payload(database, payload)
    assert result["saved"] is True
    with connect(database) as conn:
        row = conn.execute(
            "SELECT value_num,raw_unit,canonical_unit,unit_resolution_method,unit_rule_reference,validation_status "
            "FROM nutrition_observations"
        ).fetchone()
    assert tuple(row) == (
        13.25, "", "MJ/kg", "approved_metric_rule",
        "legacy_dictionary:nutrient_name:11026", "accepted",
    )


def test_digestibility_blank_unit_uses_percent_context_default(tmp_path: Path):
    database = tmp_path / "digestibility-unit.sqlite"
    batch = ContributionBatch(
        title="Digestibility unit ambiguity",
        ingredients=[IngredientGridRow("I1", "corn", values=[])],
        animals=[AnimalGridRow("A1", trial_no="1", species="pig")],
        nutrition_metrics=[MetricSpec("Lysine", "SID")],
        nutrition_rows=[NutritionGridRow("I1", "A1", values=["0.82"])],
    )
    save_contribution_batch(database, batch)
    with connect(database) as conn:
        row = conn.execute(
            "SELECT value_num,canonical_unit,unit_resolution_method,unit_mapping_status,validation_status "
            "FROM nutrition_observations"
        ).fetchone()
    assert tuple(row) == (0.82, "%", "metric_catalog_default", "approved", "accepted")
    with connect(database) as conn:
        basis = conn.execute(
            "SELECT raw_basis,canonical_basis,basis_mapping_status FROM nutrition_observations"
        ).fetchone()
    assert tuple(basis) == ("", "not applicable", "approved")


def test_ai_json_rejects_matrix_width_mismatch(tmp_path: Path):
    payload = {
        "format_version": "1.0", "metadata": {"title": "Bad matrix"},
        "composition_metrics": [{"metric": "CP"}, {"metric": "Ash"}],
        "ingredients": [{"ingredient_key": "I1", "raw_name": "corn", "values": [9.5]}],
        "animals": [], "nutrition_metrics": [], "nutrition_rows": [],
    }
    try:
        validate_payload(tmp_path / "bad.sqlite", payload)
    except ValueError as exc:
        assert "composition_metrics 有 2 项" in str(exc)
    else:
        raise AssertionError("matrix width mismatch must be rejected")


def test_json_contract_rejects_missing_extra_and_duplicate_fields(tmp_path: Path):
    complete = {
        "format_version": "1.0", "metadata": {"title": "Contract"},
        "composition_metrics": [{"metric": "CP"}],
        "ingredients": [{"ingredient_key": "I1", "raw_name": "corn", "values": [9.5]}],
        "animals": [], "nutrition_metrics": [], "nutrition_rows": [],
    }
    missing = dict(complete); missing.pop("animals")
    with pytest.raises(ValueError, match="缺少必填字段"):
        batch_from_payload(missing)
    extra = dict(complete); extra["unexpected"] = True
    with pytest.raises(ValueError, match="未定义字段"):
        batch_from_payload(extra)
    duplicate = tmp_path / "duplicate.json"
    duplicate.write_text(
        '{"format_version":"1.0","format_version":"1.0","metadata":{"title":"x"},'
        '"composition_metrics":[],"ingredients":[],"animals":[],"nutrition_metrics":[],"nutrition_rows":[]}',
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="字段重复"):
        read_payload(duplicate)


def test_empty_contribution_is_blocked_and_extreme_values_are_flagged(tmp_path: Path):
    empty = ContributionBatch(title="No scientific data")
    empty_check = validate_contribution_batch(tmp_path / "empty.sqlite", empty)
    assert any("至少需要一个有名称的原料" in error for error in empty_check.errors)
    assert any("至少需要一个原料成分值" in error for error in empty_check.errors)
    extreme = ContributionBatch(
        title="Extreme warnings",
        composition_metrics=[MetricSpec("CP", "composition", "%")],
        ingredients=[IngredientGridRow("I1", "corn", values=["1000"])],
        animals=[AnimalGridRow("A1", trial_no="1", species="pig", body_weight_kg="-5")],
        nutrition_metrics=[MetricSpec("ME", "energy", "kJ/kg")],
        nutrition_rows=[NutritionGridRow("I1", "A1", values=["-100"])],
    )
    check = validate_contribution_batch(tmp_path / "extreme.sqlite", extreme)
    assert not check.errors
    assert any("百分数超出0–100" in warning for warning in check.warnings)
    assert any("体重为负数" in warning for warning in check.warnings)
    assert any("浓度或能量为负数" in warning for warning in check.warnings)


def test_ai_json_batch_can_fill_gui_without_saving(tmp_path: Path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtWidgets import QApplication
    from feeddb.contribution_widget import ContributionEditor

    app = QApplication.instance() or QApplication([])
    database = tmp_path / "gui-fill.sqlite"
    payload = {
        "format_version": "1.0",
        "metadata": {"title": "AI preview only", "doi": None, "species": "pig"},
        "composition_metrics": [{"metric": "CP", "unit": "%", "basis": "dry matter"}],
        "ingredients": [{"ingredient_key": "I9", "raw_name": "corn", "values": [9.5]}],
        "animals": [{"animal_group_key": "A9", "trial_no": "1", "species": "pig"}],
        "nutrition_metrics": [{"metric": "ME", "context": "energy", "unit": "kJ/kg", "basis": "dry matter"}],
        "nutrition_rows": [{"ingredient_key": "I9", "animal_group_key": "A9", "values": [13500]}],
    }
    batch, check = validate_payload(database, payload)
    assert check.errors == []
    editor = ContributionEditor(database); editor.populate_batch(batch)
    visible = editor.collect_batch()
    assert visible.title == "AI preview only"
    assert visible.ingredients[0].raw_name == "corn"
    assert visible.ingredients[0].values == ["9.5"]
    assert visible.nutrition_metrics[0].unit == "kJ/kg"
    assert visible.nutrition_rows[0].values == ["13500"]
    assert not database.exists()
    editor.close(); app.processEvents()


def test_metric_selector_uses_manager_groups_and_group_headers(tmp_path: Path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtCore import Qt
    from PyQt6.QtWidgets import QApplication
    from feeddb.contribution_widget import MetricDialog

    app = QApplication.instance() or QApplication([])
    database = tmp_path / "grouped.sqlite"; initialize(database)
    with connect(database) as conn:
        seed_dictionaries(conn)
        upsert_metric_catalog(conn, "composition", "用户常用", "Custom marker", "mg/kg", "")
    dialog = MetricDialog(database, nutrition=False)
    assert dialog.group.currentText() == dialog.GROUPED_COMMON
    headers = [
        dialog.list.item(index).text().removeprefix("▾ ")
        for index in range(dialog.list.count())
        if dialog.list.item(index).data(Qt.ItemDataRole.UserRole) is None
    ]
    assert {"常用组分", "微量元素", "维生素", "抗营养因子", "霉菌毒素", "用户常用"} <= set(headers)
    dialog.search.setText("微量元素")
    visible_metrics = [
        dialog.list.item(index).text() for index in range(dialog.list.count())
        if not dialog.list.item(index).isHidden()
        and isinstance(dialog.list.item(index).data(Qt.ItemDataRole.UserRole), tuple)
    ]
    assert len(visible_metrics) == 9
    dialog.search.clear()
    dialog.group.setCurrentText("用户常用")
    assert dialog.list.count() == 1
    unit_options = {dialog.unit.itemText(index) for index in range(dialog.unit.count())}
    assert {"%", "g/kg", "mg/kg", "µg/kg", "g/cm3", "g/L", "g/g", "CFU/g", "FTU/kg", "mmol/kg", "unit/mg"} <= unit_options
    dialog.select_current_group()
    assert [(item.raw_metric, item.unit) for item in dialog.selected_specs()] == [("Custom marker", "mg/kg")]
    dialog.close(); app.processEvents()


def test_complete_legacy_excel_becomes_review_draft_before_database_write(tmp_path: Path):
    workbook = tmp_path / "legacy-review.xlsx"
    _legacy_book(workbook, "10.1/review")

    loaded = load_contribution_file(workbook)

    assert loaded.source_kind == "legacy_excel"
    assert loaded.batch.title == "Test"
    assert loaded.batch.doi == "10.1/review"
    assert [(item.raw_metric, item.unit) for item in loaded.batch.composition_metrics] == [("CP", "%")]
    assert loaded.batch.ingredients[0].raw_name == "Corn"
    assert loaded.batch.ingredients[0].values == ["9.5"]


def test_standard_excel_metadata_and_values_share_one_review_draft(tmp_path: Path):
    book = Workbook(); studies = book.active; studies.title = "Studies"
    studies.append(["study_key", "doi", "title", "journal", "publication_year", "species", "country"])
    studies.append(["S1", "", "File-owned metadata", "Journal", 2025, "pig", "China"])
    trials = book.create_sheet("Trials"); trials.append(["trial_key", "study_key", "source_trial_no"]); trials.append(["T1", "S1", "1"])
    ingredients = book.create_sheet("Ingredients"); ingredients.append(["ingredient_key", "study_key", "raw_name", "reported_basis"]); ingredients.append(["I1", "S1", "corn", "dry matter"])
    animals = book.create_sheet("AnimalGroups"); animals.append(["animal_group_key", "trial_key", "species", "breed"]); animals.append(["A1", "T1", "pig", "Duroc"])
    composition = book.create_sheet("CompositionValues"); composition.append(["ingredient_key", "raw_metric", "value", "unit", "basis"]); composition.append(["I1", "CP", 9.5, "%", "dry matter"])
    nutrition = book.create_sheet("NutritionValues"); nutrition.append(["ingredient_key", "animal_group_key", "measurement_context", "raw_metric", "value", "unit", "basis"]); nutrition.append(["I1", "A1", "SID", "Lys", 0.94, "%", "dry matter"])
    workbook = tmp_path / "standard-review.xlsx"; book.save(workbook); book.close()

    loaded = load_contribution_file(workbook)

    assert loaded.source_kind == "standard_excel"
    assert loaded.batch.title == "File-owned metadata"
    assert loaded.batch.doi == ""
    assert loaded.batch.ingredients[0].values == ["9.5"]
    assert loaded.batch.animals[0].breed == "Duroc"
    assert loaded.batch.nutrition_metrics[0].context == "SID"
    assert loaded.batch.nutrition_rows[0].values == ["0.94"]


def test_json_file_is_a_complete_review_input(tmp_path: Path):
    payload = {
        "format_version": "1.0",
        "metadata": {"title": "JSON-owned metadata", "doi": None},
        "composition_metrics": [{"metric": "Ash", "unit": "%"}],
        "ingredients": [{"ingredient_key": "I1", "raw_name": "corn", "values": [1.2]}],
        "animals": [], "nutrition_metrics": [], "nutrition_rows": [],
    }
    source = tmp_path / "contribution.json"
    source.write_text(json.dumps(payload), encoding="utf-8")

    loaded = load_contribution_file(source)

    assert loaded.source_kind == "json"
    assert loaded.batch.title == "JSON-owned metadata"
    assert loaded.batch.ingredients[0].values == ["1.2"]


def test_same_file_source_cannot_be_saved_twice(tmp_path: Path):
    database = tmp_path / "continuous.sqlite"
    batch = ContributionBatch(
        title="Idempotent file import",
        composition_metrics=[MetricSpec("CP", "composition", "%", "dry matter")],
        ingredients=[IngredientGridRow("I1", "corn", values=["9.5"])],
    )
    save_contribution_batch(database, batch, source_prefix="file-import/input.xlsx")
    with pytest.raises(ValueError, match="已经写入"):
        save_contribution_batch(database, batch, source_prefix="file-import/input.xlsx")


def test_main_window_exposes_four_user_level_functions(tmp_path: Path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtWidgets import QApplication, QTabWidget
    from feeddb.gui import MainWindow

    app = QApplication.instance() or QApplication([])
    window = MainWindow(tmp_path / "gui.sqlite")
    tabs = window.centralWidget()
    assert isinstance(tabs, QTabWidget)
    assert [tabs.tabText(index) for index in range(tabs.count())] == [
        "导入与核对", "查询数据", "词典与指标", "查看数据库",
    ]
    management = tabs.widget(2)
    assert [management.tabText(index) for index in range(management.count())] == [
        "待核对新词", "名称词典", "指标分类",
    ]
    window.close(); app.processEvents()


def test_main_window_switches_chinese_and_english_without_changing_data_codes(tmp_path: Path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtWidgets import QApplication, QDialogButtonBox
    from feeddb.gui import MainWindow

    app = QApplication.instance() or QApplication([])
    database = tmp_path / "bilingual.sqlite"
    window = MainWindow(database)
    tabs = window.centralWidget()
    window.language_combo.setCurrentIndex(window.language_combo.findData("en"))
    app.processEvents()
    assert window.windowTitle().startswith("Feed Ingredient Database Assistant")
    assert [tabs.tabText(index) for index in range(tabs.count())] == [
        "Import & Review", "Query Data", "Dictionary & Metrics", "Browse Database",
    ]
    management = tabs.widget(2)
    assert [management.tabText(index) for index in range(management.count())] == [
        "Candidate Terms", "Name Dictionary", "Metric Catalog",
    ]
    assert window.contribution.ingredients_table.horizontalHeaderItem(0).text() == "System Ingredient Key"
    assert window.term_review_manager.table.horizontalHeaderItem(1).text() == "Category"
    assert window.metric_catalog_manager.table.horizontalHeaderItem(3).text() == "Canonical Metric"
    assert window.metric_catalog_manager.context.currentText() == "composition"
    common_index = window.metric_catalog_manager.group.findData("常用组分")
    assert common_index >= 0
    assert window.metric_catalog_manager.group.itemText(common_index) == "Common Components"
    assert window.metric_catalog_manager.group.itemData(common_index) == "常用组分"
    review_common_index = window.term_review_manager.group.findData("常用组分")
    assert window.term_review_manager.group.itemText(review_common_index) == "Common Components"
    from feeddb.contribution_widget import MetricDialog
    metric_dialog = MetricDialog(database, nutrition=False)
    dialog_common_index = metric_dialog.group.findData("常用组分")
    assert metric_dialog.group.itemText(dialog_common_index) == "Common Components"
    assert metric_dialog._selected_group() == "全部常用指标（按分组）"
    assert metric_dialog.buttons.button(QDialogButtonBox.StandardButton.Ok).text() == "OK"
    window.metric_catalog_manager.clear_form()
    window.metric_catalog_manager.context.setCurrentText("energy")
    assert window.metric_catalog_manager.unit.currentText() == "MJ/kg"
    window.metric_catalog_manager.context.setCurrentText("composition")
    assert window.metric_catalog_manager.unit.currentText() == "%"
    with connect(database) as conn:
        contexts = {row[0] for row in conn.execute("SELECT DISTINCT context FROM metric_catalog")}
    assert {"composition", "energy", "AID", "SID"} <= contexts

    window.language_combo.setCurrentIndex(window.language_combo.findData("zh"))
    app.processEvents()
    assert [tabs.tabText(index) for index in range(tabs.count())] == [
        "导入与核对", "查询数据", "词典与指标", "查看数据库",
    ]
    assert window.contribution.ingredients_table.horizontalHeaderItem(0).text() == "系统原料编号"
    assert window.metric_catalog_manager.group.itemText(common_index) == "常用组分"
    metric_dialog.close()
    window.close(); app.processEvents()


def test_main_window_can_open_existing_database_read_only(tmp_path: Path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtWidgets import QApplication
    from feeddb.gui import MainWindow

    app = QApplication.instance() or QApplication([])
    database = tmp_path / "release.sqlite"
    initialize(database)
    window = MainWindow(database, read_only=True)
    assert "只读" in window.statusBar().currentMessage()
    assert str(database.resolve()) in window.statusBar().currentMessage()
    assert not window.centralWidget().widget(0).isEnabled()
    window.close(); app.processEvents()


def test_fuzzy_search_does_not_create_missing_database(tmp_path: Path):
    from feeddb.query import search_database

    missing = tmp_path / "missing.sqlite"
    with pytest.raises(FileNotFoundError):
        search_database(missing, "corn")
    assert not missing.exists()


def test_user_sql_rules_are_read_only_and_export_three_formats(tmp_path: Path):
    database = tmp_path / "rules.sqlite"
    batch = ContributionBatch(
        title="Rule query study",
        composition_metrics=[MetricSpec("CP", "composition", "%", "dry matter")],
        ingredients=[IngredientGridRow("I1", "corn", values=["9.5"])],
    )
    save_contribution_batch(database, batch)
    sql = """WITH selected AS (
                 SELECT s.title,i.canonical_name,o.canonical_metric,o.value_num
                 FROM observations o
                 JOIN ingredients i ON i.ingredient_id=o.ingredient_id
                 JOIN studies s ON s.study_id=i.study_id
                 WHERE i.canonical_name=?
             ) SELECT * FROM selected ORDER BY canonical_metric"""
    result = execute_readonly_query(database, sql, ["Corn"])
    assert result.columns == ("title", "canonical_name", "canonical_metric", "value_num")
    assert result.rows == [("Rule query study", "Corn", "CP", 9.5)]
    for suffix in (".csv", ".json", ".xlsx"):
        output = tmp_path / f"export{suffix}"
        report = export_readonly_query(database, sql, output, ["Corn"])
        assert report["rows"] == 1
        assert output.exists()
    workbook = load_workbook(tmp_path / "export.xlsx", read_only=False, data_only=True)
    assert list(workbook["QueryResult"].values) == [
        ("title", "canonical_name", "canonical_metric", "value_num"),
        ("Rule query study", "Corn", "CP", 9.5),
    ]
    sheet = workbook["QueryResult"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:D2"
    assert sheet["A1"].font.bold is True
    assert sheet.column_dimensions["A"].width >= len("title")
    workbook.close()


def test_user_sql_rules_reject_writes_and_multiple_statements(tmp_path: Path):
    database = tmp_path / "protected.sqlite"; initialize(database)
    with pytest.raises(ValueError, match="SELECT"):
        execute_readonly_query(database, "DELETE FROM studies")
    with pytest.raises(ValueError, match="一条"):
        validate_readonly_sql("SELECT * FROM studies; DELETE FROM studies")
    with connect(database) as conn:
        assert conn.execute("SELECT count(*) FROM studies").fetchone()[0] == 0


def test_gui_runs_custom_readonly_rule(tmp_path: Path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtWidgets import QApplication
    from feeddb.gui import MainWindow

    app = QApplication.instance() or QApplication([])
    database = tmp_path / "gui-rule.sqlite"
    save_contribution_batch(
        database,
        ContributionBatch(
            title="GUI rule study",
            composition_metrics=[MetricSpec("Ash", "composition", "%", "dry matter")],
            ingredients=[IngredientGridRow("I1", "corn", values=["1.2"])],
        ),
    )
    window = MainWindow(database)
    window.rule_sql.setPlainText(
        "SELECT s.title,o.canonical_metric,o.value_num FROM observations o "
        "JOIN ingredients i ON i.ingredient_id=o.ingredient_id "
        "JOIN studies s ON s.study_id=i.study_id"
    )
    window.run_rule_query()
    assert window.rule_model.rows == [("GUI rule study", "Ash", 1.2)]
    assert "只读规则执行成功" in window.rule_status.text()
    window.close(); app.processEvents()


def test_unit_resolution_priority_and_embedded_variants():
    from feeddb.unit_resolution import UnitResolver

    resolver = UnitResolver()
    for label in ("AME, MJ/kg", "AME(MJ/kg)", "AME，MJ/kg", "AME MJ/kg"):
        result = resolver.resolve(label, "energy", "AME")
        assert result.parsed_metric == "AME"
        assert result.raw_unit == "MJ/kg"
        assert result.canonical_unit == "MJ/kg"
        assert result.resolution_method == "embedded_in_metric"
    explicit = resolver.resolve("AME, MJ/kg", "energy", "AME", explicit_unit="kcal/kg")
    assert explicit.raw_unit == "kcal/kg"
    assert explicit.canonical_unit == "kcal/kg"
    assert explicit.resolution_method == "explicit_field"


def test_unit_aliases_metric_rules_and_context_defaults():
    from feeddb.unit_resolution import UnitResolver

    resolver = UnitResolver()
    for spelling in ("μg/kg", "µg/kg", "ug/kg"):
        assert resolver.canonicalize_unit(spelling) == ("µg/kg", "approved")
    ame = resolver.resolve("AME", "energy", "AME")
    assert (ame.canonical_unit, ame.resolution_method) == ("MJ/kg", "approved_metric_rule")
    tmen = resolver.resolve("TMEn", "energy", "TMEn")
    assert (tmen.canonical_unit, tmen.resolution_method, tmen.status) == ("MJ/kg", "approved_metric_rule", "approved")
    assert "user_curation" in tmen.rule_reference
    explicit_tmen = resolver.resolve("TMEn", "energy", "TMEn", explicit_unit="MJ/kg")
    assert (explicit_tmen.canonical_unit, explicit_tmen.resolution_method) == ("MJ/kg", "explicit_field")
    unknown_energy = resolver.resolve("Novel energy", "energy", "Novel energy")
    assert (unknown_energy.canonical_unit, unknown_energy.resolution_method) == ("MJ/kg", "context_default")
    ordinary_component = resolver.resolve("Novel component", "composition", "Novel component")
    assert (ordinary_component.canonical_unit, ordinary_component.resolution_method) == ("%", "context_default")
    aid = resolver.resolve("Novel AID metric", "AID", "Novel AID metric")
    assert (aid.canonical_unit, aid.resolution_method, aid.status) == ("%", "context_default", "approved")


def test_unicode_lowercase_comparison_and_latest_curated_aliases(tmp_path: Path):
    database = tmp_path / "casefold.sqlite"
    initialize(database)
    with connect(database) as conn:
        seed_dictionaries(conn)
        assert canonicalize(conn, "metric", "ＡＳＨ", "composition") == ("Ash", "approved")
        assert canonicalize(conn, "metric", "Ash", "AID") == ("Ash", "approved")
        assert canonicalize(conn, "metric", "scho", "composition") == ("Soluble carbohydrate", "approved")
        assert canonicalize(conn, "metric", "NSP TOTAL", "composition") == (
            "Total non-starch polysaccharides", "approved"
        )
        assert canonicalize(conn, "metric", "Non-starch polysaccharides", "composition") == (
            "Total non-starch polysaccharides", "approved"
        )
        assert canonicalize(conn, "metric", "Glucosinolates", "composition") == (
            "Total glucosinolates", "approved"
        )
        for raw in ("TIs", "TIUS", "Trypsin inhabitor", "trypsin inhibitor units"):
            assert canonicalize(conn, "metric", raw, "composition") == ("Trypsin inhibitor", "approved")
        assert canonicalize(conn, "metric", "Nitrogen free extract", "composition") == (
            "Nitrogen-free extract", "approved"
        )
        assert canonicalize(conn, "metric", "Phytate", "composition") == ("Phytate phosphorus", "approved")


def test_blank_digestibility_percent_still_warns_on_suspicious_value(tmp_path: Path):
    batch = ContributionBatch(
        title="Digestibility range warning",
        ingredients=[IngredientGridRow("I1", "corn")],
        animals=[AnimalGridRow("A1", species="pig")],
        nutrition_metrics=[MetricSpec("IDF", "AID", "")],
        nutrition_rows=[NutritionGridRow("I1", "A1", ["115"])],
    )
    check = validate_contribution_batch(tmp_path / "warning.sqlite", batch)
    assert any("百分数超出0–100" in warning for warning in check.warnings)


def test_energy_value_magnitude_resolution_and_complete_unit_options():
    from feeddb.unit_resolution import UnitResolver

    resolver = UnitResolver()
    magnitude_energy = resolver.resolve("Novel energy", "energy", "Novel energy", reported_value=13.6)
    assert (magnitude_energy.canonical_unit, magnitude_energy.resolution_method, magnitude_energy.status) == (
        "MJ/kg", "value_magnitude_rule", "approved"
    )
    assert "0.1–100" in magnitude_energy.rule_reference
    kilojoule = resolver.resolve("Novel energy", "energy", "Novel energy", reported_value=13_500)
    assert (kilojoule.canonical_unit, kilojoule.resolution_method) == ("kJ/kg", "value_magnitude_rule")
    kilocalorie = resolver.resolve("Novel energy", "energy", "Novel energy", reported_value=3_200)
    assert (kilocalorie.canonical_unit, kilocalorie.resolution_method) == ("kcal/kg", "value_magnitude_rule")
    ambiguous = resolver.resolve("Novel energy", "energy", "Novel energy", reported_value=250)
    assert (ambiguous.canonical_unit, ambiguous.resolution_method, ambiguous.status) == (
        "MJ/kg", "context_default", "approved"
    )
    assert "not distinctive" in ambiguous.warning
    explicit = resolver.resolve("TMEn", "energy", "TMEn", explicit_unit="kJ/kg", reported_value=13.6)
    assert (explicit.canonical_unit, explicit.resolution_method) == ("kJ/kg", "explicit_field")

    available = set(resolver.available_units)
    assert {
        "%", "g/kg", "mg/kg", "µg/kg", "MJ/kg", "kJ/kg", "kcal/kg", "Mcal/kg",
        "g/cm3", "g/L", "g/g", "CFU/g", "FTU/kg", "TIU/mg", "µmol/g", "mmol/kg", "unit/mg",
    } <= available
    toxin_options = resolver.possible_units("composition", "Aflatoxin B1", "霉菌毒素", "µg/kg")
    assert toxin_options[:4] == ["µg/kg", "mg/kg", "ppb", "ppm"]


def test_comma_suffix_unit_extraction_is_bilingual_and_conservative():
    from feeddb.unit_resolution import UnitResolver

    resolver = UnitResolver()
    assert resolver.split_embedded_unit("CP, mg/kg") == ("CP", "mg/kg")
    assert resolver.split_embedded_unit("粗蛋白，g/kg") == ("粗蛋白", "g/kg")
    assert resolver.split_embedded_unit("Novel marker, parts/100") == ("Novel marker", "parts/100")
    assert resolver.resolve("ME，kJ kg−1", "energy", "ME").canonical_unit == "kJ/kg"
    assert resolver.split_embedded_unit("Thiobarbituric acid, absorbance (532 nm)") == (
        "Thiobarbituric acid, absorbance (532 nm)", ""
    )


def test_gui_assigns_editable_defaults_and_preserves_blank_source_unit(tmp_path: Path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtWidgets import QApplication
    from feeddb.contribution_widget import ContributionEditor, UnitReviewDialog

    app = QApplication.instance() or QApplication([])
    database = tmp_path / "unit-review-gui.sqlite"
    editor = ContributionEditor(database)
    batch = ContributionBatch(
        title="Editable default unit",
        composition_metrics=[MetricSpec("CP", "composition", "")],
        ingredients=[IngredientGridRow("I1", "corn", values=["9.5"])],
    )
    prepared = editor.prepare_batch_units(batch)
    spec = prepared.composition_metrics[0]
    assert (spec.unit, spec.source_unit, spec.unit_resolution_method) == (
        "%", "", "metric_catalog_default"
    )
    dialog = UnitReviewDialog([spec])
    dialog.combos[0].setCurrentText("g/kg")
    reviewed = dialog.reviewed_specs()[0]
    assert (reviewed.unit, reviewed.source_unit, reviewed.unit_resolution_method) == (
        "g/kg", "", "user_selected"
    )
    save_contribution_batch(database, replace(batch, composition_metrics=[reviewed]))
    with connect(database) as conn:
        row = conn.execute(
            "SELECT raw_unit,canonical_unit,unit_resolution_method,unit_rule_reference FROM component_observations"
        ).fetchone()
    assert tuple(row) == ("", "g/kg", "user_selected", "GUI unit review")
    dialog.close(); editor.close(); app.processEvents()


def test_blank_tmen_unit_is_saved_with_curated_metric_rule_provenance(tmp_path: Path):
    database = tmp_path / "energy-magnitude.sqlite"
    batch = ContributionBatch(
        title="Energy magnitude inference",
        ingredients=[IngredientGridRow("I1", "corn")],
        animals=[AnimalGridRow("A1", "1", "chicken")],
        nutrition_metrics=[MetricSpec("TMEn", "energy", "")],
        nutrition_rows=[NutritionGridRow("I1", "A1", ["13.6"])],
    )
    check = validate_contribution_batch(database, batch)
    assert not check.errors
    assert any("TMEn" in item and "MJ/kg" in item for item in check.mappings)
    save_contribution_batch(database, batch)
    with connect(database) as conn:
        row = conn.execute(
            """SELECT value_num,raw_unit,canonical_unit,unit_resolution_method,unit_rule_reference
               FROM nutrition_observations"""
        ).fetchone()
    assert tuple(row[:4]) == (13.6, "", "MJ/kg", "approved_metric_rule")
    assert "user_curation" in row[4]

    suspicious_explicit = ContributionBatch(
        title="Explicit unit magnitude conflict",
        ingredients=[IngredientGridRow("I1", "corn")],
        animals=[AnimalGridRow("A1", "1", "chicken")],
        nutrition_metrics=[MetricSpec("Novel energy", "energy", "kJ/kg")],
        nutrition_rows=[NutritionGridRow("I1", "A1", ["13.6"])],
    )
    suspicious_check = validate_contribution_batch(tmp_path / "explicit-conflict.sqlite", suspicious_explicit)
    assert any("可能实际为MJ/kg" in warning for warning in suspicious_check.warnings)


def test_hydrolyzed_asparagine_and_glutamine_use_combined_amino_acid_names(tmp_path: Path):
    database = tmp_path / "name-safety.sqlite"
    initialize(database)
    with connect(database) as conn:
        seed_dictionaries(conn)
        for context in ("composition", "AID", "SID", "ATTD"):
            assert canonicalize(conn, "metric", "Asp", context) == ("Aspartic acid", "approved")
            assert canonicalize(conn, "metric", "asparitc acid", context) == ("Aspartic acid", "approved")
            assert canonicalize(conn, "metric", "Asparagine", context) == ("Aspartic acid", "approved")
        assert canonicalize(conn, "metric", "Glutamine", "composition") == ("Glutamic acid", "approved")
        assert canonicalize(conn, "metric", "Glutamate", "composition") == ("Glutamic acid", "approved")


def test_user_curated_embedded_units_and_metric_defaults():
    from feeddb.unit_resolution import UnitResolver

    resolver = UnitResolver()
    cases = [
        ("DGM (μm)", "composition", "Geometric mean diameter (DGM)", "µm", "embedded_in_metric"),
        ("Electrolyte, mEq/Kg", "composition", "Electrolyte", "mEq/kg", "embedded_in_metric"),
        ("Phytase, phytase unit/kg", "composition", "Phytase activity", "FTU/kg", "embedded_in_metric"),
        ("Urease activity", "composition", "Urease activity", "pH unit rise", "approved_metric_rule"),
        ("Color score L", "composition", "Color score L", "1", "approved_metric_rule"),
        ("Total glucosinolates", "composition", "Total glucosinolates", "µmol/g", "approved_metric_rule"),
    ]
    for label, context, metric, unit, method in cases:
        resolved = resolver.resolve(label, context, metric)
        assert (resolved.canonical_unit, resolved.resolution_method, resolved.status) == (unit, method, "approved")


def test_candidate_review_maps_alias_or_creates_new_metric_without_changing_raw_data(tmp_path: Path):
    database = tmp_path / "term-review.sqlite"
    save_contribution_batch(
        database,
        ContributionBatch(
            title="Candidate review one",
            composition_metrics=[MetricSpec("Crude protien", "composition", "%")],
            ingredients=[IngredientGridRow("I1", "corn", values=["9.5"])],
        ),
        source_prefix="candidate-review/one",
    )
    candidate = next(item for item in list_candidate_terms(database, "metric") if item.raw_term == "Crude protien")
    assert candidate.occurrences == 1
    assert "CP" in candidate.suggestions or candidate.suggestions
    resolve_candidate_as_alias(database, candidate.dictionary_id, "CP", "Confirmed spelling variant")
    with connect(database) as conn:
        row = conn.execute(
            "SELECT raw_metric,canonical_metric,metric_mapping_status,value_num FROM component_observations"
        ).fetchone()
        assert tuple(row) == ("Crude protien", "CP", "approved", 9.5)

    save_contribution_batch(
        database,
        ContributionBatch(
            title="Candidate review two",
            composition_metrics=[MetricSpec("Novel verified marker", "composition", "mg/kg")],
            ingredients=[IngredientGridRow("I1", "corn", values=["2.5"])],
        ),
        source_prefix="candidate-review/two",
    )
    new_candidate = next(
        item for item in list_candidate_terms(database, "metric") if item.raw_term == "Novel verified marker"
    )
    create_standard_from_candidate(
        database, new_candidate.dictionary_id, "Novel verified marker", "其他", "mg/kg", "dry matter",
        "Confirmed from source article",
    )
    with connect(database) as conn:
        catalog = conn.execute(
            """SELECT default_unit,default_basis FROM metric_catalog
               WHERE context='composition' AND canonical_metric='Novel verified marker'"""
        ).fetchone()
        assert tuple(catalog) == ("mg/kg", "dry matter")
        observation = conn.execute(
            """SELECT raw_metric,canonical_metric,metric_mapping_status FROM component_observations
               WHERE raw_metric='Novel verified marker'"""
        ).fetchone()
        assert tuple(observation) == ("Novel verified marker", "Novel verified marker", "approved")


def test_candidate_review_treats_context_as_metadata_and_updates_embedded_unit_variants(tmp_path: Path):
    database = tmp_path / "candidate-embedded-unit.sqlite"
    save_contribution_batch(
        database,
        ContributionBatch(
            title="Embedded-unit candidate",
            composition_metrics=[MetricSpec("Novel marker (mg/kg)", "composition", "")],
            ingredients=[IngredientGridRow("I1", "corn", values=["2.5"])],
        ),
        source_prefix="candidate-review/embedded-unit",
    )
    candidates = list_candidate_terms(database, "metric")
    candidate = next(item for item in candidates if item.raw_term == "Novel marker")
    assert candidate.context == "composition"
    assert candidate.occurrences == 1
    structural_blocks = {"composition", "aid", "sid", "attd", "sttd", "energy"}
    assert all(item.raw_term.casefold() not in structural_blocks for item in candidates)
    create_standard_from_candidate(
        database, candidate.dictionary_id, "Novel marker", "其他", "mg/kg", "dry matter",
        "Confirmed candidate with embedded unit",
    )
    with connect(database) as conn:
        row = conn.execute(
            "SELECT raw_metric,canonical_metric,metric_mapping_status,value_num FROM component_observations"
        ).fetchone()
        assert tuple(row) == ("Novel marker (mg/kg)", "Novel marker", "approved", 2.5)


def test_metric_review_is_global_while_digestibility_context_stays_on_observation(tmp_path: Path):
    database = tmp_path / "global-metric-review.sqlite"
    save_contribution_batch(
        database,
        ContributionBatch(
            title="Shared component across digestibility methods",
            ingredients=[IngredientGridRow("I1", "corn")],
            animals=[AnimalGridRow("A1", species="pig")],
            nutrition_metrics=[
                MetricSpec("Novel shared fiber", "AID"),
                MetricSpec("Novel shared fiber", "ATTD"),
            ],
            nutrition_rows=[NutritionGridRow("I1", "A1", ["4.5", "50.1"])],
        ),
        source_prefix="candidate-review/shared-component",
    )
    candidates = [item for item in list_candidate_terms(database, "metric") if item.raw_term == "Novel shared fiber"]
    assert len(candidates) == 1
    assert candidates[0].context == "AID / ATTD"
    create_standard_from_candidate(
        database, candidates[0].dictionary_id, "Shared fiber component", "纤维组分", "%", "",
        "One component name shared by AID and ATTD observations",
    )
    with connect(database) as conn:
        rows = [tuple(row) for row in conn.execute(
            "SELECT measurement_context,canonical_metric,metric_mapping_status,value_num "
            "FROM nutrition_observations ORDER BY measurement_context"
        )]
        remaining = conn.execute(
            "SELECT count(*) FROM dictionaries WHERE category='metric' AND raw_term='Novel shared fiber' AND status='candidate'"
        ).fetchone()[0]
    assert rows == [
        ("AID", "Shared fiber component", "approved", 4.5),
        ("ATTD", "Shared fiber component", "approved", 50.1),
    ]
    assert remaining == 0


def test_preflight_validation_never_creates_or_mutates_database(tmp_path: Path):
    missing_database = tmp_path / "must-not-be-created.sqlite"
    batch = ContributionBatch(
        title="Read only preflight",
        composition_metrics=[MetricSpec("Never seen metric", "composition", "mystery-unit")],
        ingredients=[IngredientGridRow("I1", "Never seen ingredient", values=["1.25"])],
    )
    check = validate_contribution_batch(missing_database, batch)
    assert not check.errors
    assert not missing_database.exists()

    existing_database = tmp_path / "existing.sqlite"
    initialize(existing_database)
    with connect(existing_database) as conn:
        seed_dictionaries(conn)
        before = conn.execute("SELECT count(*) FROM dictionaries").fetchone()[0]
    check = validate_contribution_batch(existing_database, batch)
    assert not check.errors
    with connect(existing_database) as conn:
        after = conn.execute("SELECT count(*) FROM dictionaries").fetchone()[0]
        assert after == before
        assert conn.execute(
            "SELECT count(*) FROM dictionaries WHERE raw_term='Never seen metric'"
        ).fetchone()[0] == 0


def test_formatted_and_uncertain_numbers_keep_original_text(tmp_path: Path):
    database = tmp_path / "raw-number-text.sqlite"
    batch = ContributionBatch(
        title="Raw numeric text",
        composition_metrics=[
            MetricSpec("Count marker", "composition", "mg/kg"),
            MetricSpec("Uncertain marker", "composition", "%"),
        ],
        ingredients=[IngredientGridRow("I1", "corn", values=["12,450", "0.89 ± 0.03"])],
    )
    check = validate_contribution_batch(database, batch)
    assert sum("同时保存数值和原始文本" in warning for warning in check.warnings) == 2
    save_contribution_batch(database, batch, source_prefix="numeric-text/source")
    with connect(database) as conn:
        rows = conn.execute(
            "SELECT value_num,value_text FROM component_observations ORDER BY component_observation_id"
        ).fetchall()
        assert [tuple(row) for row in rows] == [(12450.0, "12,450"), (0.89, "0.89 ± 0.03")]


def test_exact_duplicate_is_blocked_for_every_source_kind(tmp_path: Path):
    database = tmp_path / "all-source-duplicates.sqlite"
    batch = ContributionBatch(
        title="Same JSON contribution",
        composition_metrics=[MetricSpec("CP", "composition", "%")],
        ingredients=[IngredientGridRow("I1", "corn", values=["9.5"])],
    )
    save_contribution_batch(database, batch, source_prefix="ai-json")
    with pytest.raises(ValueError, match="已经写入"):
        save_contribution_batch(database, batch, source_prefix="ai-json")


def test_approved_unknown_unit_is_reused_by_future_imports(tmp_path: Path):
    database = tmp_path / "unit-review.sqlite"
    first = ContributionBatch(
        title="Unknown unit one",
        composition_metrics=[MetricSpec("CP", "composition", "parts/100")],
        ingredients=[IngredientGridRow("I1", "corn", values=["9.5"])],
    )
    save_contribution_batch(database, first, source_prefix="unit-review/one")
    candidate = next(item for item in list_candidate_terms(database, "unit") if item.raw_term == "parts/100")
    create_standard_from_candidate(database, candidate.dictionary_id, "%", note="Confirmed equivalent unit")
    with connect(database) as conn:
        first_row = conn.execute(
            "SELECT raw_unit,canonical_unit,unit_mapping_status FROM component_observations"
        ).fetchone()
        assert tuple(first_row) == ("parts/100", "%", "approved")

    second = ContributionBatch(
        title="Unknown unit two",
        composition_metrics=[MetricSpec("CP", "composition", "parts/100")],
        ingredients=[IngredientGridRow("I1", "corn", values=["10.0"])],
    )
    save_contribution_batch(database, second, source_prefix="unit-review/two")
    with connect(database) as conn:
        rows = conn.execute(
            """SELECT canonical_unit,unit_mapping_status FROM component_observations
               WHERE raw_unit='parts/100' ORDER BY component_observation_id"""
        ).fetchall()
        assert [tuple(row) for row in rows] == [("%", "approved"), ("%", "approved")]


def test_database_admin_previews_and_deletes_one_study_transactionally(tmp_path: Path):
    database = tmp_path / "working.sqlite"
    for title in ("Study one", "Study two"):
        save_contribution_batch(
            database,
            ContributionBatch(
                title=title,
                composition_metrics=[MetricSpec("CP", "composition", "%")],
                ingredients=[IngredientGridRow("I1", "corn", values=["9.5"])],
            ),
            source_prefix=f"test/{title}.xlsx",
        )
    with connect(database) as conn:
        study_id = conn.execute("SELECT study_id FROM studies WHERE title='Study one'").fetchone()[0]
    preview = preview_delete_study(database, study_id)
    assert preview.counts["studies"] == 1
    assert preview.counts["component_observations"] == 1
    assert preview.source_files_affected is False
    delete_study(database, study_id, confirm=True)
    with connect(database) as conn:
        assert conn.execute("SELECT count(*) FROM studies").fetchone()[0] == 1
        assert conn.execute("SELECT count(*) FROM component_observations").fetchone()[0] == 1
        assert conn.execute("PRAGMA foreign_key_check").fetchall() == []
        assert conn.execute("SELECT count(*) FROM operation_log").fetchone()[0] == 1


def test_database_admin_clear_and_file_delete_protections(tmp_path: Path):
    database = tmp_path / "working.db"
    save_contribution_batch(
        database,
        ContributionBatch(
            title="Clear me",
            composition_metrics=[MetricSpec("CP", "composition", "%")],
            ingredients=[IngredientGridRow("I1", "corn", values=["9.5"])],
        ),
    )
    dry_run = clear_database(database, confirm=False)
    assert dry_run.counts["studies"] == 1
    clear_database(database, keep_dictionaries=True, confirm=True)
    with connect(database) as conn:
        assert conn.execute("SELECT count(*) FROM studies").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM dictionaries").fetchone()[0] > 0
    conn.close()
    not_deleted = delete_database_file(database, confirm_filename="wrong.db")
    assert not_deleted["deleted"] is False and database.exists()
    deleted = delete_database_file(database, confirm_filename=database.name, backup=True)
    assert deleted["deleted"] is True and not database.exists()
    assert Path(deleted["backup"]).exists()


def test_rebuild_database_uses_new_file_and_keeps_excel(tmp_path: Path):
    source = tmp_path / "source"; source.mkdir()
    workbook = source / "study.xlsx"
    _legacy_book(workbook, "10.1/rebuild")
    output = tmp_path / "rebuilt.sqlite"
    report = rebuild_database(source, output)
    assert report["studies"] == 1
    assert report["source_files_deleted"] is False
    assert workbook.exists() and output.exists()
    with connect(output) as conn:
        assert conn.execute("SELECT count(*) FROM observations").fetchone()[0] == 1
