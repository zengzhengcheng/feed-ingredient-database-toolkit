from __future__ import annotations

import csv
import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .excel_io import PUBLIC_ARTICLE_METADATA_SHEET, PUBLIC_INGREDIENT_DATA_SHEET


REQUIRED_PATHS = (
    "README.md",
    "data/study_source_index.csv",
    "database/feed_ingredient_release.sqlite",
    "documentation/data_dictionary.csv",
    "documentation/release_contents.csv",
    "documentation/sheet_name_mapping.csv",
    "dictionaries/name_alias_dictionary.csv",
    "dictionaries/metric_catalog.csv",
    "dictionaries/unit_basis_mapping.csv",
    "templates/feed_ingredient_contribution_template.xlsx",
    "templates/contribution.schema.json",
    "validation/release_statistics.json",
    "validation/source_audit_summary.json",
    "validation/working_database_comparison.json",
    "validation/scope_and_rights_review.csv",
    "validation/duplicate_decisions.csv",
)
EXCHANGE_RELATIONS = (
    "studies",
    "trials",
    "ingredient_samples",
    "animal_groups",
    "ingredient_evaluations",
    "component_observations",
    "nutrition_observations",
    "source_provenance",
    "dictionaries",
    "metric_catalog",
    "import_runs",
    "validation_issues",
    "ingredients",
    "observations",
)
BLOCKED_HEADERS = {
    "luru_preson",
    "luru_person",
    "entry_person",
    "input_person",
    "录入人员",
    "录入人",
    "录入者",
    "cunyi",
    "存疑",
    "password",
    "passwd",
    "database_password",
    "local_path",
    "file_path",
    "本地路径",
    "文件路径",
}
ABSOLUTE_PATH_PATTERNS = (
    re.compile(r"[A-Za-z]:[\\/]Users[\\/]", re.I),
    re.compile(r"/Users/", re.I),
    re.compile(r"/home/", re.I),
)


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _item(code: str, passed: bool, message: str, severity: str = "error") -> dict[str, object]:
    return {
        "code": code,
        "status": "pass" if passed else ("warning" if severity == "warning" else "fail"),
        "message": message,
    }


def _check_clean_workbooks(root: Path, index_rows: list[dict[str, str]]) -> list[dict[str, object]]:
    items: list[dict[str, object]] = []
    clean_directory = root / "data" / "editable_workbooks"
    files = sorted(clean_directory.glob("*.xlsx")) if clean_directory.exists() else []
    items.append(_item("clean_workbook_count", len(files) == len(index_rows) and bool(files), f"clean={len(files)}, index={len(index_rows)}"))
    required_sheets = [PUBLIC_ARTICLE_METADATA_SHEET, PUBLIC_INGREDIENT_DATA_SHEET]
    bad_sheets: list[str] = []
    blocked_headers: list[str] = []
    missing_titles: list[str] = []
    for path in files:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet_names = [sheet.title for sheet in workbook.worksheets]
        if sheet_names != required_sheets:
            bad_sheets.append(f"{path.name}: {' | '.join(sheet_names)}")
        parameter = next(
            (sheet for sheet in workbook.worksheets if sheet.title == PUBLIC_ARTICLE_METADATA_SHEET),
            None,
        )
        if parameter is None:
            missing_titles.append(path.name)
        else:
            headers = ["" if cell.value is None else str(cell.value).strip() for cell in parameter[1]]
            for header in headers:
                if header.casefold() in BLOCKED_HEADERS:
                    blocked_headers.append(f"{path.name}:{header}")
            header_map = {header.casefold(): index for index, header in enumerate(headers)}
            title_index = header_map.get("title")
            title = parameter.cell(row=2, column=title_index + 1).value if title_index is not None else None
            if title is None or not str(title).strip():
                missing_titles.append(path.name)
        workbook.close()
    items.append(_item("clean_workbook_sheet_scope", not bad_sheets, "Only the two allowed sheets are present" if not bad_sheets else "; ".join(bad_sheets)))
    items.append(_item("clean_workbook_internal_headers", not blocked_headers, "No blocked internal metadata headers" if not blocked_headers else "; ".join(blocked_headers)))
    items.append(_item("clean_workbook_titles", not missing_titles, "All clean workbooks have a title" if not missing_titles else "; ".join(missing_titles)))
    return items


def _check_text_paths(root: Path) -> list[str]:
    hits: list[str] = []
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".csv", ".json", ".md", ".txt"}:
            continue
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            continue
        if any(pattern.search(text) for pattern in ABSOLUTE_PATH_PATTERNS):
            hits.append(path.relative_to(root).as_posix())
    return hits


def _check_database(root: Path) -> tuple[list[dict[str, object]], dict[str, int]]:
    items: list[dict[str, object]] = []
    counts: dict[str, int] = {}
    database = root / "database" / "feed_ingredient_release.sqlite"
    if not database.exists():
        return [_item("release_database", False, "Release database is missing")], counts
    with sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True) as conn:
        integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
        foreign_keys = conn.execute("PRAGMA foreign_key_check").fetchall()
        items.append(_item("sqlite_integrity", integrity == "ok", f"integrity_check={integrity}"))
        items.append(_item("sqlite_foreign_keys", not foreign_keys, f"foreign_key_errors={len(foreign_keys)}"))
        for relation in EXCHANGE_RELATIONS:
            exists = conn.execute(
                "SELECT 1 FROM sqlite_master WHERE name=? AND type IN ('table','view')", (relation,)
            ).fetchone()
            if exists:
                counts[relation] = conn.execute(f"SELECT count(*) FROM {relation}").fetchone()[0]
        source_values = [
            row[0]
            for row in conn.execute(
                """SELECT DISTINCT source_workbook FROM (
                       SELECT source_workbook FROM studies
                       UNION ALL SELECT source_workbook FROM ingredient_samples
                       UNION ALL SELECT source_workbook FROM component_observations
                       UNION ALL SELECT source_workbook FROM nutrition_observations
                   )"""
            )
        ]
        invalid_sources = [value for value in source_values if not re.fullmatch(r"study_\d{4}\.xlsx", value or "")]
        items.append(_item("public_source_locators", not invalid_sources, "All database source locators use public workbook IDs" if not invalid_sources else repr(invalid_sources)))
        invalid_sheets = conn.execute(
            """SELECT count(*) FROM source_provenance
               WHERE trim(coalesce(source_sheet,''))<>'' AND source_sheet NOT IN (?,?)""",
            (PUBLIC_ARTICLE_METADATA_SHEET, PUBLIC_INGREDIENT_DATA_SHEET),
        ).fetchone()[0]
        items.append(_item(
            "public_source_sheets",
            invalid_sheets == 0,
            f"non_public_source_sheet_rows={invalid_sheets}",
        ))
        duplicate_dois = conn.execute(
            """SELECT count(*) FROM (
                   SELECT lower(doi) FROM studies WHERE trim(coalesce(doi,''))<>''
                   GROUP BY lower(doi) HAVING count(*)>1
               )"""
        ).fetchone()[0]
        missing_titles = conn.execute("SELECT count(*) FROM studies WHERE trim(coalesce(title,''))='' ").fetchone()[0]
        held = conn.execute("SELECT count(*) FROM observations WHERE validation_status='held'").fetchone()[0]
        rejected = conn.execute("SELECT count(*) FROM observations WHERE validation_status='rejected'").fetchone()[0]
        items.append(_item("duplicate_dois", duplicate_dois == 0, f"duplicate_doi_groups={duplicate_dois}"))
        items.append(_item("database_titles", missing_titles == 0, f"missing_titles={missing_titles}"))
        items.append(_item("validation_status", held == 0 and rejected == 0, f"held={held}, rejected={rejected}"))
    return items, counts


def check_release_package(release_directory: str | Path, write_report: bool = True) -> dict[str, object]:
    root = Path(release_directory).resolve()
    if not root.is_dir():
        raise NotADirectoryError(root)
    checks: list[dict[str, object]] = []

    missing = [relative for relative in REQUIRED_PATHS if not (root / relative).exists()]
    checks.append(_item("required_files", not missing, "All required files are present" if not missing else f"Missing: {', '.join(missing)}"))

    forbidden_files = [
        path.relative_to(root).as_posix()
        for path in root.rglob("*")
        if path.is_file() and path.suffix.lower() in {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
    ]
    checks.append(_item("copyrighted_or_image_files", not forbidden_files, "No article/image files are in the data package" if not forbidden_files else "; ".join(forbidden_files)))
    software_files = [
        path.relative_to(root).as_posix()
        for path in root.rglob("*")
        if path.is_file() and (
            path.suffix.lower() in {".py", ".pyc", ".pyo", ".ipynb", ".joblib", ".pkl"}
            or path.name.casefold().startswith("test_")
        )
    ]
    checks.append(_item(
        "software_and_model_files",
        not software_files,
        "No source code, tests, or model files are in the data package" if not software_files else "; ".join(software_files),
    ))
    database_files = [
        path.relative_to(root).as_posix()
        for path in (root / "database").rglob("*")
        if path.is_file() and path.suffix.lower() != ".sqlite"
    ] if (root / "database").exists() else []
    checks.append(_item(
        "database_directory_scope",
        not database_files,
        "database/ contains SQLite data files only" if not database_files else "; ".join(database_files),
    ))
    absolute_path_hits = _check_text_paths(root)
    checks.append(_item("local_absolute_paths", not absolute_path_hits, "No user-specific absolute paths in text files" if not absolute_path_hits else "; ".join(absolute_path_hits)))

    index_rows = _csv_rows(root / "data" / "study_source_index.csv")
    checks.extend(_check_clean_workbooks(root, index_rows))
    rights_rows = _csv_rows(root / "validation" / "scope_and_rights_review.csv")
    pending_rights = [row.get("public_workbook_id", "") for row in rights_rows if row.get("rights_status", "").casefold() not in {"approved", "cleared"}]
    checks.append(_item("rights_review", bool(rights_rows) and not pending_rights, "All released workbooks have approved rights review" if rights_rows and not pending_rights else f"pending={len(pending_rights)}"))
    checks.append(_item("data_license", (root / "documentation" / "DATA_LICENSE.txt").exists(), "Data license confirmed" if (root / "documentation" / "DATA_LICENSE.txt").exists() else "Author confirmation of the data license is required"))

    database_checks, database_counts = _check_database(root)
    checks.extend(database_checks)
    exchange_mismatches: list[str] = []
    for relation, expected in database_counts.items():
        csv_path = root / "data" / "exchange_tables" / f"{relation}.csv"
        if not csv_path.exists():
            exchange_mismatches.append(f"{relation}: missing CSV")
            continue
        actual = len(_csv_rows(csv_path))
        if actual != expected:
            exchange_mismatches.append(f"{relation}: database={expected}, csv={actual}")
    checks.append(_item("exchange_table_counts", not exchange_mismatches, "All CSV row counts match SQLite" if not exchange_mismatches else "; ".join(exchange_mismatches)))

    statistics_path = root / "validation" / "release_statistics.json"
    if statistics_path.exists():
        statistics = json.loads(statistics_path.read_text(encoding="utf-8"))
        readiness = statistics.get("release_readiness", {})
        checks.append(_item("statistics_structural_checks", bool(readiness.get("structural_checks_passed")), "Release statistics structural checks passed"))
        checks.append(_item("statistics_data_validation", bool(readiness.get("data_validation_complete")), "Release statistics data validation passed"))
        candidates = statistics.get("terminology_review", {}).get("candidate_terms_requiring_review", 0)
        checks.append(_item(
            "terminology_candidates", candidates == 0,
            f"unique active candidate terms={candidates}; measurement-block labels are excluded and candidates do not invalidate data",
            "warning",
        ))

    comparison_path = root / "validation" / "working_database_comparison.json"
    if comparison_path.exists():
        comparison = json.loads(comparison_path.read_text(encoding="utf-8"))
        checks.append(_item("working_release_counts", bool(comparison.get("all_core_counts_match")), "Core working and rebuilt release counts match" if comparison.get("all_core_counts_match") else "Working and release counts differ; inspect the comparison report"))

    import_issues = _csv_rows(root / "validation" / "import_issues.csv")
    import_errors = [row for row in import_issues if row.get("level") == "error"]
    checks.append(_item("import_errors", not import_errors, f"import_errors={len(import_errors)}"))

    failed = [check for check in checks if check["status"] == "fail"]
    warnings = [check for check in checks if check["status"] == "warning"]
    report = {
        "generated_utc": _utc_now(),
        "release_directory": root.name,
        "ready_for_repository": not failed,
        "failed_checks": len(failed),
        "warnings": len(warnings),
        "checks": checks,
    }
    if write_report:
        destination = root / "validation" / "release_readiness.json"
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report
