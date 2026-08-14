from __future__ import annotations

import csv
import math
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .database import canonicalize, connect, initialize, lookup_dictionary, seed_dictionaries, upsert_dictionary
from .legacy_parser import LegacyIngredientColumn, parse_legacy_ingredient_sheet
from .unit_resolution import DEFAULT_UNIT_RESOLVER, UnitResolution


PUBLIC_ARTICLE_METADATA_SHEET = "Article Metadata"
PUBLIC_INGREDIENT_DATA_SHEET = "Ingredient Data"
PUBLIC_SHEET_NAME_MAPPING = {
    "论文参数": PUBLIC_ARTICLE_METADATA_SHEET,
    "study": PUBLIC_ARTICLE_METADATA_SHEET,
    "studies": PUBLIC_ARTICLE_METADATA_SHEET,
    "成分及营养价值数据": PUBLIC_INGREDIENT_DATA_SHEET,
    "成分及营养价值": PUBLIC_INGREDIENT_DATA_SHEET,
    "ingredients": PUBLIC_INGREDIENT_DATA_SHEET,
    "ingredient data": PUBLIC_INGREDIENT_DATA_SHEET,
}
INGREDIENT_SHEET_NAMES = {"成分及营养价值数据", "成分及营养价值", "ingredients", "ingredient data"}
DIET_SHEET_NAMES = {"配方组成及营养价值", "配方数据库", "配方成分及营养价值", "配方参数"}
PARAMETER_SHEET_NAMES = {"论文参数", "study", "studies", "article metadata"}
STANDARD_SHEETS = {
    "Studies": {"study_key", "doi", "title"},
    "Trials": {"trial_key", "study_key", "source_trial_no"},
    "Ingredients": {"ingredient_key", "study_key", "raw_name"},
    "AnimalGroups": {"animal_group_key", "trial_key"},
    "CompositionValues": {"ingredient_key", "raw_metric", "value"},
    "NutritionValues": {"ingredient_key", "animal_group_key", "measurement_context", "raw_metric", "value"},
}
UNIT_RE = re.compile(r"\((%|g/kg|mg/kg|ug/kg|µg/kg|kcal/kg|kJ/kg|MJ/kg)\)\s*$", re.I)


@dataclass
class Issue:
    level: str
    location: str
    message: str


@dataclass
class ImportResult:
    studies: int = 0
    ingredients: int = 0
    component_observations: int = 0
    nutrition_observations: int = 0
    observations: int = 0
    held: int = 0
    unmapped: int = 0
    issues: list[Issue] = field(default_factory=list)


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _source_locator(source: Path) -> str:
    """Return a portable source locator without exposing a user-specific absolute path."""
    parts = source.parts[-4:] if len(source.parts) >= 4 else source.parts
    return "/".join(parts)


def _numeric(value: object) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = _clean(value).replace(",", "")
    try:
        return float(text)
    except ValueError:
        match = re.fullmatch(r"\s*([-+]?\d+(?:\.\d+)?)\s*(?:±|\+/-).+", text)
        return float(match.group(1)) if match else None


def _split_metric_unit(raw_metric: str) -> tuple[str, str]:
    match = UNIT_RE.search(raw_metric)
    if not match:
        return raw_metric.strip(), ""
    return raw_metric[: match.start()].strip(), match.group(1)


def _find_sheet(workbook, accepted: set[str]):
    accepted_names = {name.casefold() for name in accepted}
    return next((ws for ws in workbook.worksheets if ws.title.strip().casefold() in accepted_names), None)


def _records(worksheet):
    headers = [_clean(cell.value) for cell in worksheet[1]]
    for row_no, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if any(value is not None and _clean(value) for value in row):
            yield row_no, dict(zip(headers, row))


def _study_metadata(workbook, source: Path) -> dict[str, object]:
    worksheet = _find_sheet(workbook, PARAMETER_SHEET_NAMES)
    if not worksheet:
        return {"doi": "", "title": source.stem}
    headers = [_clean(cell.value).casefold() for cell in worksheet[1]]
    values = [cell.value for cell in worksheet[2]]
    data = dict(zip(headers, values))
    return {
        "doi": _clean(data.get("doi")),
        "title": _clean(data.get("title")) or source.stem,
        "journal": _clean(data.get("journal")),
        "publication_year": data.get("publication_year"),
        "species": _clean(data.get("animal")) or _clean(data.get("species")),
        "country": _clean(data.get("country")),
    }


def validate_standard_workbook(path: str | Path) -> list[Issue]:
    source = Path(path)
    workbook = load_workbook(source, read_only=False, data_only=True)
    issues: list[Issue] = []
    try:
        for sheet_name, columns in STANDARD_SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                issues.append(Issue("error", sheet_name, "Required sheet is missing"))
                continue
            headers = {_clean(cell.value) for cell in workbook[sheet_name][1]}
            missing = columns - headers
            if missing:
                issues.append(Issue("error", sheet_name, f"Missing columns: {', '.join(sorted(missing))}"))
            merged = [str(cell_range) for cell_range in workbook[sheet_name].merged_cells.ranges]
            if merged:
                issues.append(Issue(
                    "warning", sheet_name,
                    f"Merged cells in an input table can hide values: {', '.join(merged[:10])}",
                ))
    finally:
        workbook.close()
    return issues


def _load_embedded_dictionary_candidates(workbook, conn: sqlite3.Connection) -> None:
    """Load contributor aliases as candidates; approval remains a curator action."""
    if "DictionaryCandidates" not in workbook.sheetnames:
        return
    for _, record in _records(workbook["DictionaryCandidates"]):
        category = _clean(record.get("category")).casefold()
        raw = _clean(record.get("raw_term"))
        canonical = _clean(record.get("suggested_canonical_term")) or raw
        if category not in {"ingredient", "metric", "unit", "basis", "species"} or not raw:
            continue
        context = _clean(record.get("context")) if category == "metric" else ""
        upsert_dictionary(
            conn, category, raw, canonical, "candidate",
            _clean(record.get("note")) or "Contributor-supplied candidate; curator review required",
            context,
        )


def _add_provenance(
    conn: sqlite3.Connection,
    entity_type: str,
    entity_id: int,
    source_workbook: str,
    source_sheet: str = "",
    source_row: int | None = None,
    source_column: int | None = None,
    note: str = "",
) -> None:
    conn.execute(
        """INSERT OR IGNORE INTO source_provenance(
               entity_type,entity_id,source_workbook,source_sheet,source_row,source_column,source_note
           ) VALUES(?,?,?,?,?,?,?)""",
        (entity_type, entity_id, source_workbook, source_sheet, source_row, source_column, note),
    )


def _upsert_study(conn: sqlite3.Connection, metadata: dict[str, object], source_locator: str) -> int:
    title = _clean(metadata.get("title"))
    if not title:
        raise ValueError("Study title is required; DOI may be blank")
    doi = _clean(metadata.get("doi")) or None
    raw_species = _clean(metadata.get("species"))
    canonical_species, species_status = canonicalize(conn, "species", raw_species)
    row = None
    if doi:
        row = conn.execute("SELECT study_id FROM studies WHERE CASEFOLD(doi)=CASEFOLD(?)", (doi,)).fetchone()
        if row is None:
            rows_without_doi = conn.execute(
                "SELECT study_id FROM studies WHERE doi IS NULL AND CASEFOLD(title)=CASEFOLD(?)", (title,)
            ).fetchall()
            if len(rows_without_doi) == 1:
                row = rows_without_doi[0]
    else:
        row = conn.execute("SELECT study_id FROM studies WHERE CASEFOLD(title)=CASEFOLD(?)", (title,)).fetchone()
    values = (
        doi, title, _clean(metadata.get("journal")), metadata.get("publication_year"),
        raw_species, canonical_species, species_status, _clean(metadata.get("country")), source_locator,
    )
    if row is None:
        return int(conn.execute(
            """INSERT INTO studies(
                   doi,title,journal,publication_year,raw_species,canonical_species,species_mapping_status,country,source_workbook
               ) VALUES(?,?,?,?,?,?,?,?,?) RETURNING study_id""",
            values,
        ).fetchone()[0])
    study_id = int(row[0])
    conn.execute(
        """UPDATE studies SET
               doi=coalesce(doi,?),
               title=?,
               journal=CASE WHEN ?<>'' THEN ? ELSE journal END,
               publication_year=coalesce(?,publication_year),
               raw_species=CASE WHEN ?<>'' THEN ? ELSE raw_species END,
               canonical_species=CASE WHEN ?<>'' THEN ? ELSE canonical_species END,
               species_mapping_status=CASE WHEN ?<>'' THEN ? ELSE species_mapping_status END,
               country=CASE WHEN ?<>'' THEN ? ELSE country END
           WHERE study_id=?""",
        (
            doi, title, values[2], values[2], values[3], raw_species, raw_species,
            canonical_species, canonical_species, species_status if raw_species else "",
            species_status, values[7], values[7], study_id,
        ),
    )
    return study_id


def _get_or_create_trial(
    conn: sqlite3.Connection, study_id: int, source_trial_no: object, treatment: object = "", notes: object = ""
) -> int:
    trial_no = _clean(source_trial_no)
    conn.execute(
        "INSERT OR IGNORE INTO trials(study_id,source_trial_no,treatment,notes) VALUES(?,?,?,?)",
        (study_id, trial_no, _clean(treatment), _clean(notes)),
    )
    return int(conn.execute(
        "SELECT trial_id FROM trials WHERE study_id=? AND source_trial_no=?", (study_id, trial_no)
    ).fetchone()[0])


def _observation_mapping(
    conn: sqlite3.Connection, raw_metric: object, context: str, raw_unit: object, raw_basis: object,
    raw_value: object = None, selected_unit: object | None = None,
    selected_method: str = "", selected_rule: str = "",
) -> tuple[str, str, str, str, str, str, str, str, str, str]:
    metric_label = _clean(raw_metric)
    metric, _ = DEFAULT_UNIT_RESOLVER.split_embedded_unit(metric_label)
    canonical_metric, metric_status = canonicalize(conn, "metric", metric, context)
    catalog = conn.execute(
        """SELECT default_unit FROM metric_catalog
           WHERE CASEFOLD(context)=CASEFOLD(?) AND CASEFOLD(canonical_metric)=CASEFOLD(?) AND active=1
           ORDER BY sort_order,metric_catalog_id LIMIT 1""",
        (context, canonical_metric),
    ).fetchone()
    resolution = DEFAULT_UNIT_RESOLVER.resolve(
        metric_label, context, canonical_metric, raw_unit, catalog[0] if catalog else "", _numeric(raw_value)
    )
    unit = resolution.raw_unit
    canonical_unit = resolution.canonical_unit
    unit_status = resolution.status if canonical_unit else "missing"
    if selected_unit is not None:
        chosen = _clean(selected_unit)
        canonical_unit, unit_status = DEFAULT_UNIT_RESOLVER.canonicalize_unit(chosen)
        if unit and DEFAULT_UNIT_RESOLVER.canonicalize_unit(unit)[1] == "candidate":
            known_source = lookup_dictionary(conn, "unit", unit)
            if known_source is None:
                upsert_dictionary(
                    conn, "unit", unit, canonical_unit or chosen, "candidate",
                    "Source unit assigned to an editable review choice; confirm before global approval",
                )
        resolution = UnitResolution(
            resolution.raw_metric_label, resolution.parsed_metric, unit, canonical_unit,
            selected_method or "user_selected", selected_rule or "GUI unit review", unit_status,
        )
    if unit and unit_status == "candidate":
        known_unit = lookup_dictionary(conn, "unit", unit)
        if known_unit is None:
            canonicalize(conn, "unit", unit)
        else:
            canonical_unit, unit_status = known_unit
    basis = _clean(raw_basis)
    if basis:
        canonical_basis, basis_status = canonicalize(conn, "basis", basis)
    elif context.casefold() in {"aid", "sid", "attd", "sttd", "cattd"}:
        canonical_basis, basis_status = "not applicable", "approved"
    else:
        canonical_basis, basis_status = "", "missing"
    return (
        metric, canonical_metric, metric_status, unit, canonical_unit, unit_status,
        resolution.resolution_method, resolution.rule_reference, canonical_basis, basis_status,
    )


def _value_parts(value: object) -> tuple[float | None, str | None]:
    number = _numeric(value)
    raw = _clean(value)
    if number is None:
        return None, raw
    if isinstance(value, str) and not re.fullmatch(
        r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?", raw
    ):
        # Keep formatting, qualifiers, ranges, and uncertainty text even when a
        # useful central numeric value can also be extracted.
        return number, raw
    return number, None


def _is_unmapped(*statuses: str) -> bool:
    return any(status not in {"approved", "missing"} for status in statuses)


def import_standard_workbook(path: str | Path, db_path: str | Path) -> ImportResult:
    source = Path(path)
    issues = validate_standard_workbook(source)
    if any(issue.level == "error" for issue in issues):
        return ImportResult(issues=issues)
    initialize(db_path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    source_locator = _source_locator(source)
    result = ImportResult(issues=issues)
    try:
        with connect(db_path) as conn:
            seed_dictionaries(conn)
            _load_embedded_dictionary_candidates(workbook, conn)
            study_ids: dict[str, int] = {}
            trial_ids: dict[str, int] = {}
            trial_study_ids: dict[str, int] = {}
            ingredient_ids: dict[str, int] = {}
            ingredient_study_ids: dict[str, int] = {}
            ingredient_statuses: dict[str, str] = {}
            animal_group_ids: dict[str, int] = {}
            animal_trial_ids: dict[str, int] = {}

            for row_no, record in _records(workbook["Studies"]):
                study_key = _clean(record.get("study_key"))
                if not study_key:
                    result.issues.append(Issue("error", f"Studies!{row_no}", "study_key is required"))
                    continue
                metadata = dict(record)
                metadata["doi"] = _clean(record.get("doi"))
                study_id = _upsert_study(conn, metadata, source_locator)
                study_ids[study_key] = study_id
                _add_provenance(conn, "study", study_id, source_locator, "Studies", row_no)
                result.studies += 1

            for row_no, record in _records(workbook["Trials"]):
                key = _clean(record.get("trial_key")); study_key = _clean(record.get("study_key"))
                if not key or study_key not in study_ids:
                    result.issues.append(Issue("error", f"Trials!{row_no}", "Valid trial_key and study_key are required"))
                    continue
                trial_id = _get_or_create_trial(
                    conn, study_ids[study_key], record.get("source_trial_no"), record.get("treatment"), record.get("notes")
                )
                trial_ids[key] = trial_id; trial_study_ids[key] = study_ids[study_key]
                _add_provenance(conn, "trial", trial_id, source_locator, "Trials", row_no)

            for row_no, record in _records(workbook["Ingredients"]):
                key = _clean(record.get("ingredient_key")); study_key = _clean(record.get("study_key"))
                raw_name = _clean(record.get("raw_name"))
                if not key or study_key not in study_ids or not raw_name:
                    result.issues.append(Issue("error", f"Ingredients!{row_no}", "Valid ingredient_key, study_key, and raw_name are required"))
                    continue
                canonical_name, name_status = canonicalize(conn, "ingredient", raw_name)
                cursor = conn.execute(
                    """INSERT INTO ingredient_samples(
                           study_id,source_ingredient_no,raw_name,canonical_name,name_mapping_status,feed_type,
                           processing_technic,place_of_origin,reported_basis,dry_matter_value,
                           source_workbook,source_sheet,first_source_column
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING ingredient_sample_id""",
                    (
                        study_ids[study_key], _clean(record.get("source_ingredient_no")) or key,
                        raw_name, canonical_name, name_status, _clean(record.get("feed_type")),
                        _clean(record.get("processing_technic")), _clean(record.get("place_of_origin")),
                        _clean(record.get("reported_basis")), _numeric(record.get("dry_matter_value")),
                        source_locator, "Ingredients", row_no,
                    ),
                )
                sample_id = int(cursor.fetchone()[0])
                ingredient_ids[key] = sample_id; ingredient_study_ids[key] = study_ids[study_key]
                ingredient_statuses[key] = name_status
                _add_provenance(conn, "ingredient_sample", sample_id, source_locator, "Ingredients", row_no)
                result.ingredients += 1

            for row_no, record in _records(workbook["AnimalGroups"]):
                key = _clean(record.get("animal_group_key")); trial_key = _clean(record.get("trial_key"))
                if not key or trial_key not in trial_ids:
                    result.issues.append(Issue("error", f"AnimalGroups!{row_no}", "Valid animal_group_key and trial_key are required"))
                    continue
                raw_species = _clean(record.get("species"))
                canonical_species, species_status = canonicalize(conn, "species", raw_species)
                cursor = conn.execute(
                    """INSERT INTO animal_groups(
                           trial_id,raw_species,canonical_species,species_mapping_status,breed,sex,animal_type,period,
                           age_days,age_weeks,body_weight_kg,notes,source_workbook,source_sheet,source_column
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING animal_group_id""",
                    (
                        trial_ids[trial_key], raw_species, canonical_species, species_status,
                        _clean(record.get("breed")), _clean(record.get("sex")), _clean(record.get("animal_type")),
                        _clean(record.get("period")), _numeric(record.get("age_days")),
                        _numeric(record.get("age_weeks")), _numeric(record.get("body_weight_kg")),
                        _clean(record.get("notes")), source_locator, "AnimalGroups", row_no,
                    ),
                )
                group_id = int(cursor.fetchone()[0])
                animal_group_ids[key] = group_id; animal_trial_ids[key] = trial_ids[trial_key]
                _add_provenance(conn, "animal_group", group_id, source_locator, "AnimalGroups", row_no)

            component_cache: dict[tuple[object, ...], int] = {}
            for row_no, record in _records(workbook["CompositionValues"]):
                ingredient_key = _clean(record.get("ingredient_key"))
                if ingredient_key not in ingredient_ids:
                    result.issues.append(Issue("error", f"CompositionValues!{row_no}", f"Unknown ingredient_key: {ingredient_key}"))
                    continue
                raw_metric = _clean(record.get("raw_metric"))
                if not raw_metric:
                    result.issues.append(Issue("error", f"CompositionValues!{row_no}", "raw_metric is required"))
                    continue
                mapped = _observation_mapping(conn, raw_metric, "composition", record.get("unit"), record.get("basis"), record.get("value"))
                parsed_metric, canonical_metric, metric_status, raw_unit, canonical_unit, unit_status, unit_method, unit_rule, canonical_basis, basis_status = mapped
                value_num, value_text = _value_parts(record.get("value"))
                key = (ingredient_ids[ingredient_key], raw_metric.casefold(), value_num, value_text, raw_unit, _clean(record.get("basis")), _clean(record.get("qualifier")))
                observation_id = component_cache.get(key)
                if observation_id is None:
                    cursor = conn.execute(
                        """INSERT INTO component_observations(
                               ingredient_sample_id,raw_metric,canonical_metric,value_num,value_text,raw_unit,canonical_unit,
                               unit_resolution_method,unit_rule_reference,
                               raw_basis,canonical_basis,metric_mapping_status,unit_mapping_status,basis_mapping_status,
                               qualifier,section,subsection,validation_status,source_workbook,source_sheet,source_row,source_column
                           ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING component_observation_id""",
                        (
                            ingredient_ids[ingredient_key], raw_metric, canonical_metric, value_num, value_text,
                            raw_unit, canonical_unit, unit_method, unit_rule, _clean(record.get("basis")), canonical_basis,
                            metric_status, unit_status, basis_status, _clean(record.get("qualifier")),
                            "composition", "", "accepted", source_locator, "CompositionValues", row_no, 0,
                        ),
                    )
                    observation_id = int(cursor.fetchone()[0]); component_cache[key] = observation_id
                    result.component_observations += 1; result.observations += 1
                    result.unmapped += int(_is_unmapped(ingredient_statuses[ingredient_key], metric_status, unit_status, basis_status))
                _add_provenance(conn, "component_observation", observation_id, source_locator, "CompositionValues", row_no)

            evaluation_ids: dict[tuple[int, int], int] = {}
            for row_no, record in _records(workbook["NutritionValues"]):
                ingredient_key = _clean(record.get("ingredient_key")); animal_key = _clean(record.get("animal_group_key"))
                if ingredient_key not in ingredient_ids or animal_key not in animal_group_ids:
                    result.issues.append(Issue("error", f"NutritionValues!{row_no}", "Known ingredient_key and animal_group_key are required"))
                    continue
                trial_id = animal_trial_ids[animal_key]
                trial_study_id = int(conn.execute("SELECT study_id FROM trials WHERE trial_id=?", (trial_id,)).fetchone()[0])
                if ingredient_study_ids[ingredient_key] != trial_study_id:
                    result.issues.append(Issue("error", f"NutritionValues!{row_no}", "Ingredient and animal group belong to different studies"))
                    continue
                evaluation_key = (ingredient_ids[ingredient_key], animal_group_ids[animal_key])
                evaluation_id = evaluation_ids.get(evaluation_key)
                if evaluation_id is None:
                    cursor = conn.execute(
                        """INSERT INTO ingredient_evaluations(
                               ingredient_sample_id,trial_id,animal_group_id,source_workbook,source_sheet,source_column
                           ) VALUES(?,?,?,?,?,?) RETURNING ingredient_evaluation_id""",
                        (ingredient_ids[ingredient_key], trial_id, animal_group_ids[animal_key], source_locator, "NutritionValues", row_no),
                    )
                    evaluation_id = int(cursor.fetchone()[0]); evaluation_ids[evaluation_key] = evaluation_id
                    _add_provenance(conn, "ingredient_evaluation", evaluation_id, source_locator, "NutritionValues", row_no)
                raw_context = _clean(record.get("measurement_context"))
                context = raw_context if raw_context.casefold() == "energy" else raw_context.upper()
                validation_status = "accepted" if context else "held"
                context = context or "UNSPECIFIED_NUTRITION"
                raw_metric = _clean(record.get("raw_metric"))
                mapped = _observation_mapping(conn, raw_metric, context, record.get("unit"), record.get("basis"), record.get("value"))
                parsed_metric, canonical_metric, metric_status, raw_unit, canonical_unit, unit_status, unit_method, unit_rule, canonical_basis, basis_status = mapped
                value_num, value_text = _value_parts(record.get("value"))
                cursor = conn.execute(
                    """INSERT INTO nutrition_observations(
                           ingredient_evaluation_id,measurement_context,raw_metric,canonical_metric,value_num,value_text,
                           raw_unit,canonical_unit,unit_resolution_method,unit_rule_reference,
                           raw_basis,canonical_basis,metric_mapping_status,unit_mapping_status,
                           basis_mapping_status,qualifier,section,subsection,validation_status,
                           source_workbook,source_sheet,source_row,source_column
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING nutrition_observation_id""",
                    (
                        evaluation_id, context, raw_metric, canonical_metric, value_num, value_text,
                        raw_unit, canonical_unit, unit_method, unit_rule, _clean(record.get("basis")), canonical_basis,
                        metric_status, unit_status, basis_status, _clean(record.get("qualifier")),
                        "nutrition", "", validation_status, source_locator, "NutritionValues", row_no, 0,
                    ),
                )
                observation_id = int(cursor.fetchone()[0])
                _add_provenance(conn, "nutrition_observation", observation_id, source_locator, "NutritionValues", row_no)
                result.nutrition_observations += 1; result.observations += 1
                result.held += int(validation_status == "held")
                result.unmapped += int(_is_unmapped(ingredient_statuses[ingredient_key], metric_status, unit_status, basis_status))
    finally:
        workbook.close()
    return result


def _legacy_field(fields: dict[str, object], *names: str) -> object:
    for name in names:
        if name in fields:
            return fields[name]
    return ""


def _legacy_sample(
    conn: sqlite3.Connection,
    study_id: int,
    ingredient: LegacyIngredientColumn,
    source_locator: str,
    sheet_name: str,
) -> tuple[int, str, bool]:
    raw_name = ingredient.raw_name
    canonical_name, status = canonicalize(conn, "ingredient", raw_name)
    source_no = ingredient.source_ingredient_no or str(ingredient.source_column - 1)
    cursor = conn.execute(
        """INSERT OR IGNORE INTO ingredient_samples(
               study_id,source_ingredient_no,raw_name,canonical_name,name_mapping_status,feed_type,
               processing_technic,place_of_origin,reported_basis,dry_matter_value,
               source_workbook,source_sheet,first_source_column
           ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING ingredient_sample_id""",
        (
            study_id, source_no, raw_name, canonical_name, status,
            _clean(_legacy_field(ingredient.ingredient_fields, "feedtype")),
            _clean(_legacy_field(ingredient.ingredient_fields, "processing technic")),
            _clean(_legacy_field(ingredient.ingredient_fields, "place of origin")),
            ingredient.reported_basis, _numeric(_legacy_field(ingredient.ingredient_fields, "dry matter")),
            source_locator, sheet_name, ingredient.source_column,
        ),
    )
    inserted = cursor.fetchone()
    if inserted:
        return int(inserted[0]), status, True
    sample_id = conn.execute(
        """SELECT ingredient_sample_id FROM ingredient_samples
           WHERE study_id=? AND source_workbook=? AND source_ingredient_no=? AND raw_name=?""",
        (study_id, source_locator, source_no, raw_name),
    ).fetchone()[0]
    return int(sample_id), status, False


def import_legacy_workbook(path: str | Path, db_path: str | Path) -> ImportResult:
    source = Path(path)
    initialize(db_path)
    workbook = load_workbook(source, read_only=False, data_only=True)
    worksheet = _find_sheet(workbook, INGREDIENT_SHEET_NAMES)
    if not worksheet:
        workbook.close()
        return ImportResult(issues=[Issue("error", source.name, "No recognized ingredient sheet")])
    metadata = _study_metadata(workbook, source)
    source_locator = _source_locator(source)
    result = ImportResult()
    try:
        with connect(db_path) as conn:
            seed_dictionaries(conn)
            study_id = _upsert_study(conn, metadata, source_locator)
            _add_provenance(conn, "study", study_id, source_locator, "论文参数", 2)
            result.studies = 1
            parsed = parse_legacy_ingredient_sheet(worksheet)
            component_cache: dict[tuple[object, ...], int] = {}

            for ingredient in parsed.ingredients:
                if not ingredient.raw_name:
                    continue
                sample_id, ingredient_status, inserted_sample = _legacy_sample(
                    conn, study_id, ingredient, source_locator, worksheet.title
                )
                _add_provenance(conn, "ingredient_sample", sample_id, source_locator, worksheet.title, 1, ingredient.source_column)
                result.ingredients += int(inserted_sample)

                trial_id = _get_or_create_trial(conn, study_id, ingredient.trial_no)
                _add_provenance(conn, "trial", trial_id, source_locator, worksheet.title, None, ingredient.source_column)
                raw_species = _clean(_legacy_field(ingredient.animal_fields, "animal", "animalname")) or _clean(metadata.get("species"))
                canonical_species, species_status = canonicalize(conn, "species", raw_species)
                cursor = conn.execute(
                    """INSERT OR IGNORE INTO animal_groups(
                           trial_id,raw_species,canonical_species,species_mapping_status,breed,sex,animal_type,period,
                           age_days,age_weeks,body_weight_kg,notes,source_workbook,source_sheet,source_column
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING animal_group_id""",
                    (
                        trial_id, raw_species, canonical_species, species_status,
                        _clean(_legacy_field(ingredient.animal_fields, "breed")),
                        _clean(_legacy_field(ingredient.animal_fields, "gender")),
                        _clean(_legacy_field(ingredient.animal_fields, "animaltype", "type")),
                        _clean(_legacy_field(ingredient.animal_fields, "period")),
                        _numeric(_legacy_field(ingredient.animal_fields, "days")),
                        _numeric(_legacy_field(ingredient.animal_fields, "weeks")),
                        _numeric(_legacy_field(ingredient.animal_fields, "weight")),
                        _clean(_legacy_field(ingredient.animal_fields, "cunyi")),
                        source_locator, worksheet.title, ingredient.source_column,
                    ),
                )
                inserted_group = cursor.fetchone()
                animal_group_id = int(inserted_group[0]) if inserted_group else int(conn.execute(
                    """SELECT animal_group_id FROM animal_groups
                       WHERE trial_id=? AND source_workbook=? AND source_sheet=? AND source_column=?""",
                    (trial_id, source_locator, worksheet.title, ingredient.source_column),
                ).fetchone()[0])
                _add_provenance(conn, "animal_group", animal_group_id, source_locator, worksheet.title, None, ingredient.source_column)
                cursor = conn.execute(
                    """INSERT OR IGNORE INTO ingredient_evaluations(
                           ingredient_sample_id,trial_id,animal_group_id,source_workbook,source_sheet,source_column
                       ) VALUES(?,?,?,?,?,?) RETURNING ingredient_evaluation_id""",
                    (sample_id, trial_id, animal_group_id, source_locator, worksheet.title, ingredient.source_column),
                )
                inserted_evaluation = cursor.fetchone()
                evaluation_id = int(inserted_evaluation[0]) if inserted_evaluation else int(conn.execute(
                    """SELECT ingredient_evaluation_id FROM ingredient_evaluations
                       WHERE ingredient_sample_id=? AND trial_id=? AND animal_group_id=?
                         AND source_workbook=? AND source_sheet=? AND source_column=?""",
                    (sample_id, trial_id, animal_group_id, source_locator, worksheet.title, ingredient.source_column),
                ).fetchone()[0])
                _add_provenance(conn, "ingredient_evaluation", evaluation_id, source_locator, worksheet.title, None, ingredient.source_column)

                for observation in ingredient.observations:
                    metric = _clean(observation.raw_metric)
                    raw_unit = ""
                    raw_basis = ingredient.reported_basis if observation.measurement_context in {"composition", "energy"} else ""
                    value_num, value_text = _value_parts(observation.value)
                    if observation.measurement_context == "composition":
                        mapped = _observation_mapping(conn, metric, "composition", raw_unit, raw_basis, observation.value)
                        parsed_metric, canonical_metric, metric_status, raw_unit, canonical_unit, unit_status, unit_method, unit_rule, canonical_basis, basis_status = mapped
                        logical_key = (
                            sample_id, metric.casefold(), value_num, value_text, raw_unit, raw_basis,
                            observation.section, observation.subsection or "",
                        )
                        observation_id = component_cache.get(logical_key)
                        if observation_id is None:
                            cursor = conn.execute(
                                """INSERT INTO component_observations(
                                       ingredient_sample_id,raw_metric,canonical_metric,value_num,value_text,raw_unit,canonical_unit,
                                       unit_resolution_method,unit_rule_reference,
                                       raw_basis,canonical_basis,metric_mapping_status,unit_mapping_status,basis_mapping_status,
                                       qualifier,section,subsection,validation_status,source_workbook,source_sheet,source_row,source_column
                                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING component_observation_id""",
                                (
                                    sample_id, metric, canonical_metric, value_num, value_text, raw_unit, canonical_unit,
                                    unit_method, unit_rule, raw_basis, canonical_basis, metric_status, unit_status, basis_status,
                                    observation.subsection or "", observation.section, observation.subsection or "", "accepted",
                                    source_locator, worksheet.title, observation.source_row, observation.source_column,
                                ),
                            )
                            observation_id = int(cursor.fetchone()[0]); component_cache[logical_key] = observation_id
                            result.component_observations += 1; result.observations += 1
                            result.unmapped += int(_is_unmapped(ingredient_status, metric_status, unit_status, basis_status))
                        _add_provenance(
                            conn, "component_observation", observation_id, source_locator, worksheet.title,
                            observation.source_row, observation.source_column,
                            "Repeated source cells may point to one deduplicated logical component observation",
                        )
                    else:
                        context = observation.measurement_context or "UNSPECIFIED_DIGESTIBILITY"
                        validation_status = "accepted" if observation.measurement_context is not None else "held"
                        mapped = _observation_mapping(conn, metric, context, raw_unit, raw_basis, observation.value)
                        parsed_metric, canonical_metric, metric_status, raw_unit, canonical_unit, unit_status, unit_method, unit_rule, canonical_basis, basis_status = mapped
                        cursor = conn.execute(
                            """INSERT INTO nutrition_observations(
                                   ingredient_evaluation_id,measurement_context,raw_metric,canonical_metric,value_num,value_text,
                                   raw_unit,canonical_unit,unit_resolution_method,unit_rule_reference,
                                   raw_basis,canonical_basis,metric_mapping_status,unit_mapping_status,
                                   basis_mapping_status,qualifier,section,subsection,validation_status,
                                   source_workbook,source_sheet,source_row,source_column
                               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING nutrition_observation_id""",
                            (
                                evaluation_id, context, metric, canonical_metric, value_num, value_text,
                                raw_unit, canonical_unit, unit_method, unit_rule, raw_basis, canonical_basis,
                                metric_status, unit_status, basis_status, observation.subsection or "",
                                observation.section, observation.subsection or "", validation_status,
                                source_locator, worksheet.title, observation.source_row, observation.source_column,
                            ),
                        )
                        observation_id = int(cursor.fetchone()[0])
                        _add_provenance(conn, "nutrition_observation", observation_id, source_locator, worksheet.title, observation.source_row, observation.source_column)
                        result.nutrition_observations += 1; result.observations += 1
                        result.held += int(validation_status == "held")
                        result.unmapped += int(_is_unmapped(ingredient_status, metric_status, unit_status, basis_status))
                        if validation_status == "held":
                            result.issues.append(Issue(
                                "warning", f"{source.name}!{observation.source_row}:{observation.source_column}",
                                "Digestibility value has no named AID/SID/ATTD/STTD or preserved subtype context",
                            ))
    finally:
        workbook.close()
    return result


def import_workbook(path: str | Path, db_path: str | Path, mode: str = "auto") -> ImportResult:
    source = Path(path)
    if source.name.startswith("~$"):
        return ImportResult(issues=[Issue("error", source.name, "Excel temporary lock file is not importable")])
    if mode == "legacy":
        return import_legacy_workbook(source, db_path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        is_standard = (
            set(STANDARD_SHEETS).issubset(workbook.sheetnames)
            or {"Studies", "Ingredients", "Observations"}.issubset(workbook.sheetnames)
        )
    finally:
        workbook.close()
    if is_standard:
        return import_standard_workbook(source, db_path)
    if mode in {"auto", "standard"}:
        from .bilingual_template import is_bilingual_template

        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            is_bilingual = is_bilingual_template(workbook)
        finally:
            workbook.close()
        if is_bilingual:
            from .contribution_batch import save_contribution_batch, validate_contribution_batch
            from .contribution_file import load_contribution_file

            try:
                loaded = load_contribution_file(source)
                check = validate_contribution_batch(db_path, loaded.batch)
                if check.errors:
                    return ImportResult(
                        issues=[Issue("error", source.name, message) for message in check.errors]
                    )
                saved = save_contribution_batch(
                    db_path,
                    loaded.batch,
                    source_prefix=f"excel-contribution/{source.stem}",
                )
                observations = saved.component_observations + saved.nutrition_observations
                return ImportResult(
                    studies=1,
                    ingredients=saved.ingredient_samples,
                    component_observations=saved.component_observations,
                    nutrition_observations=saved.nutrition_observations,
                    observations=observations,
                    unmapped=len(check.new_terms),
                    issues=[Issue("warning", source.name, message) for message in check.warnings],
                )
            except Exception as exc:
                return ImportResult(issues=[Issue("error", source.name, str(exc))])
    if mode == "standard":
        return ImportResult(issues=[Issue(
            "error", source.name,
            "Standard mode requires either the bilingual Article Metadata/Ingredient Data template or the supported multi-sheet schema",
        )])
    return import_legacy_workbook(source, db_path)


def _default_duplicate_resolution_file() -> Path:
    return Path(__file__).resolve().parent.parent / "config" / "duplicate_resolution.csv"


def _load_duplicate_resolution(path: str | Path | None) -> tuple[dict[str, str], set[str]]:
    resolution_path = Path(path) if path else _default_duplicate_resolution_file()
    preferred: dict[str, str] = {}
    excluded: set[str] = set()
    if not resolution_path.exists():
        return preferred, excluded
    with resolution_path.open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            doi = _clean(row.get("doi")).casefold()
            title = _clean(row.get("title")).casefold()
            preferred_path = _clean(row.get("preferred_relative_path")).replace("\\", "/")
            excluded_path = _clean(row.get("excluded_relative_path")).replace("\\", "/")
            if doi and preferred_path:
                preferred[f"doi:{doi}"] = preferred_path
            elif title and preferred_path:
                preferred[f"title:{title}"] = preferred_path
            if excluded_path:
                excluded.add(excluded_path)
    return preferred, excluded


def discover_legacy_workbooks(
    directory: str | Path,
    duplicate_resolution_file: str | Path | None = None,
) -> list[Path]:
    """Return one curated workbook per DOI using the explicit duplicate ledger."""
    root = Path(directory).resolve()
    candidates = sorted(
        path for path in root.rglob("*.xlsx")
        if not path.name.startswith("~$") and "原始" not in path.stem and "副本" not in path.stem
    )
    preferred, excluded = _load_duplicate_resolution(duplicate_resolution_file)
    selected = [path for path in candidates if path.relative_to(root).as_posix() not in excluded]
    by_identity: dict[str, list[Path]] = defaultdict(list)
    for path in selected:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            metadata = _study_metadata(workbook, path)
            doi = _clean(metadata.get("doi")).casefold()
            title = _clean(metadata.get("title")).casefold()
        finally:
            workbook.close()
        identity = f"doi:{doi}" if doi else f"title:{title}"
        by_identity[identity].append(path)
    unresolved: list[str] = []
    for identity, paths in by_identity.items():
        if len(paths) <= 1:
            continue
        preferred_relative = preferred.get(identity)
        matches = [path for path in paths if path.relative_to(root).as_posix() == preferred_relative]
        if len(matches) == 1:
            selected = [path for path in selected if path not in paths or path == matches[0]]
        else:
            unresolved.append(f"{identity}: {', '.join(path.relative_to(root).as_posix() for path in paths)}")
    if unresolved:
        raise ValueError("Unresolved duplicate DOI workbooks; update config/duplicate_resolution.csv: " + "; ".join(unresolved))
    return sorted(selected)


def import_directory(
    directory: str | Path,
    db_path: str | Path,
    duplicate_resolution_file: str | Path | None = None,
) -> ImportResult:
    workbooks = discover_legacy_workbooks(directory, duplicate_resolution_file)
    initialize(db_path)
    with connect(db_path) as conn:
        run_id = conn.execute(
            "INSERT INTO import_runs(source_path,source_kind,status,message) VALUES(?,?,?,?) RETURNING import_id",
            (_source_locator(Path(directory)), "legacy_directory", "running", f"Resolved {len(workbooks)} DOI-deduplicated workbooks"),
        ).fetchone()[0]
    total = ImportResult()
    for path in workbooks:
        try:
            current = import_legacy_workbook(path, db_path)
        except Exception as exc:
            total.issues.append(Issue("error", str(path), str(exc)))
            continue
        total.studies += current.studies
        total.ingredients += current.ingredients
        total.component_observations += current.component_observations
        total.nutrition_observations += current.nutrition_observations
        total.observations += current.observations
        total.held += current.held
        total.unmapped += current.unmapped
        total.issues.extend(current.issues)
    errors = [issue for issue in total.issues if issue.level == "error"]
    with connect(db_path) as conn:
        conn.execute(
            """UPDATE import_runs SET finished_at=CURRENT_TIMESTAMP,status=?,studies_count=?,ingredients_count=?,
                   component_observations_count=?,nutrition_observations_count=?,observations_count=?,
                   held_count=?,unmapped_count=?,message=? WHERE import_id=?""",
            (
                "completed_with_errors" if errors else "completed", total.studies, total.ingredients,
                total.component_observations, total.nutrition_observations, total.observations,
                total.held, total.unmapped, f"Processed {len(workbooks)} workbooks; errors={len(errors)}", run_id,
            ),
        )
    return total
