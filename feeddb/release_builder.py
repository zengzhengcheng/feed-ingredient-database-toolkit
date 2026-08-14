from __future__ import annotations

import csv
import json
import shutil
import sqlite3
import tempfile
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook

from .database import connect, upsert_dictionary
from .excel_io import (
    INGREDIENT_SHEET_NAMES,
    PARAMETER_SHEET_NAMES,
    PUBLIC_ARTICLE_METADATA_SHEET,
    PUBLIC_INGREDIENT_DATA_SHEET,
    PUBLIC_SHEET_NAME_MAPPING,
    _source_locator,
    discover_legacy_workbooks,
    import_directory,
)
from .release_audit import audit_source_tree
from .release_check import check_release_package
from .release_stats import collect_release_statistics, write_release_statistics


PUBLIC_TABLES = (
    "schema_metadata",
    "studies",
    "trials",
    "ingredient_samples",
    "animal_groups",
    "ingredient_evaluations",
    "component_observations",
    "nutrition_observations",
    "source_provenance",
    "dictionaries",
    "metric_catalog",
    "import_runs",
    "validation_issues",
)
PUBLIC_VIEWS = ("ingredients", "observations")
CORE_COUNT_TABLES = (
    "studies",
    "trials",
    "ingredient_samples",
    "animal_groups",
    "ingredient_evaluations",
    "component_observations",
    "nutrition_observations",
    "source_provenance",
)
SOURCE_TABLES = (
    "studies",
    "ingredient_samples",
    "animal_groups",
    "ingredient_evaluations",
    "component_observations",
    "nutrition_observations",
    "source_provenance",
    "validation_issues",
)
BLOCKED_METADATA_HEADERS = {
    "luru_preson",
    "luru_person",
    "entry_person",
    "input_person",
    "录入人员",
    "录入人",
    "录入者",
    "cunyi",
    "存疑",
    "password",
    "passwd",
    "database_password",
    "local_path",
    "file_path",
    "本地路径",
    "文件路径",
}

TABLE_DESCRIPTIONS = {
    "studies": "Bibliographic and study-level metadata; title is required and DOI is optional.",
    "trials": "Experiments reported within a study.",
    "ingredient_samples": "Feed ingredient samples described in each study.",
    "animal_groups": "Animal conditions associated with nutritional-value measurements.",
    "ingredient_evaluations": "Links an ingredient sample to a trial and animal group.",
    "component_observations": "Ingredient composition observations that do not require an animal group.",
    "nutrition_observations": "Energy and digestibility observations interpreted in an animal and trial context.",
    "source_provenance": "Public workbook, sheet, row, and column locators for released entities.",
    "dictionaries": "Alias-to-canonical-name mappings and review states.",
    "metric_catalog": "Editable common metric groups and suggested units used by contribution tools.",
    "import_runs": "Readable import-batch statistics for the release reconstruction.",
    "validation_issues": "Structured issues found during import and validation.",
    "ingredients": "Convenience view over ingredient_samples.",
    "observations": "Convenience view combining component and nutrition observations.",
}

FIELD_DESCRIPTIONS = {
    "doi": "Digital object identifier when available; may be blank.",
    "title": "Source article title; required.",
    "raw_name": "Ingredient name exactly as represented in the source data.",
    "canonical_name": "Dictionary-normalized ingredient name.",
    "name_mapping_status": "Alias mapping state; independent from data validity.",
    "raw_metric": "Measurement name exactly as represented in the source data.",
    "canonical_metric": "Dictionary-normalized measurement name within its context.",
    "measurement_context": "composition, energy, AID, SID, ATTD, STTD, or another preserved subtype.",
    "value_num": "Numeric value when the source value can be represented as a number.",
    "value_text": "Original text value when numeric formatting, qualifiers, or uncertainty must be preserved.",
    "raw_unit": "Unit as reported in the source.",
    "canonical_unit": "Dictionary-normalized unit; the raw unit remains available.",
    "raw_basis": "Reported dry-matter/as-fed basis from the source.",
    "canonical_basis": "Dictionary-normalized reporting basis.",
    "validation_status": "Data interpretation state: accepted, held, or rejected.",
    "source_workbook": "Public clean-workbook filename, not an original local path.",
    "source_sheet": "Released worksheet containing the source value.",
    "source_row": "One-based source worksheet row number when applicable.",
    "source_column": "One-based source worksheet column number when applicable.",
}


@dataclass
class ReleaseBuildResult:
    output_directory: str
    release_version: str
    selected_workbooks: int
    clean_workbooks: int
    studies: int
    ingredient_samples: int
    observations: int
    import_errors: int
    ready_for_repository: bool


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _write_csv(path: Path, rows: Iterable[dict[str, object]], fieldnames: list[str] | None = None) -> int:
    materialized = list(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = fieldnames or (list(materialized[0]) if materialized else [])
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(materialized)
    return len(materialized)


def _find_sheet(workbook, accepted: set[str]):
    names = {name.casefold() for name in accepted}
    return next((sheet for sheet in workbook.worksheets if sheet.title.strip().casefold() in names), None)


def _metadata_from_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = _find_sheet(workbook, PARAMETER_SHEET_NAMES)
        if worksheet is None:
            return {"doi": "", "title": "", "journal": "", "publication_year": "", "country": ""}
        headers = [_clean(cell.value).casefold() for cell in worksheet[1]]
        values = [cell.value for cell in worksheet[2]]
        data = dict(zip(headers, values))
        return {
            "doi": _clean(data.get("doi")),
            "title": _clean(data.get("title")),
            "journal": _clean(data.get("journal")),
            "publication_year": _clean(data.get("publication_year")),
            "country": _clean(data.get("country")),
        }
    finally:
        workbook.close()


def _copy_parameter_sheet(source, target) -> list[str]:
    headers = [_clean(cell.value) for cell in source[1]]
    kept_columns = [
        index
        for index, header in enumerate(headers, start=1)
        if header.casefold() not in BLOCKED_METADATA_HEADERS
    ]
    removed = [headers[index - 1] for index in range(1, len(headers) + 1) if index not in kept_columns]
    for row in source.iter_rows(values_only=True):
        target.append([row[index - 1] if index <= len(row) else None for index in kept_columns])
    return removed


def _copy_value_sheet(source, target) -> None:
    for row in source.iter_rows(values_only=True):
        target.append(list(row))


def _worksheet_part(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    namespace = workbook_root.tag.partition("}")[0].lstrip("{")
    relationship_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    relationship_id = None
    for sheet in workbook_root.findall(f".//{{{namespace}}}sheet"):
        if sheet.get("name") == sheet_name:
            relationship_id = sheet.get(f"{{{relationship_namespace}}}id")
            break
    if not relationship_id:
        raise ValueError(f"worksheet not found in workbook XML: {sheet_name}")

    relationships_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    for relationship in relationships_root:
        if relationship.get("Id") != relationship_id:
            continue
        target = relationship.get("Target", "").replace("\\", "/")
        if target.startswith("/"):
            return target.lstrip("/")
        return f"xl/{target}" if not target.startswith("xl/") else target
    raise ValueError(f"worksheet relationship not found: {sheet_name}")


def _restore_numeric_xml_values(
    source_path: Path,
    source_sheet_name: str,
    destination: Path,
    destination_sheet_name: str,
) -> int:
    """Restore source numeric XML lexemes after openpyxl's value-only copy.

    Excel can store numerically close source cells with different decimal
    lexemes.  A load/save cycle may normalize those lexemes and make the
    importer collapse distinct historical observations.  Public workbooks
    therefore retain the exact numeric ``<v>`` text from the source worksheet
    while still excluding non-public sheets and internal metadata columns.
    """

    spreadsheet_namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    cell_tag = f"{{{spreadsheet_namespace}}}c"
    value_tag = f"{{{spreadsheet_namespace}}}v"
    with zipfile.ZipFile(source_path) as source_archive:
        source_part = _worksheet_part(source_archive, source_sheet_name)
        source_root = ET.fromstring(source_archive.read(source_part))
    numeric_values: dict[str, str] = {}
    for cell in source_root.iter(cell_tag):
        if cell.get("t") not in (None, "n"):
            continue
        value = cell.find(value_tag)
        if cell.get("r") and value is not None and value.text is not None:
            numeric_values[cell.get("r", "")] = value.text

    with zipfile.ZipFile(destination) as target_archive:
        target_part = _worksheet_part(target_archive, destination_sheet_name)
        entries = [(item, target_archive.read(item.filename)) for item in target_archive.infolist()]

    changed = 0
    rewritten: list[tuple[zipfile.ZipInfo, bytes]] = []
    for item, payload in entries:
        if item.filename != target_part:
            rewritten.append((item, payload))
            continue
        target_root = ET.fromstring(payload)
        for cell in target_root.iter(cell_tag):
            source_value = numeric_values.get(cell.get("r", ""))
            value = cell.find(value_tag)
            if source_value is None or value is None or value.text == source_value:
                continue
            value.text = source_value
            changed += 1
        payload = ET.tostring(target_root, encoding="utf-8", xml_declaration=False)
        rewritten.append((item, payload))

    with tempfile.NamedTemporaryFile(
        prefix=f".{destination.stem}-", suffix=".xlsx", dir=destination.parent, delete=False
    ) as stream:
        temporary_path = Path(stream.name)
    try:
        with zipfile.ZipFile(temporary_path, "w") as target_archive:
            for item, payload in rewritten:
                target_archive.writestr(item, payload)
        temporary_path.replace(destination)
    finally:
        temporary_path.unlink(missing_ok=True)
    return changed


def create_clean_workbook(source_path: Path, destination: Path) -> dict[str, object]:
    source = load_workbook(source_path, read_only=True, data_only=True)
    parameter = _find_sheet(source, PARAMETER_SHEET_NAMES)
    ingredient = _find_sheet(source, INGREDIENT_SHEET_NAMES)
    if parameter is None or ingredient is None:
        source.close()
        missing = []
        if parameter is None:
            missing.append("论文参数")
        if ingredient is None:
            missing.append("成分及营养价值数据")
        raise ValueError(f"{source_path.name}: missing allowed sheet(s): {', '.join(missing)}")

    target = Workbook()
    parameter_target = target.active
    parameter_target.title = PUBLIC_ARTICLE_METADATA_SHEET
    ingredient_target = target.create_sheet(PUBLIC_INGREDIENT_DATA_SHEET)
    removed_headers = _copy_parameter_sheet(parameter, parameter_target)
    _copy_value_sheet(ingredient, ingredient_target)
    source.close()

    destination.parent.mkdir(parents=True, exist_ok=True)
    target.save(destination)
    target.close()
    restored_numeric_cells = _restore_numeric_xml_values(
        source_path,
        ingredient.title,
        destination,
        PUBLIC_INGREDIENT_DATA_SHEET,
    )
    return {
        "public_file_name": destination.name,
        "sheets": f"{PUBLIC_ARTICLE_METADATA_SHEET} | {PUBLIC_INGREDIENT_DATA_SHEET}",
        "removed_internal_metadata_fields": len(removed_headers),
        "restored_numeric_xml_cells": restored_numeric_cells,
    }


def _load_rights_review(path: str | Path | None) -> tuple[dict[str, tuple[str, str]], dict[str, tuple[str, str]]]:
    by_workbook: dict[str, tuple[str, str]] = {}
    by_doi: dict[str, tuple[str, str]] = {}
    if path is None:
        return by_workbook, by_doi
    review_path = Path(path)
    if not review_path.exists():
        raise FileNotFoundError(review_path)
    with review_path.open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            status = _clean(row.get("rights_status")).casefold() or "pending"
            note = _clean(row.get("rights_note"))
            workbook_id = _clean(row.get("public_workbook_id"))
            doi = _clean(row.get("doi")).casefold()
            if workbook_id:
                by_workbook[workbook_id] = (status, note)
            if doi:
                by_doi[doi] = (status, note)
    return by_workbook, by_doi


def _write_public_duplicate_decisions(
    duplicate_resolution_file: str | Path | None,
    source_index: list[dict[str, object]],
    output_path: Path,
) -> int:
    by_doi = {
        _clean(row.get("doi")).casefold(): _clean(row.get("public_workbook_id"))
        for row in source_index
        if _clean(row.get("doi"))
    }
    rows: list[dict[str, object]] = []
    if duplicate_resolution_file is not None and Path(duplicate_resolution_file).exists():
        with Path(duplicate_resolution_file).open(encoding="utf-8-sig", newline="") as stream:
            for index, row in enumerate(csv.DictReader(stream), start=1):
                doi = _clean(row.get("doi"))
                title = _clean(row.get("title"))
                if not doi and not title:
                    continue
                rows.append(
                    {
                        "duplicate_group_id": f"duplicate_{index:04d}",
                        "identity_type": "DOI" if doi else "title",
                        "identity_value": doi or title,
                        "retained_public_workbook_id": by_doi.get(doi.casefold(), ""),
                        "excluded_candidate_id": f"excluded_duplicate_{index:04d}",
                        "decision": "retain_more_complete_structured_record",
                        "public_reason": _clean(row.get("public_reason"))
                        or "The excluded workbook represents the same study; the retained structured record is more complete.",
                        "excluded_article_file_included": "no",
                    }
                )
    return _write_csv(
        output_path,
        rows,
        [
            "duplicate_group_id",
            "identity_type",
            "identity_value",
            "retained_public_workbook_id",
            "excluded_candidate_id",
            "decision",
            "public_reason",
            "excluded_article_file_included",
        ],
    )


def _overlay_working_dictionary(release_database: Path, working_database: Path) -> dict[str, int]:
    if not working_database.is_file():
        raise FileNotFoundError(working_database)
    with sqlite3.connect(working_database) as working:
        working.row_factory = sqlite3.Row
        version = working.execute(
            "SELECT metadata_value FROM schema_metadata WHERE metadata_key='schema_version'"
        ).fetchone()
        if version is None:
            raise RuntimeError("Working database has no schema version")
        dictionary_rows = working.execute(
            "SELECT category,context,raw_term,canonical_term,status,note FROM dictionaries"
        ).fetchall()
        catalog_rows = working.execute(
            """SELECT context,group_name,canonical_metric,default_unit,default_basis,
                      sort_order,active,note,updated_at FROM metric_catalog ORDER BY metric_catalog_id"""
        ).fetchall()

    applied = 0
    with connect(release_database) as release:
        release_version = release.execute(
            "SELECT metadata_value FROM schema_metadata WHERE metadata_key='schema_version'"
        ).fetchone()
        if release_version is None or release_version[0] != version[0]:
            raise RuntimeError("Working and release database schemas are not compatible")
        for row in dictionary_rows:
            exists = release.execute(
                "SELECT 1 FROM dictionaries WHERE category=? AND context=? AND CASEFOLD(raw_term)=CASEFOLD(?)",
                (row["category"], row["context"], row["raw_term"]),
            ).fetchone()
            if row["status"] not in {"approved", "deprecated"} and exists is None:
                continue
            upsert_dictionary(
                release,
                row["category"],
                row["raw_term"],
                row["canonical_term"],
                row["status"],
                row["note"],
                row["context"],
            )
            applied += 1

        release.execute("DELETE FROM metric_catalog")
        release.executemany(
            """INSERT INTO metric_catalog(
                   context,group_name,canonical_metric,default_unit,default_basis,
                   sort_order,active,note,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?)""",
            [tuple(row) for row in catalog_rows],
        )
        _synchronize_canonical_fields(release)
    return {"dictionary_rows_applied": applied, "metric_catalog_rows": len(catalog_rows)}


def _synchronize_canonical_fields(conn: sqlite3.Connection) -> None:
    conn.execute(
        """UPDATE ingredient_samples SET
               canonical_name=coalesce((SELECT canonical_term FROM dictionaries d
                   WHERE d.category='ingredient' AND d.context='' AND CASEFOLD(d.raw_term)=CASEFOLD(ingredient_samples.raw_name)),canonical_name),
               name_mapping_status=coalesce((SELECT status FROM dictionaries d
                   WHERE d.category='ingredient' AND d.context='' AND CASEFOLD(d.raw_term)=CASEFOLD(ingredient_samples.raw_name)),name_mapping_status)"""
    )
    for table in ("studies", "animal_groups"):
        conn.execute(
            f"""UPDATE {table} SET
                   canonical_species=coalesce((SELECT canonical_term FROM dictionaries d
                       WHERE d.category='species' AND d.context='' AND CASEFOLD(d.raw_term)=CASEFOLD({table}.raw_species)),canonical_species),
                   species_mapping_status=coalesce((SELECT status FROM dictionaries d
                       WHERE d.category='species' AND d.context='' AND CASEFOLD(d.raw_term)=CASEFOLD({table}.raw_species)),species_mapping_status)"""
        )
    conn.execute(
        """UPDATE component_observations SET
               canonical_metric=coalesce((SELECT canonical_term FROM dictionaries d
                   WHERE d.category='metric' AND d.context='composition' AND CASEFOLD(d.raw_term)=CASEFOLD(component_observations.raw_metric)),canonical_metric),
               metric_mapping_status=coalesce((SELECT status FROM dictionaries d
                   WHERE d.category='metric' AND d.context='composition' AND CASEFOLD(d.raw_term)=CASEFOLD(component_observations.raw_metric)),metric_mapping_status)"""
    )
    conn.execute(
        """UPDATE nutrition_observations SET
               canonical_metric=coalesce((SELECT canonical_term FROM dictionaries d
                   WHERE d.category='metric' AND d.context=nutrition_observations.measurement_context
                     AND CASEFOLD(d.raw_term)=CASEFOLD(nutrition_observations.raw_metric)),canonical_metric),
               metric_mapping_status=coalesce((SELECT status FROM dictionaries d
                   WHERE d.category='metric' AND d.context=nutrition_observations.measurement_context
                     AND CASEFOLD(d.raw_term)=CASEFOLD(nutrition_observations.raw_metric)),metric_mapping_status)"""
    )
    for table in ("component_observations", "nutrition_observations"):
        conn.execute(
            f"""UPDATE {table} SET
                   canonical_unit=coalesce((SELECT canonical_term FROM dictionaries d
                       WHERE d.category='unit' AND d.context='' AND CASEFOLD(d.raw_term)=CASEFOLD({table}.raw_unit)),canonical_unit),
                   unit_mapping_status=coalesce((SELECT status FROM dictionaries d
                       WHERE d.category='unit' AND d.context='' AND CASEFOLD(d.raw_term)=CASEFOLD({table}.raw_unit)),unit_mapping_status),
                   canonical_basis=coalesce((SELECT canonical_term FROM dictionaries d
                       WHERE d.category='basis' AND d.context='' AND CASEFOLD(d.raw_term)=CASEFOLD({table}.raw_basis)),canonical_basis),
                   basis_mapping_status=coalesce((SELECT status FROM dictionaries d
                       WHERE d.category='basis' AND d.context='' AND CASEFOLD(d.raw_term)=CASEFOLD({table}.raw_basis)),basis_mapping_status)"""
        )


def _remap_source_locators(database: Path, mapping: dict[str, str], release_version: str) -> int:
    updated = 0
    with connect(database) as conn:
        for old, new in mapping.items():
            for table in SOURCE_TABLES:
                columns = {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}
                if "source_workbook" not in columns:
                    continue
                cursor = conn.execute(
                    f"UPDATE {table} SET source_workbook=? WHERE source_workbook=?", (new, old)
                )
                updated += cursor.rowcount
        for table in SOURCE_TABLES:
            columns = {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}
            if "source_sheet" not in columns:
                continue
            for old_sheet, public_sheet in PUBLIC_SHEET_NAME_MAPPING.items():
                cursor = conn.execute(
                    f"UPDATE {table} SET source_sheet=? WHERE CASEFOLD(source_sheet)=CASEFOLD(?)",
                    (public_sheet, old_sheet),
                )
                updated += cursor.rowcount
        conn.execute(
            "UPDATE import_runs SET source_path='historical_release_input', source_kind='historical_release_rebuild'"
        )
        conn.execute(
            "INSERT INTO schema_metadata(metadata_key,metadata_value) VALUES('release_version',?) "
            "ON CONFLICT(metadata_key) DO UPDATE SET metadata_value=excluded.metadata_value",
            (release_version,),
        )
        conn.execute(
            "INSERT INTO schema_metadata(metadata_key,metadata_value) VALUES('release_generated_utc',?) "
            "ON CONFLICT(metadata_key) DO UPDATE SET metadata_value=excluded.metadata_value",
            (_utc_now(),),
        )
        remaining = conn.execute(
            """SELECT count(*) FROM (
                   SELECT source_workbook FROM studies
                   UNION ALL SELECT source_workbook FROM ingredient_samples
                   UNION ALL SELECT source_workbook FROM component_observations
                   UNION ALL SELECT source_workbook FROM nutrition_observations
               ) WHERE source_workbook NOT GLOB 'study_[0-9][0-9][0-9][0-9].xlsx'"""
        ).fetchone()[0]
        if remaining:
            raise RuntimeError(f"{remaining} public database source locators were not remapped")
        invalid_sheets = conn.execute(
            """SELECT count(*) FROM source_provenance
               WHERE trim(coalesce(source_sheet,''))<>''
                 AND source_sheet NOT IN (?,?)""",
            (PUBLIC_ARTICLE_METADATA_SHEET, PUBLIC_INGREDIENT_DATA_SHEET),
        ).fetchone()[0]
        if invalid_sheets:
            raise RuntimeError(f"{invalid_sheets} provenance rows still use non-public sheet names")
    return updated


def _export_sqlite_relations(database: Path, output_directory: Path) -> dict[str, int]:
    counts: dict[str, int] = {}
    with sqlite3.connect(database) as conn:
        conn.row_factory = sqlite3.Row
        for relation in (*PUBLIC_TABLES, *PUBLIC_VIEWS):
            columns = [row[1] for row in conn.execute(f"PRAGMA table_info({relation})")]
            rows = [dict(row) for row in conn.execute(f"SELECT * FROM {relation}")]
            counts[relation] = _write_csv(output_directory / f"{relation}.csv", rows, columns)
    return counts


def _write_data_dictionary(database: Path, output: Path) -> int:
    rows: list[dict[str, object]] = []
    with sqlite3.connect(database) as conn:
        for relation in (*PUBLIC_TABLES, *PUBLIC_VIEWS):
            for column in conn.execute(f"PRAGMA table_info({relation})"):
                rows.append(
                    {
                        "table": relation,
                        "table_description": TABLE_DESCRIPTIONS.get(relation, "Release database relation."),
                        "field": column[1],
                        "sqlite_type": column[2],
                        "required": "yes" if column[3] else "no",
                        "primary_key": "yes" if column[5] else "no",
                        "description": FIELD_DESCRIPTIONS.get(column[1], "See the data model and source schema."),
                    }
                )
    return _write_csv(output, rows)


def _database_counts(database: Path) -> dict[str, int]:
    with sqlite3.connect(database) as conn:
        return {table: conn.execute(f"SELECT count(*) FROM {table}").fetchone()[0] for table in CORE_COUNT_TABLES}


def _copy_public_assets(output: Path) -> None:
    project = Path(__file__).resolve().parent.parent
    copies = {
        project / "templates" / "feed_ingredient_contribution_template.xlsx": output / "templates" / "feed_ingredient_contribution_template.xlsx",
        project / "integrations" / "contribution.schema.json": output / "templates" / "contribution.schema.json",
        project / "integrations" / "ai_contribution_example.json": output / "templates" / "ai_contribution_example.json",
        project / "integrations" / "AI_AGENT_GUIDE.md": output / "documentation" / "AI_AGENT_GUIDE.md",
        project / "docs" / "data_model.md": output / "documentation" / "data_model.md",
        project / "docs" / "dictionary_guide.md": output / "documentation" / "dictionary_guide.md",
        project / "docs" / "excel_entry_guide.md": output / "documentation" / "excel_entry_guide.md",
        project / "docs" / "database_query_and_export_guide.md": output / "documentation" / "database_query_and_export_guide.md",
    }
    for source, destination in copies.items():
        if source.exists():
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, destination)


def _write_release_readme(output: Path, version: str, license_name: str) -> None:
    license_line = license_name if license_name.casefold() != "pending" else "Pending author confirmation"
    text = f"""# Historical Feed Ingredient Dataset {version}

This candidate package contains historical feed ingredient metadata, composition observations, and animal-context nutritional values.

Included public worksheets: `{PUBLIC_ARTICLE_METADATA_SHEET}` and `{PUBLIC_INGREDIENT_DATA_SHEET}` only.

Excluded: diet/formulation worksheets, source materials outside the selected release scope, article PDFs, screenshots, OCR files, local paths, passwords, and internal entry-person fields.

The `data/editable_workbooks` files preserve source values in sequentially named public workbooks. `data/exchange_tables` and `database/feed_ingredient_release.sqlite` provide machine-readable normalized representations. Source articles are referenced in `data/study_source_index.csv`; article files are not redistributed.

The SQLite file can be opened directly by standard SQLite clients, Python, R, or other compatible tools. Editable read-only SQL query rules and software examples are distributed in the separate GitHub code package, not in this ScienceDB data package.

Related article or project: Not applicable.

Data license: {license_line}.

Final repository citation and DOI/CSTR must be added after the ScienceDB record is created.
"""
    destination = output / "README.md"
    destination.write_text(text, encoding="utf-8")


def _write_license_files(output: Path, license_name: str) -> None:
    documentation = output / "documentation"
    documentation.mkdir(parents=True, exist_ok=True)
    if license_name.casefold() == "pending":
        (documentation / "DATA_LICENSE_REVIEW_REQUIRED.txt").write_text(
            "The data license has not yet been confirmed by the authors. "
            "Create DATA_LICENSE.txt only after the authors confirm the release license.\n",
            encoding="utf-8",
        )
    else:
        (documentation / "DATA_LICENSE.txt").write_text(
            f"Dataset license selected by the authors: {license_name}\n",
            encoding="utf-8",
        )


def _write_release_contents(output: Path) -> int:
    rows: list[dict[str, object]] = []
    for path in sorted(item for item in output.rglob("*") if item.is_file()):
        if path.name == "release_contents.csv":
            continue
        relative = path.relative_to(output).as_posix()
        role = relative.split("/", 1)[0] if "/" in relative else "package_root"
        rows.append(
            {
                "relative_path": relative,
                "role": role,
                "size_bytes": path.stat().st_size,
                "description": "See README and data dictionary.",
            }
        )
    return _write_csv(output / "documentation" / "release_contents.csv", rows)


def _write_sheet_name_mapping(output: Path) -> int:
    rows = [
        {
            "historical_sheet_name": "论文参数",
            "public_sheet_name": PUBLIC_ARTICLE_METADATA_SHEET,
            "purpose": "study-level bibliographic metadata",
        },
        {
            "historical_sheet_name": "成分及营养价值数据",
            "public_sheet_name": PUBLIC_INGREDIENT_DATA_SHEET,
            "purpose": "ingredient composition and animal-context nutritional values",
        },
    ]
    return _write_csv(output / "documentation" / "sheet_name_mapping.csv", rows)


def _export_dictionary_products(database: Path, output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(database) as conn:
        conn.row_factory = sqlite3.Row
        dictionaries = [dict(row) for row in conn.execute("SELECT * FROM dictionaries ORDER BY category,context,raw_term")]
        catalog = [dict(row) for row in conn.execute("SELECT * FROM metric_catalog ORDER BY context,sort_order,canonical_metric")]
    _write_csv(output / "name_alias_dictionary.csv", dictionaries)
    _write_csv(output / "metric_catalog.csv", catalog)
    _write_csv(
        output / "unit_basis_mapping.csv",
        [row for row in dictionaries if row.get("category") in {"unit", "basis"}],
    )


def build_release_package(
    historical_directory: str | Path,
    working_database: str | Path,
    output_directory: str | Path,
    duplicate_resolution_file: str | Path | None = None,
    rights_review_file: str | Path | None = None,
    release_version: str = "v0.5-candidate",
    data_license: str = "PENDING",
) -> ReleaseBuildResult:
    source_root = Path(historical_directory).resolve()
    working_db = Path(working_database).resolve()
    output = Path(output_directory).resolve()
    if not source_root.is_dir():
        raise NotADirectoryError(source_root)
    if not working_db.is_file():
        raise FileNotFoundError(working_db)
    if output.exists() and any(output.iterdir()):
        raise FileExistsError(f"Output directory is not empty: {output}")
    output.mkdir(parents=True, exist_ok=True)

    workbooks = discover_legacy_workbooks(source_root, duplicate_resolution_file)
    by_workbook, by_doi = _load_rights_review(rights_review_file)
    source_mapping: dict[str, str] = {}
    source_index: list[dict[str, object]] = []
    clean_reports: list[dict[str, object]] = []
    for index, source in enumerate(workbooks, start=1):
        workbook_id = f"study_{index:04d}"
        public_name = f"{workbook_id}.xlsx"
        metadata = _metadata_from_workbook(source)
        rights_status, rights_note = by_workbook.get(
            workbook_id,
            by_doi.get(_clean(metadata["doi"]).casefold(), ("pending", "Author confirmation required")),
        )
        public_workbook = output / "data" / "editable_workbooks" / public_name
        report = create_clean_workbook(source, public_workbook)
        clean_reports.append({"public_workbook_id": workbook_id, **report})
        source_mapping[_source_locator(public_workbook)] = public_name
        doi = _clean(metadata["doi"])
        source_index.append(
            {
                "public_workbook_id": workbook_id,
                "public_file_name": public_name,
                "doi": doi,
                "title": metadata["title"],
                "journal": metadata["journal"],
                "publication_year": metadata["publication_year"],
                "country": metadata["country"],
                "source_reference": f"https://doi.org/{doi}" if doi else _clean(metadata["title"]),
                "released_sheets": f"{PUBLIC_ARTICLE_METADATA_SHEET} | {PUBLIC_INGREDIENT_DATA_SHEET}",
                "article_file_included": "no",
                "rights_status": rights_status,
                "rights_note": rights_note,
            }
        )

    _write_csv(output / "data" / "study_source_index.csv", source_index)
    _write_public_duplicate_decisions(
        duplicate_resolution_file,
        source_index,
        output / "validation" / "duplicate_decisions.csv",
    )
    _write_csv(output / "validation" / "clean_workbook_report.csv", clean_reports)
    _write_csv(
        output / "validation" / "scope_and_rights_review.csv",
        [
            {
                "public_workbook_id": row["public_workbook_id"],
                "doi": row["doi"],
                "title": row["title"],
                "rights_status": row["rights_status"],
                "rights_note": row["rights_note"],
                "article_file_included": "no",
                "diet_or_formulation_data_included": "no",
                "new_23_article_corpus_included": "no",
            }
            for row in source_index
        ],
    )

    release_database = output / "database" / "feed_ingredient_release.sqlite"
    import_result = import_directory(output / "data" / "editable_workbooks", release_database, None)
    overlay = _overlay_working_dictionary(release_database, working_db)
    remapped_rows = _remap_source_locators(release_database, source_mapping, release_version)
    _export_sqlite_relations(release_database, output / "data" / "exchange_tables")
    _write_data_dictionary(release_database, output / "documentation" / "data_dictionary.csv")
    _write_sheet_name_mapping(output)
    _export_dictionary_products(release_database, output / "dictionaries")

    statistics = write_release_statistics(
        release_database, output / "validation" / "release_statistics.json"
    )
    records, collections, audit_summary = audit_source_tree(source_root)
    sanitized_audit = {
        "generated_utc": _utc_now(),
        "scope_rule": "All source files were counted, but only the two allowed workbook sheets are released.",
        "total_source_files": len(records),
        "total_source_collections": len(collections),
        "source_file_roles": audit_summary["file_roles"],
        "selected_workbooks": len(workbooks),
        "clean_public_workbooks": len(clean_reports),
        "article_files_in_public_package": 0,
        "local_source_paths_in_public_index": 0,
    }
    _write_json(output / "validation" / "source_audit_summary.json", sanitized_audit)
    _write_csv(
        output / "validation" / "import_issues.csv",
        [asdict(issue) for issue in import_result.issues],
        ["level", "location", "message"],
    )

    working_counts = _database_counts(working_db)
    release_counts = _database_counts(release_database)
    comparison = {
        "generated_utc": _utc_now(),
        "note": "The working database is editable; differences block this historical release candidate until explained.",
        "working_database_file": working_db.name,
        "release_database_file": release_database.name,
        "tables": {
            table: {
                "working": working_counts[table],
                "release": release_counts[table],
                "matches": working_counts[table] == release_counts[table],
            }
            for table in CORE_COUNT_TABLES
        },
    }
    comparison["all_core_counts_match"] = all(row["matches"] for row in comparison["tables"].values())
    _write_json(output / "validation" / "working_database_comparison.json", comparison)
    _write_json(
        output / "validation" / "build_summary.json",
        {
            "generated_utc": _utc_now(),
            "release_version": release_version,
            "selected_workbooks": len(workbooks),
            "clean_workbooks": len(clean_reports),
            "dictionary_overlay": overlay,
            "source_locator_rows_remapped": remapped_rows,
            "import_result": {
                **asdict(import_result),
                "issues": [asdict(issue) for issue in import_result.issues],
            },
        },
    )

    _copy_public_assets(output)
    _write_release_readme(output, release_version, data_license)
    _write_license_files(output, data_license)
    _write_release_contents(output)
    readiness = check_release_package(output, write_report=True)
    _write_release_contents(output)

    result = ReleaseBuildResult(
        output_directory=output.as_posix(),
        release_version=release_version,
        selected_workbooks=len(workbooks),
        clean_workbooks=len(clean_reports),
        studies=statistics["table_counts"]["studies"],
        ingredient_samples=statistics["table_counts"]["ingredient_samples"],
        observations=statistics["table_counts"]["observations"],
        import_errors=sum(issue.level == "error" for issue in import_result.issues),
        ready_for_repository=readiness["ready_for_repository"],
    )
    return result
