from __future__ import annotations

from collections import Counter, OrderedDict, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .contribution_api import batch_from_payload, read_payload
from .contribution_batch import (
    AnimalGridRow,
    ContributionBatch,
    IngredientGridRow,
    MetricSpec,
    NutritionGridRow,
)
from .excel_io import (
    INGREDIENT_SHEET_NAMES,
    STANDARD_SHEETS,
    _clean,
    _find_sheet,
    _legacy_field,
    _split_metric_unit,
    _study_metadata,
    validate_standard_workbook,
)
from .legacy_parser import parse_legacy_ingredient_sheet
from .bilingual_template import batch_from_bilingual_workbook, is_bilingual_template


@dataclass
class ContributionFileResult:
    batch: ContributionBatch
    source_path: Path
    source_kind: str
    warnings: list[str] = field(default_factory=list)


def _records(worksheet) -> list[dict[str, Any]]:
    headers = [_clean(cell.value) for cell in worksheet[1]]
    return [
        dict(zip(headers, row))
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None and _clean(value) for value in row)
    ]


def _value(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _metric_occurrences(
    rows: list[tuple[str, tuple[str, str, str, str], Any]],
) -> tuple[list[tuple[str, str, str, str, int]], dict[str, dict[tuple[str, str, str, str, int], str]]]:
    """Pivot long observations without overwriting repeated metric rows."""
    order: OrderedDict[tuple[str, str, str, str, int], None] = OrderedDict()
    values: dict[str, dict[tuple[str, str, str, str, int], str]] = defaultdict(dict)
    seen: Counter[tuple[str, tuple[str, str, str, str]]] = Counter()
    for owner_key, metric_key, raw_value in rows:
        seen[(owner_key, metric_key)] += 1
        occurrence_key = (*metric_key, seen[(owner_key, metric_key)])
        order.setdefault(occurrence_key, None)
        values[owner_key][occurrence_key] = _value(raw_value)
    return list(order), values


def _standard_excel_batch(workbook, source: Path) -> ContributionFileResult:
    issues = validate_standard_workbook(source)
    errors = [f"{item.location}: {item.message}" for item in issues if item.level == "error"]
    if errors:
        raise ValueError("标准 Excel 结构不完整：" + "；".join(errors))

    studies = _records(workbook["Studies"])
    if len(studies) != 1:
        raise ValueError(f"当前核对工作台要求一个文件对应一篇论文；Studies 中发现 {len(studies)} 条记录")
    study = studies[0]
    study_key = _clean(study.get("study_key"))
    if not study_key:
        raise ValueError("Studies.study_key 不能为空")
    study_key_folded = study_key.casefold()

    trial_records = _records(workbook["Trials"])
    wrong_trial_studies = sorted({
        _clean(row.get("study_key")) for row in trial_records
        if _clean(row.get("study_key")).casefold() != study_key_folded
    })
    if wrong_trial_studies:
        raise ValueError(f"Trials 引用了当前文件中不存在的研究编号：{', '.join(wrong_trial_studies)}")
    trial_key_list = [_clean(row.get("trial_key")) for row in trial_records]
    duplicate_trials = sorted(key for key, count in Counter(key.casefold() for key in trial_key_list if key).items() if count > 1)
    if duplicate_trials:
        raise ValueError(f"Trials.trial_key 重复：{', '.join(duplicate_trials)}")
    trials = {_clean(row.get("trial_key")).casefold(): row for row in trial_records}

    ingredient_records = _records(workbook["Ingredients"])
    wrong_ingredient_studies = sorted({
        _clean(row.get("study_key")) for row in ingredient_records
        if _clean(row.get("study_key")).casefold() != study_key_folded
    })
    if wrong_ingredient_studies:
        raise ValueError(f"Ingredients 引用了当前文件中不存在的研究编号：{', '.join(wrong_ingredient_studies)}")
    ingredient_key_list = [_clean(row.get("ingredient_key")) for row in ingredient_records]
    duplicate_ingredients = sorted(key for key, count in Counter(key.casefold() for key in ingredient_key_list if key).items() if count > 1)
    if duplicate_ingredients:
        raise ValueError(f"Ingredients.ingredient_key 重复：{', '.join(duplicate_ingredients)}")
    ingredient_keys = {key.casefold(): key for key in ingredient_key_list if key}

    animal_records = _records(workbook["AnimalGroups"])
    wrong_animal_trials = sorted({
        _clean(row.get("trial_key")) for row in animal_records
        if _clean(row.get("trial_key")).casefold() not in trials
    })
    if wrong_animal_trials:
        raise ValueError(f"AnimalGroups 引用了不存在的试验编号：{', '.join(wrong_animal_trials)}")
    animal_key_list = [_clean(row.get("animal_group_key")) for row in animal_records]
    duplicate_animals = sorted(key for key, count in Counter(key.casefold() for key in animal_key_list if key).items() if count > 1)
    if duplicate_animals:
        raise ValueError(f"AnimalGroups.animal_group_key 重复：{', '.join(duplicate_animals)}")
    animal_keys = {key.casefold(): key for key in animal_key_list if key}

    component_long: list[tuple[str, tuple[str, str, str, str], Any]] = []
    for row in _records(workbook["CompositionValues"]):
        owner_input = _clean(row.get("ingredient_key"))
        owner = ingredient_keys.get(owner_input.casefold())
        if owner is None:
            raise ValueError(f"CompositionValues 引用了不存在的原料编号：{owner_input}")
        metric = _clean(row.get("raw_metric"))
        if not metric:
            raise ValueError("CompositionValues.raw_metric 不能为空")
        component_long.append((owner, (metric, "composition", _clean(row.get("unit")), _clean(row.get("basis"))), row.get("value")))
    component_order, component_values = _metric_occurrences(component_long)

    nutrition_long: list[tuple[str, tuple[str, str, str, str], Any]] = []
    for row in _records(workbook["NutritionValues"]):
        ingredient_input = _clean(row.get("ingredient_key"))
        animal_input = _clean(row.get("animal_group_key"))
        ingredient_key = ingredient_keys.get(ingredient_input.casefold())
        animal_key = animal_keys.get(animal_input.casefold())
        if ingredient_key is None or animal_key is None:
            raise ValueError(f"NutritionValues 引用了不存在的原料或动物组：{ingredient_input} / {animal_input}")
        metric = _clean(row.get("raw_metric"))
        context = _clean(row.get("measurement_context"))
        if not metric or not context:
            raise ValueError("NutritionValues 的 measurement_context 和 raw_metric 不能为空")
        owner = f"{ingredient_key}\u241f{animal_key}"
        nutrition_long.append((owner, (metric, context, _clean(row.get("unit")), _clean(row.get("basis"))), row.get("value")))
    nutrition_order, nutrition_values = _metric_occurrences(nutrition_long)

    batch = ContributionBatch(
        title=_clean(study.get("title")),
        doi=_clean(study.get("doi")),
        journal=_clean(study.get("journal")),
        publication_year=_clean(study.get("publication_year")),
        species=_clean(study.get("species")),
        country=_clean(study.get("country")),
        composition_metrics=[MetricSpec(metric, context, unit, basis) for metric, context, unit, basis, _ in component_order],
        ingredients=[
            IngredientGridRow(
                ingredient_key=_clean(row.get("ingredient_key")),
                raw_name=_clean(row.get("raw_name")),
                source_ingredient_no=_clean(row.get("source_ingredient_no")),
                feed_type=_clean(row.get("feed_type")),
                processing_technic=_clean(row.get("processing_technic")),
                place_of_origin=_clean(row.get("place_of_origin")),
                reported_basis=_clean(row.get("reported_basis")),
                dry_matter_value=_clean(row.get("dry_matter_value")),
                values=[component_values[_clean(row.get("ingredient_key"))].get(key, "") for key in component_order],
            )
            for row in ingredient_records
        ],
        animals=[
            AnimalGridRow(
                animal_group_key=_clean(row.get("animal_group_key")),
                trial_no=_clean(trials[_clean(row.get("trial_key")).casefold()].get("source_trial_no")),
                species=_clean(row.get("species")),
                breed=_clean(row.get("breed")),
                sex=_clean(row.get("sex")),
                animal_type=_clean(row.get("animal_type")),
                period=_clean(row.get("period")),
                age_days=_clean(row.get("age_days")),
                age_weeks=_clean(row.get("age_weeks")),
                body_weight_kg=_clean(row.get("body_weight_kg")),
            )
            for row in animal_records
        ],
        nutrition_metrics=[MetricSpec(metric, context, unit, basis) for metric, context, unit, basis, _ in nutrition_order],
        nutrition_rows=[
            NutritionGridRow(owner.split("\u241f", 1)[0], owner.split("\u241f", 1)[1], [values.get(key, "") for key in nutrition_order])
            for owner, values in nutrition_values.items()
        ],
    )
    warnings = [f"{item.location}: {item.message}" for item in issues if item.level != "error"]
    return ContributionFileResult(batch, source, "standard_excel", warnings)


def _legacy_observation_rows(ingredients, composition: bool):
    rows: list[tuple[str, tuple[str, str, str, str], Any]] = []
    for index, ingredient in enumerate(ingredients, start=1):
        owner = f"I{index:03d}" if composition else f"I{index:03d}\u241fA{index:03d}"
        for observation in ingredient.observations:
            is_composition = observation.measurement_context == "composition"
            if is_composition != composition:
                continue
            metric, unit = _split_metric_unit(observation.raw_metric)
            context = observation.measurement_context or "UNSPECIFIED_NUTRITION"
            basis = ingredient.reported_basis if context in {"composition", "energy"} else ""
            rows.append((owner, (metric, context, unit, basis), observation.value))
    return rows


def _legacy_excel_batch(workbook, source: Path) -> ContributionFileResult:
    worksheet = _find_sheet(workbook, INGREDIENT_SHEET_NAMES)
    if worksheet is None:
        raise ValueError("未找到成分及营养价值数据 sheet")
    metadata = _study_metadata(workbook, source)
    parsed = parse_legacy_ingredient_sheet(worksheet)
    ingredients = [item for item in parsed.ingredients if item.raw_name]
    component_order, component_values = _metric_occurrences(_legacy_observation_rows(ingredients, True))
    nutrition_order, nutrition_values = _metric_occurrences(_legacy_observation_rows(ingredients, False))

    ingredient_rows: list[IngredientGridRow] = []
    animal_rows: list[AnimalGridRow] = []
    for index, ingredient in enumerate(ingredients, start=1):
        ingredient_key = f"I{index:03d}"
        ingredient_rows.append(IngredientGridRow(
            ingredient_key=ingredient_key,
            raw_name=ingredient.raw_name,
            source_ingredient_no=ingredient.source_ingredient_no,
            feed_type=_clean(_legacy_field(ingredient.ingredient_fields, "feedtype")),
            processing_technic=_clean(_legacy_field(ingredient.ingredient_fields, "processing technic")),
            place_of_origin=_clean(_legacy_field(ingredient.ingredient_fields, "place of origin")),
            reported_basis=ingredient.reported_basis,
            dry_matter_value=_clean(_legacy_field(ingredient.ingredient_fields, "dry matter")),
            values=[component_values[ingredient_key].get(key, "") for key in component_order],
        ))
        owner = f"{ingredient_key}\u241fA{index:03d}"
        if owner in nutrition_values:
            raw_species = _clean(_legacy_field(ingredient.animal_fields, "animal", "animalname")) or _clean(metadata.get("species"))
            animal_rows.append(AnimalGridRow(
                animal_group_key=f"A{index:03d}",
                trial_no=ingredient.trial_no,
                species=raw_species or "not reported",
                breed=_clean(_legacy_field(ingredient.animal_fields, "breed")),
                sex=_clean(_legacy_field(ingredient.animal_fields, "gender")),
                animal_type=_clean(_legacy_field(ingredient.animal_fields, "animaltype", "type")),
                period=_clean(_legacy_field(ingredient.animal_fields, "period")),
                age_days=_clean(_legacy_field(ingredient.animal_fields, "days")),
                age_weeks=_clean(_legacy_field(ingredient.animal_fields, "weeks")),
                body_weight_kg=_clean(_legacy_field(ingredient.animal_fields, "weight")),
            ))

    batch = ContributionBatch(
        title=_clean(metadata.get("title")),
        doi=_clean(metadata.get("doi")),
        journal=_clean(metadata.get("journal")),
        publication_year=_clean(metadata.get("publication_year")),
        species=_clean(metadata.get("species")),
        country=_clean(metadata.get("country")),
        composition_metrics=[MetricSpec(metric, context, unit, basis) for metric, context, unit, basis, _ in component_order],
        ingredients=ingredient_rows,
        animals=animal_rows,
        nutrition_metrics=[MetricSpec(metric, context, unit, basis) for metric, context, unit, basis, _ in nutrition_order],
        nutrition_rows=[
            NutritionGridRow(owner.split("\u241f", 1)[0], owner.split("\u241f", 1)[1], [values.get(key, "") for key in nutrition_order])
            for owner, values in nutrition_values.items()
        ],
    )
    warnings = []
    unspecified = sum(
        observation.measurement_context is None
        for ingredient in ingredients
        for observation in ingredient.observations
    )
    if unspecified:
        warnings.append(f"{unspecified} 条营养价值缺少 AID/SID/ATTD/STTD 等明确语境，请在核对表中确认")
    return ContributionFileResult(batch, source, "legacy_excel", warnings)


def load_contribution_file(path: str | Path) -> ContributionFileResult:
    source = Path(path).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if source.name.startswith("~$"):
        raise ValueError("Excel 临时锁文件不能导入")
    if source.suffix.casefold() == ".json":
        return ContributionFileResult(batch_from_payload(read_payload(source)), source, "json")
    if source.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("仅支持 .xlsx、.xlsm 或 .json 文件")

    workbook = load_workbook(source, read_only=False, data_only=True)
    try:
        if is_bilingual_template(workbook):
            return ContributionFileResult(
                batch_from_bilingual_workbook(workbook, source),
                source,
                "bilingual_two_sheet_excel",
            )
        standard_overlap = set(STANDARD_SHEETS).intersection(workbook.sheetnames)
        if set(STANDARD_SHEETS).issubset(workbook.sheetnames) or len(standard_overlap) >= 3:
            return _standard_excel_batch(workbook, source)
        return _legacy_excel_batch(workbook, source)
    finally:
        workbook.close()
