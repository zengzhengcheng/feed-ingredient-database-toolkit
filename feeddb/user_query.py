from __future__ import annotations

import csv
import json
import re
import sqlite3
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .normalization import register_sqlite_functions


@dataclass(frozen=True)
class QueryResult:
    columns: tuple[str, ...]
    rows: list[tuple[object, ...]]
    truncated: bool = False


_LEADING_COMMENTS = re.compile(r"\A(?:\s+|--[^\n]*(?:\n|\Z)|/\*.*?\*/)*", re.DOTALL)


def validate_readonly_sql(sql: str) -> str:
    """Accept one SELECT/WITH statement and reject obvious multi-statement input."""
    statement = sql.strip()
    if not statement:
        raise ValueError("查询规则不能为空。")
    if statement.endswith(";"):
        statement = statement[:-1].rstrip()
    if ";" in statement:
        raise ValueError("一次只能执行一条只读 SQL 语句。")
    visible = _LEADING_COMMENTS.sub("", statement, count=1).lstrip()
    first_keyword = visible.split(None, 1)[0].casefold() if visible else ""
    if first_keyword not in {"select", "with"}:
        raise ValueError("只允许 SELECT 或 WITH 开头的只读查询。")
    return statement


def readonly_connection(database_path: str | Path) -> sqlite3.Connection:
    path = Path(database_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"数据库文件不存在：{path}")
    uri = f"file:{quote(path.as_posix(), safe='/:')}?mode=ro"
    conn = sqlite3.connect(uri, uri=True)
    conn.row_factory = sqlite3.Row
    register_sqlite_functions(conn)
    conn.execute("PRAGMA query_only = ON")

    denied_names = (
        "SQLITE_INSERT", "SQLITE_UPDATE", "SQLITE_DELETE", "SQLITE_CREATE_INDEX",
        "SQLITE_CREATE_TABLE", "SQLITE_CREATE_TEMP_INDEX", "SQLITE_CREATE_TEMP_TABLE",
        "SQLITE_CREATE_TEMP_TRIGGER", "SQLITE_CREATE_TEMP_VIEW", "SQLITE_CREATE_TRIGGER",
        "SQLITE_CREATE_VIEW", "SQLITE_DROP_INDEX", "SQLITE_DROP_TABLE",
        "SQLITE_DROP_TEMP_INDEX", "SQLITE_DROP_TEMP_TABLE", "SQLITE_DROP_TEMP_TRIGGER",
        "SQLITE_DROP_TEMP_VIEW", "SQLITE_DROP_TRIGGER", "SQLITE_DROP_VIEW",
        "SQLITE_ALTER_TABLE", "SQLITE_REINDEX", "SQLITE_ANALYZE", "SQLITE_ATTACH",
        "SQLITE_DETACH", "SQLITE_CREATE_VTABLE", "SQLITE_DROP_VTABLE", "SQLITE_PRAGMA",
    )
    denied_actions = {getattr(sqlite3, name) for name in denied_names if hasattr(sqlite3, name)}

    def authorize(action: int, _arg1: str | None, _arg2: str | None, _db: str | None, _trigger: str | None) -> int:
        return sqlite3.SQLITE_DENY if action in denied_actions else sqlite3.SQLITE_OK

    conn.set_authorizer(authorize)
    return conn


def execute_readonly_query(
    database_path: str | Path,
    sql: str,
    parameters: Sequence[object] | None = None,
    max_rows: int = 10_000,
    timeout_seconds: float = 30.0,
) -> QueryResult:
    """Run a bounded read-only query against an existing SQLite database."""
    statement = validate_readonly_sql(sql)
    if max_rows < 1:
        raise ValueError("max_rows 必须大于 0。")
    if max_rows > 1_000_000:
        raise ValueError("单次查询最多导出 1,000,000 行。")
    deadline = time.monotonic() + max(float(timeout_seconds), 0.1)
    with readonly_connection(database_path) as conn:
        conn.set_progress_handler(lambda: int(time.monotonic() > deadline), 1000)
        cursor = conn.execute(statement, tuple(parameters or ()))
        if cursor.description is None:
            raise ValueError("查询没有返回数据表。")
        fetched = cursor.fetchmany(max_rows + 1)
        columns = tuple(item[0] for item in cursor.description)
    truncated = len(fetched) > max_rows
    return QueryResult(columns, [tuple(row) for row in fetched[:max_rows]], truncated)


def write_query_result(
    output_path: str | Path,
    columns: Sequence[str],
    rows: Iterable[Sequence[object]],
) -> Path:
    """Write query rows as UTF-8 CSV, JSON, or XLSX according to the file suffix."""
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    materialized = [tuple(row) for row in rows]
    suffix = destination.suffix.casefold()
    if suffix == ".csv":
        with destination.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(columns)
            writer.writerows(materialized)
    elif suffix == ".json":
        records = [dict(zip(columns, row)) for row in materialized]
        destination.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    elif suffix == ".xlsx":
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet("QueryResult")
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        header_fill = PatternFill("solid", fgColor="0F766E")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_cells = []
        for column in columns:
            cell = WriteOnlyCell(worksheet, value=column)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            header_cells.append(cell)
        worksheet.append(header_cells)
        worksheet.row_dimensions[1].height = 32
        for row in materialized:
            worksheet.append(list(row))
        last_column = get_column_letter(len(columns))
        worksheet.auto_filter.ref = f"A1:{last_column}{len(materialized) + 1}"
        wide_tokens = ("title", "source_workbook", "source_sheet", "rule_reference", "note", "treatment")
        for index, column in enumerate(columns, start=1):
            values = [column, *(row[index - 1] for row in materialized)]
            longest = max((len(str(value)) for value in values if value is not None), default=len(column))
            maximum = 48 if any(token in column.casefold() for token in wide_tokens) else 32
            worksheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), maximum)
        workbook.save(destination)
    else:
        raise ValueError("导出文件扩展名必须是 .csv、.json 或 .xlsx。")
    return destination


def export_readonly_query(
    database_path: str | Path,
    sql: str,
    output_path: str | Path,
    parameters: Sequence[object] | None = None,
    max_rows: int = 1_000_000,
) -> dict[str, object]:
    result = execute_readonly_query(database_path, sql, parameters, max_rows=max_rows)
    destination = write_query_result(output_path, result.columns, result.rows)
    return {
        "output": str(destination),
        "rows": len(result.rows),
        "columns": len(result.columns),
        "truncated": result.truncated,
    }
